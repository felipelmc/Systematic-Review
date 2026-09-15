"""Funil formal declarado: filtros sequenciais sobre dados/registros_unicos.csv que etiquetam por padrão.

USO
    python3 rs.py filtrar --config 01-busca/filtros_v1.json [--ancoras 00-protocolo/ancoras_validacao.csv]
    python3 rs.py filtrar --config 01-busca/filtros_v1.json --amostra-elusao 300 --semente 7
    python3 rs.py filtrar --calcular-elusao 02-triagem/validacao/elusao_filtros_v1_cega.xlsx [--desenho ...]

CONFIGURAÇÃO (JSON)
    {
      "versao": "filtros_v1",                 # opcional; padrão = nome do arquivo
      "previsto_no_protocolo": false,         # vale para todos os filtros; cada filtro pode sobrescrever
      "filtros": [
        {"nome": "ano", "tipo": "ano", "min": 2001, "max": 2025},
        {"nome": "tipo", "tipo": "tipo", "recusar": ["editorial", "errata"]},
        {"nome": "idioma", "tipo": "idioma", "aceitar": ["pt", "en", "es"]},
        {"nome": "metodo", "tipo": "dicionario", "dicionario": "metodo_pt_en_es",
         "grupos": ["quantitativo", "experimental"], "campos": ["titulo", "resumo", "palavras_chave"],
         "modo": "etiquetar"}
      ]
    }
    tipo: ano (min/max inclusivos) | tipo (aceitar ou recusar valores de
    esquema.TIPOS_PUBLICACAO) | idioma (aceitar ou recusar códigos ISO 639-1) |
    dicionario (passa quem cita ao menos um termo dos grupos escolhidos).
    modo: "etiquetar" (padrão; quem falha segue no funil com etiqueta) ou
    "excluir" (quem falha sai do funil e não é avaliado pelos filtros seguintes).
    `dicionario` aceita o nome de um arquivo em assets/dicionarios/ ou um caminho.

POR QUE ASSIM (Apêndice D da base de conhecimento, item 1; references/03-organizacao-triagem.md §4)
    Dicionários e filtros formais ranqueiam/etiquetam; excluir é exceção:
      - modo "excluir" só com "previsto_no_protocolo": true (senão exit 2);
      - dicionário em modo "excluir" exige também "validacao_elusao" (referência à
        amostra de elusão que mediu o que o filtro perde), porque o funil de exemplo do
        REFIS excluía por dicionário sem medir recall. Se a referência for o JSON de
        métricas gerado por `--calcular-elusao`, ele é conferido: o filtro precisa
        constar nele e o dicionário, os grupos e os campos precisam ser os mesmos
        da amostra (assinatura sha256); senão exit 2. Referências em texto livre
        continuam aceitas, com aviso de que não foram verificadas;
      - idioma em modo "excluir", e tipo em modo "excluir" que atinja literatura
        cinzenta (tese, dissertação, relatório, evento, preprint, livro,
        capítulo), exigem "justificativa".
    Campo ausente -> resultado `sem_dado` e o registro segue. Registro sem resumo
    (ou com resumo truncado) nunca é excluído nem etiquetado por filtro textual:
    sem termo no título/palavras-chave, fica `sem_dado`.

CONJUNTO ATIVO
    Linhas de registros_unicos.csv com a flag `busca_inativa` (clusters que só
    tinham registros de buscas substituídas; ver dedup.py) ficam fora do funil, das
    âncoras e das amostras. O total ignorado vai para o log (`n_inativos_ignorados`).

CASAMENTO DE TERMOS
    O script de filtragem do REFIS usava regex sem fronteira de palavra: "sem" (SEM) casava
    em 46 de 181 resumos, "ols" casava dentro de "bolsa", "refis" dentro de
    "Refiscalizar". Aqui cada termo vira `(?<!\\w)termo(?!\\w)` sobre o texto com
    acentos dobrados (NFKD -> ASCII); `*` é truncamento (`\\w*`); espaço e hífen
    são intercambiáveis ("quasi-experiment*" casa "quasi experimental"); termos
    com sensivel_caixa=1 (siglas: SEM, OLS, PCA, RCT...) só casam com a caixa
    exata e são ignorados em campos inteiramente em maiúsculas, onde a caixa não
    distingue sigla de palavra.

SAÍDAS
    02-triagem/filtro_formal.csv (id_rs, filtro, resultado, detalhe), uma linha
    por registro por filtro que o avaliou; resultado passa|etiqueta|exclui|sem_dado.
    02-triagem/filtro_formal_contagens.json (entrada/saída por filtro; base do PRISMA).
    Evento `filtro_formal` e `versoes_ativas.filtros` no estado.

ÂNCORAS E RECALL RELATIVO (Apêndice D da base de conhecimento, item 9; references/02-busca.md §7)
    CSV com doi e/ou titulo + ano (colunas extras como id/chave são aceitas).
    Casamento por DOI normalizado ou título normalizado (título ou título
    alternativo) + ano igual. Âncora excluída por algum filtro -> saídas gravadas
    para auditoria e exit 2. Âncoras não encontradas no corpus são relatadas.
    Recall relativo em 01-busca/recall_ancoras.json: cada âncora encontrada é
    ligada às buscas que a acharam (id_rs -> ids_registro -> busca_id, via
    registros.csv; sem ele, pelo prefixo do id_registro). Relata:
      - `combinado`: âncoras achadas por ao menos uma busca de método `base`
        (as strings nas bases), sobre as âncoras indexadas;
      - `combinado_todos_metodos`: idem, contando citação, cinzenta e manual;
      - `por_busca`: R_b de cada busca ativa (inclusive as que não acharam nada);
      - `por_base` (só com a coluna `indexada_em`): R por item declarado.
    Coluna opcional `indexada_em`: nomes de base (openalex|scopus|wos, o recomendado:
    sobrevive a substituições de busca) ou ids de busca (B01|B02), separados por `|`,
    `;` ou `,`; vazio ou `nenhuma` = âncora não indexada em nenhuma base (problema de
    fonte, fica fora dos denominadores e vai para `nao_indexadas`). Id de busca
    substituída conta como a busca que a substituiu (cadeia `substituida_por` em
    estado.buscas), com aviso sugerindo o nome da base. Com a coluna, o denominador de
    R_b só tem as âncoras indexadas naquela busca/base; sem ela, todas as âncoras com
    chave de casamento. O arquivo de âncoras é planilha humana: CSV (`,` `;` tab;
    UTF-8, UTF-16 ou cp1252) ou xlsx, pelo leitor único de planilhas.py.
    IC 95% de Clopper-Pearson. Referência de ordem de grandeza: 0,95 (não é limiar;
    nenhum código de saída depende do recall). Imprime ids, nunca títulos.

AMOSTRA DE ELUSÃO DO DICIONÁRIO
    `--amostra-elusao N --semente S`: depois de aplicar os filtros (rode em modo
    etiquetar), sorteia N registros entre os etiquetados ou excluídos pelos filtros de
    dicionário (todos, se houver menos de N) e grava a planilha CEGA
    02-triagem/validacao/elusao_<versao>_cega.xlsx (título, resumo e metadados;
    sem resultado do filtro, termos ou etiquetas; ordem sorteada) e o desenho
    elusao_<versao>_desenho.json (população, semente, ids, assinatura dos
    filtros). Reexecutar com os mesmos parâmetros não sorteia de novo nem
    sobrescreve a planilha (que pode estar sendo codificada); parâmetros diferentes
    para a mesma versão -> exit 1.
    `--calcular-elusao PLANILHA`: lê as decisões humanas (decisao_h1, decisao_h2
    opcional, decisao_consenso; vazio = não revisado, fora das contas; xlsx ou CSV em
    qualquer delimitador e codificação aceitos por planilhas.py), calcula a
    taxa de elusão = relevantes / revisados com IC 95% de Clopper-Pearson (regra
    conservadora: incluir ou incerto contam como relevante perdido; a taxa estrita,
    só incluir, também é relatada) e os relevantes perdidos estimados na população
    (taxa x N). Grava elusao_<versao>_metricas.json, registra `validacao_calculada`
    (finalidade "elusao") e preenche "validacao_elusao" dos filtros amostrados na
    configuração, se ela estiver dentro do projeto, não congelada e com a mesma
    definição do filtro. Não muda o modo: excluir continua exigindo previsão no
    protocolo, e a decisão sobre a taxa aceitável é humana.
"""

import csv
import io
import json
import random
import re
import sys
from pathlib import Path

from . import esquema, estado, normalizar, planilhas
from .importar import buscas as _buscas

ETAPA = "05_organizacao"
ARQ_CONTAGENS = "02-triagem/filtro_formal_contagens.json"
ARQ_RECALL_ANCORAS = esquema.ARQ_RECALL_ANCORAS
DIR_VALIDACAO = esquema.DIR_VALIDACAO
DIR_DICIONARIOS = Path(__file__).resolve().parents[2] / "assets" / "dicionarios"
DICIONARIO_PADRAO = "metodo_pt_en_es"
COLUNAS_DICIONARIO = ["termo", "idioma", "grupo", "sensivel_caixa"]
TIPOS_FILTRO = ("ano", "tipo", "idioma", "dicionario")
MODOS = ("etiquetar", "excluir")
CAMPOS_TEXTO_PADRAO = ["titulo", "resumo", "palavras_chave"]
CAMPOS_TEXTO_PERMITIDOS = {"titulo", "titulo_alt", "resumo", "palavras_chave", "veiculo"}
TIPOS_CINZENTA = {"tese", "dissertacao", "relatorio", "evento", "preprint", "livro", "capitulo"}
_TRACOS = re.compile(r"[‐-―−]")
MAX_TERMOS_DETALHE = 8
TIPO_METRICAS_ELUSAO = "elusao_filtro"
TIPO_DESENHO_ELUSAO = "elusao_filtro_desenho"
COLUNAS_REGISTRO_ELUSAO = ["ordem", "id_rs", "titulo", "resumo", "palavras_chave", "ano", "veiculo", "idioma",
                           "tipo_publicacao"]
COLUNAS_CODIGO_ELUSAO = ["decisao_h1", "decisao_h2", "decisao_consenso"]
CODIGOS_ELUSAO = ["incluir", "excluir", "incerto", "nao_revisado"]
_SEM_INDEXACAO = {"", "nenhuma", "nenhum", "none", "nao", "não", "-", "na"}


class ErroFiltro(RuntimeError):
    """Erro de uso ou de dados (exit 1)."""


class ErroMetodologico(RuntimeError):
    """Configuração que viola uma regra metodológica (exit 2)."""


# ---------------------------------------------------------------------------
# E/S
# ---------------------------------------------------------------------------
def _ler_csv(caminho):
    with open(caminho, encoding="utf-8-sig", newline="") as f:
        leitor = csv.DictReader(f)
        linhas = [{k: (v if v is not None else "") for k, v in l.items() if k is not None} for l in leitor]
        return list(leitor.fieldnames or []), linhas


def _gravar_atomico(caminho, texto):
    """Grava só se mudou (sha256 estável em reexecuções), via estado.escrever_atomico (0666 menos a umask)."""
    caminho = Path(caminho)
    if caminho.exists() and caminho.read_bytes() == texto.encode("utf-8"):
        return False
    estado.escrever_atomico(caminho, texto, newline="")
    return True


def _ler_humana(caminho, obrigatorias=(), alternativas=None, aba=None):
    """Planilha humana pelo leitor único (planilhas.py); chaves em minúsculas. Erros viram ErroFiltro/ErroDependencia."""
    try:
        colunas, linhas, info = planilhas.ler_tabela_humana(caminho, obrigatorias, alternativas, aba=aba)
    except planilhas.ErroDependencia:
        raise
    except planilhas.ErroUso as e:
        raise ErroFiltro(str(e)) from None
    colunas = [c.strip().lower() for c in colunas]
    linhas = [{(k or "").strip().lower(): (v if v is not None else "") for k, v in l.items() if k} for l in linhas]
    return colunas, linhas, info


def _texto_csv(colunas, linhas):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=colunas, extrasaction="ignore", lineterminator="\n")
    w.writeheader()
    for l in linhas:
        w.writerow({c: l.get(c, "") for c in colunas})
    return buf.getvalue()


def _json(dados):
    return json.dumps(dados, ensure_ascii=False, indent=2) + "\n"


# ---------------------------------------------------------------------------
# Dicionário e casamento
# ---------------------------------------------------------------------------
def dobrar(valor):
    """Texto limpo com acentos dobrados para ASCII, preservando a caixa (siglas dependem dela)."""
    return normalizar.ascii_fold(_TRACOS.sub("-", normalizar.texto(valor)))


class Termo:
    """Termo compilado. `sonda` é um trecho literal usado como pré-filtro barato por substring."""

    __slots__ = ("termo", "idioma", "grupo", "sensivel", "regex", "sonda")

    def __init__(self, termo, idioma, grupo, sensivel):
        self.termo, self.idioma, self.grupo, self.sensivel = termo, idioma, grupo, sensivel
        base = dobrar(termo)
        if not sensivel:
            base = base.lower()
        tokens = [t for t in re.split(r"[\s\-]+", base) if t]
        if not tokens or not any(t.strip("*") for t in tokens):
            raise ErroFiltro(f"termo vazio ou só com '*' no dicionário: '{termo}'")
        partes = ["\\w*".join(re.escape(p) for p in t.split("*")) for t in tokens]
        self.regex = re.compile(r"(?<!\w)" + r"[\s\-]+".join(partes) + r"(?!\w)")
        literais = [p for t in tokens for p in t.split("*") if p]
        self.sonda = max(literais, key=len)


def carregar_dicionario(caminho):
    cabecalho, linhas = _ler_csv(caminho)
    faltando = [c for c in COLUNAS_DICIONARIO if c not in cabecalho]
    if faltando:
        raise ErroFiltro(f"{caminho}: colunas ausentes {faltando} (esperado {','.join(COLUNAS_DICIONARIO)})")
    termos = []
    for n, l in enumerate(linhas, start=2):
        termo = (l.get("termo") or "").strip()
        if not termo:
            continue
        sens = (l.get("sensivel_caixa") or "0").strip()
        if sens not in {"0", "1"}:
            raise ErroFiltro(f"{caminho}:{n}: sensivel_caixa deve ser 0 ou 1")
        termos.append(Termo(termo, (l.get("idioma") or "").strip(), (l.get("grupo") or "").strip(), sens == "1"))
    if not termos:
        raise ErroFiltro(f"{caminho}: dicionário sem termos")
    return termos


def _todo_maiusculo(texto):
    letras = [c for c in texto if c.isalpha()]
    return len(letras) >= 8 and sum(c.isupper() for c in letras) / len(letras) > 0.8


def casar(termos, valores):
    """Termos que casam em algum dos textos. Devolve lista ordenada de (grupo, termo), sem repetição."""
    achados = set()
    for valor in valores:
        original = dobrar(valor)
        if not original:
            continue
        minusculo = original.lower()
        caixa_alta = _todo_maiusculo(original)
        for t in termos:
            if t.sensivel:
                if caixa_alta or t.sonda not in original or not t.regex.search(original):
                    continue
            elif t.sonda not in minusculo or not t.regex.search(minusculo):
                continue
            achados.add((t.grupo, t.termo))
    return sorted(achados)


# ---------------------------------------------------------------------------
# Configuração
# ---------------------------------------------------------------------------
def _resolver(caminho, *bases):
    p = Path(caminho)
    if p.is_absolute():
        return p
    for base in bases:
        if base is not None and (Path(base) / p).exists():
            return Path(base) / p
    return p


def resolver_dicionario(valor, dir_config, raiz):
    valor = valor or DICIONARIO_PADRAO
    if "/" not in valor and "\\" not in valor and not valor.endswith(".csv"):
        candidato = DIR_DICIONARIOS / f"{valor}.csv"
        if candidato.exists():
            return candidato
    p = _resolver(valor, Path.cwd(), dir_config, raiz, DIR_DICIONARIOS)
    if not p.exists():
        raise ErroFiltro(f"dicionário não encontrado: {valor}")
    return p


def _lista(f, chave, nome):
    v = f.get(chave)
    if not isinstance(v, list) or not v or not all(isinstance(x, str) and x.strip() for x in v):
        raise ErroFiltro(f"filtro '{nome}': '{chave}' deve ser lista não vazia de textos")
    return [x.strip().lower() for x in v]


def assinatura_filtro(f):
    """sha256 do que define o que um filtro de dicionário deixa passar (conteúdo do dicionário, grupos, campos)."""
    return estado.sha256_texto(json.dumps({
        "tipo": "dicionario", "dicionario_sha256": f["dicionario_sha256"],
        "grupos": sorted({t.grupo for t in f["termos"]}), "campos": list(f["campos"]),
    }, sort_keys=True))


def validar_config(config, dir_config, raiz, checar=True):
    """Normaliza a configuração e aplica as regras metodológicas antes de tocar em qualquer arquivo.

    `checar=False` só normaliza (usado para recalcular assinaturas ao habilitar a validação de elusão).
    """
    if not isinstance(config, dict) or not isinstance(config.get("filtros"), list) or not config["filtros"]:
        raise ErroFiltro("configuração precisa de uma lista 'filtros' não vazia")
    previsto_global = config.get("previsto_no_protocolo", False)
    nomes, filtros = set(), []
    for pos, bruto in enumerate(config["filtros"], start=1):
        if not isinstance(bruto, dict):
            raise ErroFiltro(f"filtro {pos}: deve ser um objeto")
        tipo = bruto.get("tipo")
        if tipo not in TIPOS_FILTRO:
            raise ErroFiltro(f"filtro {pos}: tipo '{tipo}' inválido (use {', '.join(TIPOS_FILTRO)})")
        nome = str(bruto.get("nome") or tipo)
        if nome in nomes:
            raise ErroFiltro(f"nome de filtro repetido: '{nome}'")
        nomes.add(nome)
        modo = bruto.get("modo", "etiquetar")
        if modo not in MODOS:
            raise ErroFiltro(f"filtro '{nome}': modo '{modo}' inválido (use etiquetar|excluir)")
        f = {"nome": nome, "tipo": tipo, "modo": modo,
             "previsto_no_protocolo": bruto.get("previsto_no_protocolo", previsto_global) is True,
             "justificativa": str(bruto.get("justificativa") or "").strip(),
             "validacao_elusao": str(bruto.get("validacao_elusao") or "").strip()}
        if tipo == "ano":
            for lim in ("min", "max"):
                v = bruto.get(lim)
                if v is not None and (not isinstance(v, int) or isinstance(v, bool)):
                    raise ErroFiltro(f"filtro '{nome}': '{lim}' deve ser inteiro")
                f[lim] = v
            if f["min"] is None and f["max"] is None:
                raise ErroFiltro(f"filtro '{nome}': informe 'min' e/ou 'max'")
        elif tipo in ("tipo", "idioma"):
            tem_aceitar, tem_recusar = "aceitar" in bruto, "recusar" in bruto
            if tem_aceitar == tem_recusar:
                raise ErroFiltro(f"filtro '{nome}': use exatamente um de 'aceitar' ou 'recusar'")
            chave_lista = "aceitar" if tem_aceitar else "recusar"
            valores = _lista(bruto, chave_lista, nome)
            if tipo == "tipo":
                invalidos = sorted(set(valores) - set(esquema.TIPOS_PUBLICACAO))
                if invalidos:
                    raise ErroFiltro(f"filtro '{nome}': tipos desconhecidos {invalidos}")
            f[chave_lista] = set(valores)
        else:
            f["dicionario"] = resolver_dicionario(bruto.get("dicionario"), dir_config, raiz)
            f["termos"] = carregar_dicionario(f["dicionario"])
            grupos_disp = {t.grupo for t in f["termos"]}
            grupos = bruto.get("grupos")
            if grupos is not None:
                grupos = set(_lista(bruto, "grupos", nome))
                desconhecidos = sorted(grupos - grupos_disp)
                if desconhecidos:
                    raise ErroFiltro(f"filtro '{nome}': grupos desconhecidos {desconhecidos}; "
                                     f"disponíveis {sorted(grupos_disp)}")
                f["termos"] = [t for t in f["termos"] if t.grupo in grupos]
            campos = bruto.get("campos", CAMPOS_TEXTO_PADRAO)
            if not isinstance(campos, list) or not campos or set(campos) - CAMPOS_TEXTO_PERMITIDOS:
                raise ErroFiltro(f"filtro '{nome}': 'campos' deve ser lista dentro de {sorted(CAMPOS_TEXTO_PERMITIDOS)}")
            f["campos"] = list(campos)
            f["dicionario_sha256"] = estado.sha256_arquivo(f["dicionario"])
            f["assinatura"] = assinatura_filtro(f)
        if modo == "excluir" and checar:
            checar_exclusao(f)
            if tipo == "dicionario":
                f["elusao_verificada"] = conferir_validacao_elusao(f, dir_config, raiz)
        filtros.append(f)
    return filtros


def checar_exclusao(f):
    """Item 1 do Apêndice D da base de conhecimento: exclusão é exceção prevista, validada e justificada."""
    nome = f["nome"]
    if not f["previsto_no_protocolo"]:
        raise ErroMetodologico(
            f"filtro '{nome}' em modo excluir sem \"previsto_no_protocolo\": true. Filtros formais etiquetam "
            "por padrão; exclua só o que o protocolo congelado prevê (senão use modo etiquetar).")
    if f["tipo"] == "dicionario" and not f["validacao_elusao"]:
        raise ErroMetodologico(
            f"filtro '{nome}': exclusão por dicionário exige \"validacao_elusao\" (referência à amostra de "
            "elusão que mediu o que o filtro perde). Rode primeiro em modo etiquetar e valide "
            "(--amostra-elusao N --semente S, depois --calcular-elusao).")
    if f["tipo"] == "idioma" and not f["justificativa"]:
        raise ErroMetodologico(f"filtro '{nome}': exclusão por idioma exige \"justificativa\" (relatada como limitação).")
    if f["tipo"] == "tipo":
        recusados = f["recusar"] if "recusar" in f else set(esquema.TIPOS_PUBLICACAO) - f["aceitar"]
        if recusados & TIPOS_CINZENTA and not f["justificativa"]:
            raise ErroMetodologico(
                f"filtro '{nome}': excluir {sorted(recusados & TIPOS_CINZENTA)} remove literatura cinzenta; "
                "exige \"justificativa\".")


def conferir_validacao_elusao(f, dir_config, raiz):
    """Confere o JSON de métricas de `--calcular-elusao` citado em validacao_elusao.

    Devolve um resumo das métricas, ou None quando a referência não é um arquivo de métricas
    verificável (texto livre, CSV, JSON inexistente ou de outro tipo), caso aceito com aviso.
    """
    ref = f["validacao_elusao"]
    if not ref.lower().endswith(".json"):
        return None
    caminho = _resolver(ref, raiz, dir_config, Path.cwd())
    if not caminho.is_file():
        return None
    try:
        metricas = json.loads(caminho.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if not isinstance(metricas, dict) or metricas.get("tipo") != TIPO_METRICAS_ELUSAO:
        return None
    medidos = {x.get("nome"): x for x in metricas.get("filtros") or [] if isinstance(x, dict)}
    if f["nome"] not in medidos:
        raise ErroMetodologico(f"filtro '{f['nome']}': {ref} mediu a elusão de {sorted(medidos)}, não deste filtro")
    if medidos[f["nome"]].get("assinatura") != f["assinatura"]:
        raise ErroMetodologico(
            f"filtro '{f['nome']}': dicionário, grupos ou campos mudaram desde a amostra de elusão de {ref}; "
            "a validação não vale para a definição atual. Gere e codifique uma nova amostra.")
    return {"arquivo": ref, "taxa_elusao": metricas.get("taxa_elusao"), "ic95": metricas.get("ic95"),
            "n_revisados": metricas.get("n_revisados"), "perdidos_estimados": metricas.get("perdidos_estimados")}


# ---------------------------------------------------------------------------
# Avaliação
# ---------------------------------------------------------------------------
def avaliar(f, linha):
    """Devolve (situacao, detalhe) com situacao em passa|falha|sem_dado."""
    tipo = f["tipo"]
    if tipo == "ano":
        ano = normalizar.ano(linha.get("ano"))
        if not ano:
            return "sem_dado", "ano_ausente"
        ano = int(ano)
        ok = (f["min"] is None or ano >= f["min"]) and (f["max"] is None or ano <= f["max"])
        faixa = f"{f['min'] if f['min'] is not None else ''}..{f['max'] if f['max'] is not None else ''}"
        return ("passa" if ok else "falha"), f"ano={ano};faixa={faixa}"
    if tipo in ("tipo", "idioma"):
        if tipo == "tipo":
            valor = normalizar.texto(linha.get("tipo_publicacao")).lower()
            if not valor or (valor == "outro" and not normalizar.texto(linha.get("tipo_publicacao_orig"))):
                return "sem_dado", "tipo_ausente"
        else:
            valor = normalizar.idioma(linha.get("idioma")) or normalizar.texto(linha.get("idioma")).lower()
            if not valor:
                return "sem_dado", "idioma_ausente"
        ok = valor in f["aceitar"] if "aceitar" in f else valor not in f["recusar"]
        return ("passa" if ok else "falha"), f"{tipo}={valor}"
    achados = casar(f["termos"], [linha.get(c) for c in f["campos"]])
    if achados:
        grupos = sorted({g for g, _ in achados})
        termos = sorted({t for _, t in achados})
        extra = f"(+{len(termos) - MAX_TERMOS_DETALHE})" if len(termos) > MAX_TERMOS_DETALHE else ""
        return "passa", f"grupos={','.join(grupos)};termos={','.join(termos[:MAX_TERMOS_DETALHE])}{extra}"
    # Vale mesmo quando o filtro não lê o resumo: sem ele o registro tem informação
    # insuficiente para ser reprovado por texto, e a triagem T/A decide (como "incerto").
    if not normalizar.texto(linha.get("resumo")):
        return "sem_dado", "sem_resumo;nenhum_termo_no_restante"
    if normalizar.texto(linha.get("resumo_truncado")) == "1":
        return "sem_dado", "resumo_truncado;nenhum_termo_no_restante"
    return "falha", "nenhum_termo"


def aplicar_filtros(unicos, filtros):
    """Aplica os filtros em sequência. Devolve (linhas de filtro_formal, contagens, exclusões, etiquetas)."""
    ativos = list(unicos)
    saida, contagens = [], []
    excluidos, etiquetados = {}, {}
    for f in filtros:
        c = {"filtro": f["nome"], "tipo": f["tipo"], "modo": f["modo"], "entrada": len(ativos),
             "passa": 0, "etiqueta": 0, "exclui": 0, "sem_dado": 0}
        seguintes = []
        for linha in ativos:
            situacao, detalhe = avaliar(f, linha)
            if situacao == "falha":
                resultado = "exclui" if f["modo"] == "excluir" else "etiqueta"
            else:
                resultado = situacao
            c[resultado] += 1
            saida.append({"id_rs": linha["id_rs"], "filtro": f["nome"], "resultado": resultado, "detalhe": detalhe})
            if resultado == "exclui":
                excluidos[linha["id_rs"]] = f["nome"]
            else:
                if resultado == "etiqueta":
                    etiquetados.setdefault(linha["id_rs"], []).append(f["nome"])
                seguintes.append(linha)
        c["saida"] = len(seguintes)
        contagens.append(c)
        ativos = seguintes
    return saida, contagens, excluidos, etiquetados


# ---------------------------------------------------------------------------
# Âncoras
# ---------------------------------------------------------------------------
def ler_ancoras(caminho):
    """(colunas, linhas, avisos) do arquivo de âncoras: CSV (`,` `;` tab; UTF-8, UTF-16 ou cp1252) ou xlsx."""
    try:
        colunas, linhas, info = _ler_humana(caminho, alternativas=[["doi"], ["titulo"]])
    except ErroFiltro:
        raise ErroFiltro(f"{caminho}: âncoras precisam de coluna 'doi' e/ou 'titulo' (+ 'ano')") from None
    return colunas, linhas, info["avisos"]


def checar_ancoras(caminho, unicos, excluidos, etiquetados):
    _, linhas, _ = ler_ancoras(caminho)
    por_doi, por_titulo = {}, {}
    for u in unicos:
        d = normalizar.doi(u.get("doi"))
        if d:
            por_doi.setdefault(d, set()).add(u["id_rs"])
        ano = normalizar.ano(u.get("ano"))
        for campo in ("titulo", "titulo_alt"):
            tn = normalizar.titulo_normalizado(u.get(campo))
            if tn and ano:
                por_titulo.setdefault((tn, ano), set()).add(u["id_rs"])
    relatorio = []
    for n, l in enumerate(linhas, start=1):
        rotulo = normalizar.texto(l.get("id") or l.get("ancora") or l.get("chave")) or f"A{n:03d}"
        d = normalizar.doi(l.get("doi"))
        tn, ano = normalizar.titulo_normalizado(l.get("titulo")), normalizar.ano(l.get("ano"))
        ids = set()
        if d:
            ids |= por_doi.get(d, set())
        if tn and ano:
            ids |= por_titulo.get((tn, ano), set())
        item = {"ancora": rotulo, "ids_rs": sorted(ids)}
        if not ids:
            item["situacao"] = "nao_encontrada" if (d or (tn and ano)) else "sem_chave_de_casamento"
        elif any(i in excluidos for i in ids):
            item["situacao"] = "excluida"
            item["filtros"] = sorted({excluidos[i] for i in ids if i in excluidos})
        elif any(i in etiquetados for i in ids):
            item["situacao"] = "etiquetada"
            item["filtros"] = sorted({f for i in ids for f in etiquetados.get(i, [])})
        else:
            item["situacao"] = "mantida"
        relatorio.append(item)
    return relatorio


def _ic95(k, n):
    """IC 95% de Clopper-Pearson (validacao.clopper_pearson); None se n = 0 ou indisponível."""
    if not n:
        return None
    try:
        from .validacao import clopper_pearson
        lo, hi = clopper_pearson(k, n)
    except Exception:  # noqa: BLE001 - IC é complemento; a taxa pontual segue sem ele
        return None
    return [round(lo, 4), round(hi, 4)]


def _tokens_indexada(valor):
    tokens = set()
    for t in re.split(r"[|;,]", str(valor or "")):
        t = t.strip()
        if normalizar.ascii_fold(t).lower() in _SEM_INDEXACAO:
            continue
        tokens.add(t.upper() if _buscas.RE_BUSCA_ID.match(t.upper()) else t.lower())
    return tokens


def informacoes_buscas(raiz, registros, unicos):
    """busca_id -> {fonte, metodo, ativa} a partir do estado, de registros.csv e dos ids dos únicos."""
    info = {}
    contagem = {}
    for l in registros.values():
        b = l.get("busca_id") or _buscas.busca_de_id_registro(l.get("id_registro"))
        if b:
            c = contagem.setdefault(b, {"fonte": {}, "metodo": {}})
            for campo, valor in (("fonte", normalizar.texto(l.get("fonte")).lower()),
                                 ("metodo", normalizar.texto(l.get("metodo_identificacao")).lower())):
                if valor:
                    c[campo][valor] = c[campo].get(valor, 0) + 1
    for b, c in contagem.items():
        info[b] = {"fonte": max(c["fonte"], key=c["fonte"].get) if c["fonte"] else "",
                   "metodo": max(c["metodo"], key=c["metodo"].get) if c["metodo"] else _buscas.metodo_por_prefixo(b),
                   "ativa": True}
    try:
        est = estado.carregar_estado(raiz)
    except (estado.ErroProjeto, OSError, ValueError):
        est = {}
    for b in est.get("buscas", []) if isinstance(est, dict) else []:
        if not isinstance(b, dict) or not b.get("id"):
            continue
        d = info.setdefault(b["id"], {"fonte": "", "metodo": _buscas.metodo_por_prefixo(b["id"]), "ativa": True})
        d["fonte"] = normalizar.texto(b.get("fonte")).lower() or d["fonte"]
        d["metodo"] = normalizar.texto(b.get("metodo_identificacao")).lower() or d["metodo"]
        d["ativa"] = _buscas.busca_ativa(b)
        if b.get("substituida_por"):
            d["substituida_por"] = str(b["substituida_por"]).strip().upper()
    for u in unicos:
        for rid in normalizar.texto(u.get("ids_registro")).split("|"):
            b = (registros.get(rid) or {}).get("busca_id") or _buscas.busca_de_id_registro(rid)
            if b:
                info.setdefault(b, {"fonte": "", "metodo": _buscas.metodo_por_prefixo(b), "ativa": True})
    return info


def recall_ancoras(raiz, caminho, relatorio, unicos):
    """Recall relativo das âncoras por busca, combinado e (com indexada_em) por base. Ver docstring."""
    cabecalho, linhas, avisos_leitura = ler_ancoras(caminho)
    tem_indexada = "indexada_em" in cabecalho
    caminho_registros = Path(raiz) / esquema.ARQ_REGISTROS
    registros = {}
    if caminho_registros.exists():
        registros = {l["id_registro"]: l for l in _ler_csv(caminho_registros)[1] if l.get("id_registro")}
    info = informacoes_buscas(raiz, registros, unicos)
    ativas = sorted(b for b, d in info.items() if d["ativa"])
    de_base = {b for b in ativas if info[b]["metodo"] == "base"}
    por_id = {u["id_rs"]: u for u in unicos}
    avisos = list(avisos_leitura)

    itens = []
    for l, item in zip(linhas, relatorio):
        achadas = set()
        for id_rs in item["ids_rs"]:
            for rid in normalizar.texto(por_id.get(id_rs, {}).get("ids_registro")).split("|"):
                b = (registros.get(rid) or {}).get("busca_id") or _buscas.busca_de_id_registro(rid)
                if b and info.get(b, {}).get("ativa", True):
                    achadas.add(b)
        itens.append({"ancora": item["ancora"], "situacao": item["situacao"], "ids_rs": item["ids_rs"],
                      "buscas": sorted(achadas),
                      "indexada_em": sorted(_tokens_indexada(l.get("indexada_em"))) if tem_indexada else None})

    redirecionados = {}

    def cobertura(token):
        if token in info:
            if info[token]["ativa"]:
                return {token}
            # busca substituída: conta a busca que a substituiu (cadeia substituida_por), se ela está ativa
            atual, vistos = token, set()
            while atual in info and not info[atual]["ativa"] and info[atual].get("substituida_por") \
                    and atual not in vistos:
                vistos.add(atual)
                atual = info[atual]["substituida_por"]
            if atual in info and info[atual]["ativa"]:
                redirecionados[token] = atual
                return {atual}
            return set()
        return {b for b in ativas if info[b]["fonte"] == token}

    def bloco(denominador, achou):
        rotulos = [a["ancora"] for a in denominador]
        encontradas = [a["ancora"] for a in denominador if achou(a)]
        n, k = len(rotulos), len(encontradas)
        return {"n_ancoras": n, "n_encontradas": k, "recall": round(k / n, 4) if n else None, "ic95": _ic95(k, n),
                "perdidas": [r for r in rotulos if r not in set(encontradas)]}

    elegiveis = [a for a in itens if a["situacao"] != "sem_chave_de_casamento"]
    if tem_indexada:
        indexadas = [a for a in elegiveis if a["indexada_em"]]
        nao_indexadas = [a["ancora"] for a in elegiveis if not a["indexada_em"]]
    else:
        indexadas, nao_indexadas = elegiveis, []

    combinado = dict(bloco(indexadas, lambda a: bool(set(a["buscas"]) & de_base)), buscas=sorted(de_base))
    todos = dict(bloco(indexadas, lambda a: bool(a["buscas"])), buscas=ativas)
    por_busca = {}
    for b in ativas:
        if tem_indexada:
            den = [a for a in indexadas if any(b in cobertura(t) for t in a["indexada_em"])]
        else:
            den = elegiveis
        por_busca[b] = dict(bloco(den, lambda a, b=b: b in a["buscas"]), fonte=info[b]["fonte"],
                            metodo=info[b]["metodo"])
    por_base = None
    if tem_indexada:
        por_base = {}
        for token in sorted({t for a in indexadas for t in a["indexada_em"]}):
            cob = cobertura(token)
            if not cob:
                situacao = "inativa" if token in info else "sem busca correspondente"
                avisos.append(f"indexada_em '{token}': {situacao} (nem id de busca ativa nem fonte de uma busca)")
            den = [a for a in indexadas if token in a["indexada_em"]]
            por_base[token] = dict(bloco(den, lambda a, cob=cob: bool(set(a["buscas"]) & cob)), buscas=sorted(cob))
        for antigo, novo in sorted(redirecionados.items()):
            avisos.append(f"indexada_em '{antigo}': busca substituída por {novo}, contada como {novo}; escreva o nome "
                          f"da base ({info[novo]['fonte'] or 'openalex, scopus, wos...'}) em vez do id da busca, que muda "
                          "a cada substituição")
        fora = [a["ancora"] for a in indexadas
                if set(a["buscas"]) - set().union(*[cobertura(t) for t in a["indexada_em"]])]
        if fora:
            avisos.append(f"âncoras achadas em buscas fora do declarado em indexada_em (confira a coluna): {fora}")
        if nao_indexadas:
            avisos.append(f"{len(nao_indexadas)} âncoras sem indexação declarada ficam fora dos denominadores: "
                          f"{nao_indexadas}")
    rel = None
    try:
        rel = str(Path(caminho).resolve().relative_to(Path(raiz).resolve()))
    except ValueError:
        pass
    caminho_unicos = Path(raiz) / esquema.ARQ_UNICOS
    return {
        "ancoras_arquivo": rel or Path(caminho).name, "ancoras_sha256": estado.sha256_arquivo(caminho),
        "unicos_sha256": estado.sha256_arquivo(caminho_unicos) if caminho_unicos.exists() else None,
        "tem_indexada_em": tem_indexada, "n_ancoras": len(itens),
        "n_sem_chave_de_casamento": len(itens) - len(elegiveis),
        "combinado": combinado, "combinado_todos_metodos": todos, "nao_indexadas": nao_indexadas,
        "por_busca": por_busca, "por_base": por_base, "ancoras": itens, "avisos": avisos,
        "nota": "recall relativo ao conjunto de âncoras; 0,95 é referência de ordem de grandeza, não limiar",
    }


# ---------------------------------------------------------------------------
# Elusão do dicionário
# ---------------------------------------------------------------------------
def _nome_versao(versao):
    return re.sub(r"[^\w.-]+", "_", str(versao)).strip("_") or "filtros"


def _relativo(caminho, raiz):
    try:
        return str(Path(caminho).resolve().relative_to(Path(raiz).resolve()))
    except ValueError:
        return None


def planejar_amostra_elusao(raiz, filtros, linhas, n, semente, versao, caminho_config):
    """Sorteio (sem gravar) entre etiquetados/excluídos pelos dicionários; reusa desenho idêntico."""
    if n is None or n < 1:
        raise ErroFiltro("--amostra-elusao exige N >= 1")
    if semente is None:
        raise ErroFiltro("--amostra-elusao exige --semente (sorteio reprodutível)")
    dicionarios = [f for f in filtros if f["tipo"] == "dicionario"]
    if not dicionarios:
        raise ErroFiltro("--amostra-elusao exige ao menos um filtro de tipo dicionario na configuração")
    nomes = {f["nome"] for f in dicionarios}
    populacao = sorted({l["id_rs"] for l in linhas if l["filtro"] in nomes and l["resultado"] in ("etiqueta", "exclui")})
    if not populacao:
        raise ErroFiltro("nenhum registro etiquetado ou excluído pelos filtros de dicionário: não há elusão a medir")
    ids = random.Random(semente).sample(populacao, min(n, len(populacao)))
    base = _nome_versao(versao)
    rel_desenho = f"{DIR_VALIDACAO}/elusao_{base}{esquema.SUFIXO_DESENHO}"
    rel_planilha = f"{DIR_VALIDACAO}/elusao_{base}{esquema.SUFIXO_PLANILHA_CEGA}"
    desenho = {
        "tipo": TIPO_DESENHO_ELUSAO, "versao": versao,
        "config": _relativo(caminho_config, raiz) or str(Path(caminho_config).resolve()),
        "filtros": [{"nome": f["nome"], "modo": f["modo"],
                     "dicionario": _relativo(f["dicionario"], raiz) or f"assets/dicionarios/{f['dicionario'].name}",
                     "dicionario_sha256": f["dicionario_sha256"], "assinatura": f["assinatura"]}
                    for f in dicionarios],
        "resultados_amostrados": ["etiqueta", "exclui"],
        "populacao_n": len(populacao), "populacao_sha256": estado.sha256_texto("\n".join(populacao)),
        "n_solicitado": n, "n_amostra": len(ids), "semente": semente, "ids": ids,
        "planilha": rel_planilha, "gerado_em": estado.agora(),
    }
    reutilizada = False
    existente = Path(raiz) / rel_desenho
    if existente.exists():
        antigo = json.loads(existente.read_text(encoding="utf-8"))

        def chave(d):
            return ([(x.get("nome"), x.get("assinatura")) for x in d.get("filtros") or []],
                    d.get("populacao_sha256"), d.get("n_solicitado"), d.get("semente"))
        if chave(antigo) != chave(desenho):
            raise ErroFiltro(
                f"já existe amostra de elusão em {rel_desenho} com outros parâmetros (semente, tamanho, população "
                "ou definição dos filtros). Mova o desenho e a planilha para gerar outra, ou mude 'versao' na "
                "configuração.")
        desenho, reutilizada = antigo, True
    return {"desenho": desenho, "rel_desenho": rel_desenho, "rel_planilha": desenho.get("planilha", rel_planilha),
            "reutilizada": reutilizada}


def escrever_planilha_elusao(caminho, linhas, meta):
    """Planilha cega em xlsx: só dados bibliográficos + colunas vazias de decisão."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as e:
        raise ErroFiltro("openpyxl ausente: pip install openpyxl") from e
    colunas = COLUNAS_REGISTRO_ELUSAO + COLUNAS_CODIGO_ELUSAO + ["observacoes"]
    wb = Workbook()
    ws = wb.active
    ws.title = "codificacao"
    ws.append(colunas)
    for c in ws[1]:
        c.font = Font(bold=True)
    for linha in linhas:
        ws.append([linha.get(c) for c in colunas])
    lista = DataValidation(type="list", formula1='"' + ",".join(CODIGOS_ELUSAO) + '"', allow_blank=True)
    ws.add_data_validation(lista)
    ultima = max(2, len(linhas) + 1)
    for i, nome in enumerate(colunas, start=1):
        letra = ws.cell(row=1, column=i).column_letter
        if nome in COLUNAS_CODIGO_ELUSAO:
            lista.add(f"{letra}2:{letra}{ultima}")
        ws.column_dimensions[letra].width = {"titulo": 50, "resumo": 90, "palavras_chave": 30,
                                             "observacoes": 30}.get(nome, 14)
    for celulas in ws.iter_rows(min_row=2):
        for c in celulas:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    inst = wb.create_sheet("instrucoes")
    for texto in [
        "Amostra de elusão do funil formal. Para cada registro, decida pelos critérios de título/resumo (T/A) do "
        "protocolo: incluir, excluir ou incerto. Célula vazia = não revisado (fica fora das contas).",
        "Decida só pelo que está na linha; não procure saber o que o filtro fez com o registro.",
        "decisao_h1 é obrigatória para contar; decisao_h2 (segundo codificador) e decisao_consenso são opcionais. "
        "Com consenso, ele vale; sem consenso, discordância entre incluir/incerto e excluir fica fora das contas.",
        "Não altere as colunas ordem e id_rs nem a aba _meta. Depois: rs.py filtrar --calcular-elusao <esta planilha>",
    ]:
        inst.append([texto])
    inst.column_dimensions["A"].width = 120
    ws_meta = wb.create_sheet("_meta")
    for k, v in meta.items():
        ws_meta.append([k, v])
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


def gravar_amostra_elusao(raiz, plano, unicos):
    """Grava desenho (se novo) e planilha (se ausente; nunca sobrescreve uma planilha em codificação)."""
    raiz = Path(raiz)
    desenho = plano["desenho"]
    escritos = []
    if _gravar_atomico(raiz / plano["rel_desenho"], _json(desenho)):
        escritos.append(plano["rel_desenho"])
    caminho_planilha = raiz / plano["rel_planilha"]
    if not caminho_planilha.exists():
        por_id = {u["id_rs"]: u for u in unicos}
        linhas = []
        for ordem, id_rs in enumerate(desenho["ids"], start=1):
            u = por_id.get(id_rs, {})
            linhas.append({"ordem": ordem, "id_rs": id_rs, **{c: u.get(c, "") for c in COLUNAS_REGISTRO_ELUSAO[2:]}})
        escrever_planilha_elusao(caminho_planilha, linhas, {"tipo": TIPO_METRICAS_ELUSAO, "versao": desenho["versao"],
                                                            "desenho": plano["rel_desenho"]})
        escritos.append(plano["rel_planilha"])
    return escritos


def ler_planilha_elusao(caminho):
    """(meta, linhas com chaves em minúsculas, números das linhas) de uma planilha xlsx (aba codificacao) ou CSV.

    Leitor único de planilhas.py: CSV com `,` `;` ou tab, UTF-8, UTF-16 ou cp1252; xlsx com a aba `_meta`.
    """
    caminho = Path(caminho)
    _, linhas, info = _ler_humana(caminho, ["id_rs"], aba="codificacao")
    return planilhas.ler_meta_xlsx(caminho), linhas, info["numeros"]


def _referencia(codigos):
    """Decisão de referência: consenso; senão codificadores concordes (incluir x incerto -> incerto); senão None."""
    if codigos.get("decisao_consenso") is not None:
        return codigos["decisao_consenso"]
    valores = [codigos[c] for c in ("decisao_h1", "decisao_h2") if codigos.get(c) is not None]
    if not valores:
        return None
    if len({v in ("incluir", "incerto") for v in valores}) > 1:
        return None
    return valores[0] if len(set(valores)) == 1 else "incerto"


def calcular_elusao(raiz, caminho_planilha, caminho_desenho=None):
    """Métricas de elusão a partir da planilha codificada. Devolve (metricas, rel_desenho)."""
    from .validacao import clopper_pearson, normalizar_codigo

    raiz = Path(raiz)
    caminho_planilha = Path(caminho_planilha)
    if not caminho_planilha.is_file():
        raise ErroFiltro(f"planilha não encontrada: {caminho_planilha}")
    meta, linhas, numeros = ler_planilha_elusao(caminho_planilha)
    ref_desenho = caminho_desenho or meta.get("desenho")
    if not ref_desenho:
        irmao = caminho_planilha.with_name(re.sub(r"_cega$", "", caminho_planilha.stem) + "_desenho.json")
        ref_desenho = str(irmao)
    p_desenho = _resolver(ref_desenho, raiz, caminho_planilha.parent, Path.cwd())
    if not p_desenho.is_file():
        raise ErroFiltro(f"desenho da amostra não encontrado ({ref_desenho}); informe --desenho")
    desenho = json.loads(p_desenho.read_text(encoding="utf-8"))
    if desenho.get("tipo") != TIPO_DESENHO_ELUSAO:
        raise ErroFiltro(f"{ref_desenho} não é desenho de amostra de elusão do filtrar")
    ids_desenho = list(desenho.get("ids") or [])
    conjunto = set(ids_desenho)
    codigos, erros = {}, []
    for n, l in zip(numeros, linhas):
        id_rs = str(l.get("id_rs") or "").strip().upper()
        if not id_rs:
            continue
        if id_rs not in conjunto:
            erros.append(f"linha {n}: {id_rs} não pertence à amostra")
            continue
        if id_rs in codigos:
            erros.append(f"linha {n}: {id_rs} repetido")
            continue
        valores = {}
        for coluna in COLUNAS_CODIGO_ELUSAO:
            try:
                valores[coluna] = normalizar_codigo(l.get(coluna))
            except ValueError as e:
                erros.append(f"linha {n} ({id_rs}), {coluna}: {e}")
        codigos[id_rs] = valores
    if erros:
        raise ErroFiltro(f"planilha com {len(erros)} problema(s): " + "; ".join(erros[:10]))
    referencia = {i: _referencia(v) for i, v in codigos.items()}
    revisados = {i: r for i, r in referencia.items() if r is not None}
    discordancias = sum(1 for i, v in codigos.items() if referencia[i] is None
                        and any(v.get(c) is not None for c in COLUNAS_CODIGO_ELUSAO))
    n_rev = len(revisados)
    if not n_rev:
        raise ErroFiltro("nenhum registro codificado na planilha (decisao_h1/decisao_consenso vazias)")
    k_incluir = sum(1 for r in revisados.values() if r == "incluir")
    k_incerto = sum(1 for r in revisados.values() if r == "incerto")
    k_pos = k_incluir + k_incerto
    populacao = int(desenho.get("populacao_n") or 0)
    lo, hi = clopper_pearson(k_pos, n_rev)
    lo_e, hi_e = clopper_pearson(k_incluir, n_rev)
    taxa = k_pos / n_rev
    metricas = {
        "tipo": TIPO_METRICAS_ELUSAO, "finalidade": "elusao", "versao": desenho.get("versao"),
        "desenho": _relativo(p_desenho, raiz) or str(p_desenho),
        "planilha": _relativo(caminho_planilha, raiz) or caminho_planilha.name,
        "planilha_sha256": estado.sha256_arquivo(caminho_planilha),
        "filtros": desenho.get("filtros"), "populacao_n": populacao, "semente": desenho.get("semente"),
        "n_amostra": len(ids_desenho), "n_revisados": n_rev, "n_nao_revisados": len(ids_desenho) - n_rev,
        "n_discordancias_sem_consenso": discordancias,
        "n_incluir": k_incluir, "n_incerto": k_incerto, "n_excluir": n_rev - k_pos,
        "regra": "incluir ou incerto = relevante perdido (conservadora)",
        "taxa_elusao": round(taxa, 4), "ic95": [round(lo, 4), round(hi, 4)],
        "taxa_elusao_estrita": round(k_incluir / n_rev, 4), "ic95_estrita": [round(lo_e, 4), round(hi_e, 4)],
        "perdidos_estimados": round(taxa * populacao, 1),
        "perdidos_ic95": [round(lo * populacao, 1), round(hi * populacao, 1)],
        "censo": len(ids_desenho) == populacao and n_rev == populacao,
        "ids_relevantes": sorted(i for i, r in revisados.items() if r in ("incluir", "incerto")),
    }
    return metricas, (_relativo(p_desenho, raiz) or str(p_desenho))


def habilitar_validacao_na_config(raiz, desenho_ref, rel_metricas, filtros_desenho):
    """Preenche validacao_elusao dos filtros amostrados na configuração (dentro do projeto, não congelada)."""
    raiz = Path(raiz)
    p = Path(desenho_ref)
    p = p if p.is_absolute() else raiz / p
    rel = _relativo(p, raiz)
    snippet = f"\"validacao_elusao\": \"{rel_metricas}\""
    if rel is None:
        return {"atualizada": False, "motivo": f"configuração fora do projeto ({p}); acrescente à mão {snippet}"}
    if not p.is_file():
        return {"atualizada": False, "motivo": f"configuração {rel} não encontrada; acrescente à mão {snippet}"}
    info = (estado.carregar_estado(raiz).get("artefatos") or {}).get(rel) or {}
    if isinstance(info, dict) and info.get("congelado_em"):
        return {"atualizada": False, "config": rel,
                "motivo": f"{rel} está congelada ({info['congelado_em']}); registre emenda e acrescente {snippet}"}
    try:
        config = json.loads(p.read_text(encoding="utf-8"))
        atuais = {f["nome"]: f.get("assinatura") for f in validar_config(config, p.parent, raiz, checar=False)
                  if f["tipo"] == "dicionario"}
    except (json.JSONDecodeError, ErroFiltro) as e:
        return {"atualizada": False, "config": rel, "motivo": f"configuração inválida ({e}); acrescente à mão {snippet}"}
    alvo = {x["nome"]: x.get("assinatura") for x in filtros_desenho or []}
    mudados, alterados = [], []
    for bruto in config["filtros"]:
        nome = str(bruto.get("nome") or bruto.get("tipo"))
        if nome not in alvo:
            continue
        if atuais.get(nome) != alvo[nome]:
            alterados.append(nome)
            continue
        if bruto.get("validacao_elusao") != rel_metricas:
            bruto["validacao_elusao"] = rel_metricas
            mudados.append(nome)
    if mudados:
        _gravar_atomico(p, _json(config))
    saida = {"atualizada": bool(mudados), "config": rel, "filtros": mudados}
    if alterados:
        saida["definicao_alterada"] = alterados
        saida["motivo"] = f"filtros {alterados} mudaram desde a amostra; validação não aplicada a eles"
    return saida


def executar_calculo_elusao(args):
    try:
        raiz = estado.exigir_projeto(args.dir)
        if args.config or args.ancoras or args.amostra_elusao is not None or args.semente is not None:
            raise ErroFiltro("--calcular-elusao não se combina com --config, --ancoras, --amostra-elusao ou --semente")
        caminho = _resolver(args.calcular_elusao, Path.cwd(), raiz)
        metricas, rel_desenho = calcular_elusao(raiz, caminho, args.desenho)
    except estado.ErroProjeto as e:
        print(f"erro: {e}", file=sys.stderr)
        return 1
    except planilhas.ErroDependencia as e:
        print(f"erro: {e}", file=sys.stderr)
        estado.resumo({"comando": "filtrar", "ok": False, "erro": str(e), "dependencia_ausente": True})
        return 3
    except ErroFiltro as e:
        print(f"erro: {e}", file=sys.stderr)
        estado.resumo({"comando": "filtrar", "ok": False, "erro": str(e)})
        return 1
    desenho = json.loads((_resolver(rel_desenho, raiz)).read_text(encoding="utf-8"))
    rel_metricas = f"{DIR_VALIDACAO}/elusao_{_nome_versao(metricas['versao'])}{esquema.SUFIXO_METRICAS}"
    _gravar_atomico(raiz / rel_metricas, _json(metricas))
    config = habilitar_validacao_na_config(raiz, desenho.get("config") or "", rel_metricas, desenho.get("filtros"))
    artefatos = [a for a in (rel_desenho, rel_metricas, metricas["planilha"], config.get("config")
                             if config.get("atualizada") else None) if a and (raiz / a).exists()]
    dados = {k: v for k, v in metricas.items() if k not in ("ids_relevantes", "filtros")}
    dados.update(etapa_validada="filtro_formal", filtros=[f.get("nome") for f in metricas["filtros"] or []],
                 metricas=rel_metricas, config_habilitada=config)
    estado.registrar_evento(raiz, "validacao_calculada", ETAPA, "script", "filtrar", dados=dados, artefatos=artefatos)
    print(f"elusão ({metricas['versao']}): {metricas['n_incluir'] + metricas['n_incerto']} relevantes em "
          f"{metricas['n_revisados']} revisados = {metricas['taxa_elusao']:.3f} (IC 95% {metricas['ic95'][0]:.3f}-"
          f"{metricas['ic95'][1]:.3f}); ~{metricas['perdidos_estimados']} perdidos em {metricas['populacao_n']}")
    if config.get("motivo"):
        print(f"aviso: {config['motivo']}", file=sys.stderr)
    estado.resumo({
        "comando": "filtrar", "ok": True, "calcular_elusao": True, "versao": metricas["versao"],
        "n_revisados": metricas["n_revisados"], "n_nao_revisados": metricas["n_nao_revisados"],
        "taxa_elusao": metricas["taxa_elusao"], "ic95": metricas["ic95"],
        "taxa_elusao_estrita": metricas["taxa_elusao_estrita"], "perdidos_estimados": metricas["perdidos_estimados"],
        "perdidos_ic95": metricas["perdidos_ic95"], "populacao_n": metricas["populacao_n"],
        "config_habilitada": config, "arquivos": [rel_metricas],
    })
    return 0


# ---------------------------------------------------------------------------
# Comando
# ---------------------------------------------------------------------------
def carregar_unicos(raiz):
    """Todas as linhas de registros_unicos.csv (inclusive `busca_inativa`; quem chama separa o conjunto ativo)."""
    caminho = Path(raiz) / esquema.ARQ_UNICOS
    if not caminho.exists():
        raise ErroFiltro(f"{esquema.ARQ_UNICOS} não existe; rode `rs.py dedup` antes")
    cabecalho, linhas = _ler_csv(caminho)
    if "id_rs" not in cabecalho:
        raise ErroFiltro(f"{esquema.ARQ_UNICOS}: sem coluna id_rs")
    return linhas


def executar(args):
    if getattr(args, "calcular_elusao", None):
        return executar_calculo_elusao(args)
    amostra_n = getattr(args, "amostra_elusao", None)
    semente = getattr(args, "semente", None)
    try:
        raiz = estado.exigir_projeto(args.dir)
        if not args.config:
            raise ErroFiltro("--config é obrigatório (só --calcular-elusao dispensa)")
        if amostra_n is None and semente is not None:
            raise ErroFiltro("--semente só vale com --amostra-elusao")
        if getattr(args, "desenho", None):
            raise ErroFiltro("--desenho só vale com --calcular-elusao")
        caminho_config = _resolver(args.config, Path.cwd(), raiz)
        if not caminho_config.exists():
            raise ErroFiltro(f"configuração não encontrada: {args.config}")
        try:
            config = json.loads(caminho_config.read_text(encoding="utf-8"))
        except json.JSONDecodeError as e:
            raise ErroFiltro(f"{caminho_config}: JSON inválido ({e})")
        filtros = validar_config(config, caminho_config.parent, raiz)
        todos = carregar_unicos(raiz)
        unicos = [u for u in todos if not _buscas.cluster_inativo(u)]
        n_inativos = len(todos) - len(unicos)
        caminho_ancoras = None
        if args.ancoras:
            caminho_ancoras = _resolver(args.ancoras, Path.cwd(), raiz)
            if not caminho_ancoras.exists():
                raise ErroFiltro(f"arquivo de âncoras não encontrado: {args.ancoras}")
        linhas, contagens, excluidos, etiquetados = aplicar_filtros(unicos, filtros)
        ancoras = checar_ancoras(caminho_ancoras, unicos, excluidos, etiquetados) if caminho_ancoras else None
        recall = recall_ancoras(raiz, caminho_ancoras, ancoras, unicos) if caminho_ancoras else None
        versao = str(config.get("versao") or caminho_config.stem)
        plano = None
        if amostra_n is not None:
            plano = planejar_amostra_elusao(raiz, filtros, linhas, amostra_n, semente, versao, caminho_config)
    except estado.ErroProjeto as e:
        print(f"erro: {e}", file=sys.stderr)
        return 1
    except ErroFiltro as e:
        print(f"erro: {e}", file=sys.stderr)
        return 1
    except planilhas.ErroDependencia as e:
        print(f"erro: {e}", file=sys.stderr)
        estado.resumo({"comando": "filtrar", "ok": False, "erro": str(e), "dependencia_ausente": True})
        return 3
    except ErroMetodologico as e:
        print(f"checagem metodológica falhou: {e}", file=sys.stderr)
        estado.resumo({"comando": "filtrar", "ok": False, "erro": str(e)})
        return 2

    avisos = []
    sha_config = estado.sha256_arquivo(caminho_config)
    dicionarios = {}
    validacoes = {}
    for f in filtros:
        if f["tipo"] == "dicionario":
            rel = _relativo(f["dicionario"], raiz) or f"assets/dicionarios/{f['dicionario'].name}"
            dicionarios[rel] = f["dicionario_sha256"]
            if f["modo"] == "excluir":
                if f.get("elusao_verificada"):
                    validacoes[f["nome"]] = f["elusao_verificada"]
                else:
                    avisos.append(f"filtro '{f['nome']}': validacao_elusao '{f['validacao_elusao']}' não é métrica "
                                  "gerada por --calcular-elusao; não foi verificada")
    n_saida = contagens[-1]["saida"] if contagens else len(unicos)
    ancoras_excluidas = [a for a in (ancoras or []) if a["situacao"] == "excluida"]
    resumo_contagens = {
        "versao": versao, "config_sha256": sha_config, "dicionarios": dicionarios,
        "n_entrada": len(unicos), "n_saida": n_saida, "n_excluidos": len(excluidos),
        "n_etiquetados": len(etiquetados), "filtros": contagens,
    }
    if n_inativos:
        resumo_contagens["n_inativos_ignorados"] = n_inativos
    if ancoras is not None:
        resumo_contagens["ancoras"] = ancoras

    mudou_csv = _gravar_atomico(raiz / esquema.ARQ_FILTRO_FORMAL,
                                _texto_csv(esquema.COLUNAS_FILTRO_FORMAL, linhas))
    mudou_json = _gravar_atomico(raiz / ARQ_CONTAGENS, _json(resumo_contagens))
    mudou_recall = False
    if recall is not None:
        mudou_recall = _gravar_atomico(raiz / ARQ_RECALL_ANCORAS, _json(recall))
        avisos += recall["avisos"]
    escritos_amostra = []
    if plano is not None:
        try:
            escritos_amostra = gravar_amostra_elusao(raiz, plano, unicos)
        except ErroFiltro as e:
            print(f"erro: {e}", file=sys.stderr)
            return 1

    est = estado.carregar_estado(raiz)
    est.setdefault("versoes_ativas", {})["filtros"] = versao
    artefatos = [esquema.ARQ_UNICOS, esquema.ARQ_FILTRO_FORMAL, ARQ_CONTAGENS]
    for extra in (caminho_config, caminho_ancoras):
        rel = _relativo(extra, raiz) if extra else None
        if rel:
            artefatos.append(rel)
    if recall is not None:
        artefatos.append(ARQ_RECALL_ANCORAS)
    if plano is not None:
        artefatos += [plano["rel_desenho"], plano["rel_planilha"]]
    dados = {
        "config": _relativo(caminho_config, raiz) or caminho_config.name, **resumo_contagens,
        "modos": {f["nome"]: f["modo"] for f in filtros},
        "checagem_ancoras": None if ancoras is None else ("falhou" if ancoras_excluidas else "ok"),
        "n_inativos_ignorados": n_inativos,
        "sem_mudancas": not (mudou_csv or mudou_json or mudou_recall or escritos_amostra),
    }
    dados.pop("ancoras", None)
    if ancoras is not None:
        dados["ancoras_excluidas"] = ancoras_excluidas
        dados["ancoras_nao_encontradas"] = [a["ancora"] for a in ancoras if a["situacao"] == "nao_encontrada"]
    if recall is not None:
        dados["recall_ancoras"] = {
            "arquivo": ARQ_RECALL_ANCORAS, "combinado": recall["combinado"]["recall"],
            "combinado_todos_metodos": recall["combinado_todos_metodos"]["recall"],
            "por_busca": {b: r["recall"] for b, r in recall["por_busca"].items()},
            "nao_indexadas": recall["nao_indexadas"],
        }
    if validacoes:
        dados["validacoes_elusao"] = validacoes
    if plano is not None:
        d = plano["desenho"]
        dados["amostra_elusao"] = {"desenho": plano["rel_desenho"], "planilha": plano["rel_planilha"],
                                   "populacao_n": d["populacao_n"], "n_amostra": d["n_amostra"],
                                   "semente": d["semente"], "reutilizada": plano["reutilizada"],
                                   "filtros": [x["nome"] for x in d["filtros"]]}
    estado.registrar_evento(raiz, "filtro_formal", ETAPA, "script", "filtrar", dados=dados,
                            artefatos=artefatos, estado=est)

    print(f"filtrar ({versao}): {len(unicos)} -> {n_saida} registros; "
          f"{len(excluidos)} excluídos, {len(etiquetados)} etiquetados"
          + (f"; {n_inativos} de buscas substituídas ignorados" if n_inativos else ""))
    for c in contagens:
        print(f"  {c['filtro']:<16} [{c['modo']}] entrada {c['entrada']}: passa {c['passa']}, "
              f"etiqueta {c['etiqueta']}, exclui {c['exclui']}, sem_dado {c['sem_dado']} -> {c['saida']}")
    if recall is not None:
        comb = recall["combinado"]
        print(f"  recall das âncoras (buscas em bases): {comb['n_encontradas']}/{comb['n_ancoras']}; "
              f"por busca: " + ", ".join(f"{b} {r['n_encontradas']}/{r['n_ancoras']}"
                                         for b, r in recall["por_busca"].items()))
    if plano is not None:
        print(f"  amostra de elusão: {plano['desenho']['n_amostra']} de {plano['desenho']['populacao_n']} "
              f"({'reutilizada' if plano['reutilizada'] else 'sorteada'}) -> {plano['rel_planilha']}")
    for aviso in avisos:
        print(f"aviso: {aviso}", file=sys.stderr)
    codigo = 0
    if ancoras_excluidas:
        codigo = 2
        for a in ancoras_excluidas:
            print(f"ÂNCORA EXCLUÍDA: {a['ancora']} ({','.join(a['ids_rs'])}) pelo filtro {','.join(a['filtros'])}",
                  file=sys.stderr)
        print("checagem metodológica falhou: estudo-âncora excluído pelo funil; revise o filtro "
              "(ou use modo etiquetar) antes de seguir.", file=sys.stderr)
    resumo = {
        "comando": "filtrar", "ok": codigo == 0, "versao": versao, "n_entrada": len(unicos), "n_saida": n_saida,
        "n_excluidos": len(excluidos), "n_etiquetados": len(etiquetados), "n_inativos_ignorados": n_inativos,
        "por_filtro": {c["filtro"]: {k: c[k] for k in ("passa", "etiqueta", "exclui", "sem_dado")} for c in contagens},
        "sem_mudancas": dados["sem_mudancas"], "arquivos": [esquema.ARQ_FILTRO_FORMAL, ARQ_CONTAGENS],
        "avisos": avisos,
    }
    if ancoras is not None:
        resumo["ancoras"] = {s: sum(1 for a in ancoras if a["situacao"] == s)
                             for s in ("mantida", "etiquetada", "excluida", "nao_encontrada", "sem_chave_de_casamento")}
    if recall is not None:
        resumo["recall_ancoras"] = dados["recall_ancoras"]
        resumo["arquivos"].append(ARQ_RECALL_ANCORAS)
    if plano is not None:
        resumo["amostra_elusao"] = dados["amostra_elusao"]
        resumo["arquivos"] += [plano["rel_desenho"], plano["rel_planilha"]]
    estado.resumo(resumo)
    return codigo


def registrar(subparsers):
    p = subparsers.add_parser(
        "filtrar", help="funil formal declarado (ano, tipo, idioma, dicionário) que etiqueta por padrão",
        description="Aplica filtros sequenciais declarados em JSON a dados/registros_unicos.csv e grava "
                    "02-triagem/filtro_formal.csv. Exclusão só com previsão no protocolo; âncora excluída -> exit 2. "
                    "Também calcula o recall relativo das âncoras e sorteia/calcula a amostra de elusão do dicionário.")
    p.add_argument("--config", help="JSON de filtros (ex.: 01-busca/filtros_v1.json); obrigatório, salvo com "
                                    "--calcular-elusao")
    p.add_argument("--ancoras", help="CSV de estudos-âncora (doi e/ou titulo + ano; opcional indexada_em = B01|scopus "
                                     "ou nenhuma); grava 01-busca/recall_ancoras.json")
    p.add_argument("--amostra-elusao", type=int, metavar="N",
                   help="sorteia N registros etiquetados/excluídos pelos dicionários numa planilha cega (xlsx)")
    p.add_argument("--semente", type=int, help="semente do sorteio (obrigatória com --amostra-elusao)")
    p.add_argument("--calcular-elusao", metavar="PLANILHA",
                   help="calcula a taxa de elusão (IC 95%%) da planilha codificada e habilita validacao_elusao")
    p.add_argument("--desenho", help="desenho JSON da amostra (padrão: o indicado na aba _meta da planilha)")
    p.set_defaults(func=executar)
