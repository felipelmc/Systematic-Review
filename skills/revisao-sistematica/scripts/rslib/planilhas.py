"""Leitura tolerante de planilhas e textos preenchidos por humanos (um leitor só para toda a skill).

USO (programático)
    from rslib import planilhas
    colunas, linhas, info = planilhas.ler_tabela_humana(caminho, ["id_rs", "decisao_humana"])
    meta = planilhas.ler_meta_xlsx(caminho)            # aba _meta das planilhas cegas ({} em CSV)
    texto, codificacao, avisos = planilhas.ler_texto_humano(caminho)

O QUE ACEITA
    - xlsx/xlsm (openpyxl; aba pedida se existir, senão a primeira). `.xls` antigo é recusado com instrução.
    - Texto delimitado por vírgula, ponto e vírgula ou tabulação (escolhido pelo cabeçalho que contém as
      colunas obrigatórias; desempate por csv.Sniffer), com ou sem BOM, em UTF-8, UTF-16 (com BOM, ou sem
      BOM quando o arquivo tem bytes nulos alternados) ou cp1252 (Excel pt-BR "CSV separado por ponto e
      vírgula"). CRLF e linhas vazias são tolerados.
    - Nomes de coluna comparados sem caixa e sem BOM; as colunas pedidas voltam com o nome canônico.

QUEM USA
    `triagem override --fila`, `triagem consolidar` e `triagem fila` (fila humana), listas de IDs
    (`--ids`, `--excluir-ids`), `validar calcular` e `validar segunda-leitura` (planilhas codificadas),
    `filtrar --calcular-elusao` e `filtrar --ancoras`, e as leituras humanas de `textos`
    (via triagem_lotes.ler_tabela_humana, alias deste módulo).

POR QUE UM MÓDULO SÓ
    Cada comando tinha o próprio leitor e cada um aceitava uma parte dos formatos: a fila humana devolvida
    pelo Excel em cp1252 com ponto e vírgula passava pelo `override` e quebrava o `consolidar` logo depois.
    Com um leitor único, o que um comando aceita os outros também aceitam.

ERROS
    ErroUso (código 1): arquivo ausente, formato antigo, cabeçalho sem as colunas obrigatórias (o arquivo
    nunca é alterado). ErroDependencia (subclasse, código 3): openpyxl ausente para ler xlsx.
"""

import csv
import datetime as _dt
import io
from pathlib import Path

DELIMITADORES_HUMANOS = (",", ";", "\t")
SUFIXOS_XLSX = (".xlsx", ".xlsm")
BOM = "﻿"


class ErroUso(RuntimeError):
    """Erro de uso ou de dados que deve encerrar o comando com código 1."""


class ErroDependencia(ErroUso):
    """Pacote obrigatório ausente (ex.: openpyxl para ler xlsx): código de saída 3."""


def _ajustar_limite_csv():
    """Resumos longos estouram o limite padrão de 131 kB por campo do módulo csv."""
    import sys

    limite = sys.maxsize
    while True:
        try:
            csv.field_size_limit(limite)
            return
        except OverflowError:
            limite = int(limite / 10)


_ajustar_limite_csv()


# ---------------------------------------------------------------------------
# Texto
# ---------------------------------------------------------------------------
def _parece_utf16_sem_bom(bruto):
    amostra = bruto[:4096]
    if len(amostra) < 4:
        return None
    pares, impares = amostra[1::2], amostra[0::2]
    if pares.count(0) > len(pares) * 0.3 and impares.count(0) < len(impares) * 0.05:
        return "utf-16-le"
    if impares.count(0) > len(impares) * 0.3 and pares.count(0) < len(pares) * 0.05:
        return "utf-16-be"
    return None


def decodificar(bruto):
    """(texto, codificação): BOM UTF-16, UTF-16 sem BOM, UTF-8 com ou sem BOM, senão cp1252 (Excel pt-BR)."""
    if bruto.startswith((b"\xff\xfe", b"\xfe\xff")):
        return bruto.decode("utf-16"), "utf-16"
    sem_bom = _parece_utf16_sem_bom(bruto)
    if sem_bom:
        try:
            return bruto.decode(sem_bom), "utf-16"
        except UnicodeDecodeError:
            pass
    try:
        return bruto.decode("utf-8-sig"), "utf-8"
    except UnicodeDecodeError:
        return bruto.decode("cp1252", errors="replace"), "cp1252"


def aviso_codificacao(nome, codificacao):
    if codificacao == "cp1252":
        return f"{nome}: arquivo não está em UTF-8; lido como cp1252 (Excel). Confira os acentos"
    return None


def ler_texto_humano(caminho):
    """Texto de um arquivo escrito por humano (Markdown, CSV): (texto sem BOM, codificação, avisos)."""
    caminho = Path(caminho)
    if not caminho.is_file():
        raise ErroUso(f"arquivo não existe: {caminho}")
    texto, codificacao = decodificar(caminho.read_bytes())
    aviso = aviso_codificacao(caminho.name, codificacao)
    return texto.lstrip(BOM), codificacao, [aviso] if aviso else []


# ---------------------------------------------------------------------------
# Tabelas
# ---------------------------------------------------------------------------
def limpar_coluna(nome):
    return str(nome if nome is not None else "").replace(BOM, "").strip()


def celula_texto(valor):
    """Valor de célula xlsx como texto: 2020.0 -> "2020", data -> AAAA-MM-DD, vazio -> ""."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "1" if valor else "0"
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, _dt.datetime):
        return valor.date().isoformat() if valor.time() == _dt.time(0) else valor.isoformat()
    if isinstance(valor, _dt.date):
        return valor.isoformat()
    return str(valor)


def _atende_cabecalho(colunas, obrigatorias, alternativas):
    presentes = {c.lower() for c in colunas}
    if any(o.lower() not in presentes for o in obrigatorias):
        return False
    return not alternativas or any(all(c.lower() in presentes for c in grupo) for grupo in alternativas)


def _primeira_linha(texto, delimitador):
    for linha in csv.reader(io.StringIO(texto), delimiter=delimitador):
        if any(c.strip() for c in linha):
            return [limpar_coluna(c) for c in linha]
    return []


def _erro_cabecalho(nome, colunas, obrigatorias, alternativas, detalhe=""):
    exigidas = list(obrigatorias)
    if alternativas:
        exigidas.append(" ou ".join("+".join(g) for g in alternativas))
    lidas = ", ".join(c for c in colunas[:15] if c) or "nenhuma"
    return ErroUso(
        f"{nome}: o cabeçalho não tem as colunas obrigatórias ({', '.join(exigidas)}); colunas lidas: {lidas}{detalhe}. "
        "Mantenha o cabeçalho gerado pela skill e salve como CSV (vírgula, ponto e vírgula ou tabulação) ou xlsx; "
        "o arquivo não foi alterado")


def _renomear_canonicas(colunas, obrigatorias, alternativas):
    canonicas = {c.lower(): c for c in list(obrigatorias) + [c for g in (alternativas or []) for c in g]}
    saida, vistas = [], set()
    for c in colunas:
        nome = canonicas.get(c.lower(), c)
        saida.append("" if nome in vistas else nome)  # coluna repetida: vale a primeira
        vistas.add(nome)
    return saida


def _openpyxl():
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ErroDependencia("openpyxl ausente para ler xlsx: python3 -m pip install openpyxl") from e
    return load_workbook


def ler_tabela_humana(caminho, obrigatorias=(), alternativas=None, aba=None, rotulo=None):
    """Lê uma planilha preenchida por humano sem supor o formato de gravação. Devolve (colunas, linhas, info).

    - xlsx/xlsm (openpyxl; aba `aba` se existir, senão a primeira) ou texto delimitado (ver docstring do módulo);
    - `obrigatorias`: colunas exigidas; `alternativas`: grupos de colunas, dos quais um precisa estar completo.
      Nomes comparados sem caixa e renomeados para a forma pedida. Cabeçalho que não bate → ErroUso, sem
      alterar o arquivo.
    `info` traz formato, delimitador, codificação, `numeros` (linha de cada registro no arquivo) e avisos.
    Linhas totalmente vazias são ignoradas; células ausentes viram "".
    """
    caminho = Path(caminho)
    nome = rotulo or caminho.name
    if not caminho.is_file():
        raise ErroUso(f"arquivo não existe: {caminho}")
    sufixo = caminho.suffix.lower()
    info = {"formato": "csv", "delimitador": None, "codificacao": None, "numeros": [], "avisos": []}
    if sufixo == ".xls":
        raise ErroUso(f"{nome}: formato .xls antigo não é lido; salve como .xlsx ou CSV")
    if sufixo in SUFIXOS_XLSX:
        load_workbook = _openpyxl()
        info["formato"] = "xlsx"
        wb = load_workbook(caminho, read_only=True, data_only=True)
        try:
            ws = wb[aba] if aba and aba in wb.sheetnames else wb.worksheets[0]
            brutas = [[celula_texto(v) for v in (valores or ())] for valores in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
        registros = [(n, linha) for n, linha in enumerate(brutas, start=1)]
    else:
        texto, codificacao = decodificar(caminho.read_bytes())
        texto = texto.lstrip(BOM)
        info["codificacao"] = codificacao
        aviso = aviso_codificacao(nome, codificacao)
        if aviso:
            info["avisos"].append(aviso)
        exige = bool(obrigatorias or alternativas)
        candidatos = [d for d in DELIMITADORES_HUMANOS
                      if not exige or _atende_cabecalho(_primeira_linha(texto, d), obrigatorias, alternativas)]
        if not candidatos:
            melhor = max(DELIMITADORES_HUMANOS, key=lambda d: len(_primeira_linha(texto, d)))
            raise _erro_cabecalho(nome, _primeira_linha(texto, melhor), obrigatorias, alternativas)
        delimitador = candidatos[0]
        if len(candidatos) > 1:
            try:
                delimitador = csv.Sniffer().sniff(texto[:8192], delimiters="".join(candidatos)).delimiter
            except csv.Error:
                delimitador = max(candidatos, key=lambda d: (len(_primeira_linha(texto, d)), d == ","))
        info["delimitador"] = delimitador
        leitor = csv.reader(io.StringIO(texto), delimiter=delimitador)
        registros = []
        for linha in leitor:
            registros.append((leitor.line_num, linha))
    cabecalho, corpo = None, []
    for numero, valores in registros:
        if cabecalho is None:
            if any(str(v).strip() for v in valores):
                cabecalho = [limpar_coluna(v) for v in valores]
            continue
        if any(str(v).strip() for v in valores):
            corpo.append((numero, valores))
    cabecalho = cabecalho or []
    if not _atende_cabecalho(cabecalho, obrigatorias, alternativas):
        raise _erro_cabecalho(nome, cabecalho, obrigatorias, alternativas)
    colunas = _renomear_canonicas(cabecalho, obrigatorias, alternativas)
    linhas = []
    for numero, valores in corpo:
        linha = {}
        for i, coluna in enumerate(colunas):
            if coluna and coluna not in linha:
                linha[coluna] = valores[i] if i < len(valores) and valores[i] is not None else ""
        linhas.append(linha)
        info["numeros"].append(numero)
    return [c for c in colunas if c], linhas, info


def ler_meta_xlsx(caminho, aba="_meta"):
    """{chave: valor} da aba de metadados de uma planilha xlsx gerada pela skill; {} em CSV ou sem a aba."""
    caminho = Path(caminho)
    if caminho.suffix.lower() not in SUFIXOS_XLSX or not caminho.is_file():
        return {}
    load_workbook = _openpyxl()
    meta = {}
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        if aba in wb.sheetnames:
            for valores in wb[aba].iter_rows(values_only=True):
                if valores and valores[0] is not None:
                    meta[str(valores[0])] = valores[1] if len(valores) > 1 else None
    finally:
        wb.close()
    return meta
