"""Validação humana da triagem por IA: amostra cega, métricas com IC, elusão e estabilidade.

USO
    rs.py validar amostrar --etapa ta --rodada ta_v2 --n 100 --semente 11 \\
        [--enriquecer-incluidos 60] [--elusao 300] [--regra consenso] [--codificadores 2] \\
        [--finalidade validacao|desenvolvimento|calibracao] [--excluir-ids usados.csv]
    rs.py validar amostrar --rodada calib_v1 --finalidade calibracao --n 100 --semente 3 \\
        [--sem-ia] [--ids candidatos.csv] [--criterios 02-triagem/prompts/ta_v1.md]
    rs.py validar elusao --rodada ta_v2 --n 300 --semente 12 [--finalidade elusao]
    rs.py validar calcular --planilha 02-triagem/validacao/ta_v2/amostra01_cega.xlsx [--desenho ...] \\
        [--finalidade ...]
    rs.py validar estabilidade --rodada ta_v2 --amostra 0.1 [--semente 7] [--criterios ...] [--finalidade ...]
    rs.py validar amostrar --rodada ta_v1 --n <>= 20% da rodada> --semente 5 --atalho-rapida   (variante rápida)
    rs.py validar segunda-leitura --planilha 02-triagem/validacao/ta_v1/elusao01_cega.xlsx [--rodada ta_v1]
    rs.py validar calcular --planilha .../amostra01_cega.xlsx [--atalho-rapida]

ATALHO DA VARIANTE RÁPIDA (references/tipos-de-revisao.md, seção 6)
    Com projeto.variante = rapida e a amostra de validação marcada (`amostrar --atalho-rapida`, ou `calcular
    --atalho-rapida`), `calcular` grava em `validacao_calculada` os campos de esquema.CAMPOS_ATALHO_RAPIDA:
    atalho_rapida = true, fracao_dupla_humana = n_dupla_humana / n_populacao (registros com >= 2 códigos humanos
    sobre os registros ativos da rodada), kappa_humanos, n_excluidos_ia (decisão só IA, fora inativos e funil),
    n_excluidos_relidos (lidos por humano na amostra ou em `validar segunda-leitura`) e
    segunda_leitura_excluidos. O resumo sai com código 2 enquanto faltar dupla >= 20%, κ ou a releitura de
    todos os excluídos (recall não é exigido nesse caminho). `segunda-leitura` grava
    <DIR_VALIDACAO>/<rodada>/segunda_leitura.json (leituras acumulam entre execuções), a fila
    segunda_leitura_resgatados.csv (formato da fila humana, para `triagem override --fila`) com os excluídos que
    a leitura manda seguir, e um `validacao_calculada` com tipo segunda_leitura e finalidade elusao (não decide o
    G4). Depois dela, rode `calcular` de novo para gravar a segunda leitura no evento da validação.

FINALIDADE (esquema.FINALIDADES_VALIDACAO)
    Todo `validacao_calculada` grava `dados.finalidade`. Só a última validação com finalidade
    `validacao` da rodada ativa (versoes_ativas.rodada_ta) decide o G4 e a marca de rascunho da
    declaração de IA; calibração e desenvolvimento ficam como histórico.
    - calibracao (A): dupla humana calibra os critérios; limiares κ ≥ 0,6, concordância ≥ 75% e
      tamanho ≥ 100 registros ou ≥ 10 incluídos humanos. Métricas da IA, se houver, são só
      informativas. Com `--sem-ia` (ou quando a rodada ainda não tem decisões de IA) a amostra sai
      de registros_unicos.csv menos o funil formal, e o cálculo mede só a concordância entre humanos.
      Reprovada → pendência `calibracao_reprovada` (autopiloto).
    - desenvolvimento (B): conjunto em que o prompt é ajustado; mesmas métricas da validação, sem
      pendência (reprovar aqui é o esperado enquanto o prompt evolui).
    - validacao (C/D): amostra nova, limiares de recall; reprovada → `validacao_triagem_reprovada`.
    - elusao (F) e estabilidade (G): relatadas, sem limiar vinculante.
    Padrão: `validacao` em amostrar, `elusao` em elusao, `estabilidade` em estabilidade; `calcular`
    herda a do desenho (divergir do desenho é erro). Compatíveis: amostra → calibracao,
    desenvolvimento, validacao; elusao → elusao, desenvolvimento; estabilidade → estabilidade,
    desenvolvimento.
    `--excluir-ids` tira do quadro amostral (e da população usada nos pesos) IDs já usados em outra
    finalidade ou rodada, para que a validação seja feita em amostra nova (references/ia-validacao.md, 4 C).

PROTOCOLO (references/ia-validacao.md, seção 4, A–H; resumo no Apêndice D da base de conhecimento)
    (C) amostra nova com dupla codificação humana, enriquecida para ≥ 60 incluídos;
        métricas: sensibilidade, especificidade, precisão, VPN, κ, PABAK, WSS@95, IC
        Clopper-Pearson.
    (D) limiar: recall ≥ 0,95 com limite inferior do IC ≥ 0,90; concordância humana
        κ ≥ 0,6 e ≥ 75% (A). Resultado em `atende_limiares`; falha → exit 2 e lista de
        falsos negativos para a análise que antecede os critérios vN+1 (E).
    (F) elusão: amostra dos excluídos pela IA (n ≥ 300 ou todos) → incluídos perdidos.
    (G) estabilidade: reexecução de 5–10% dos registros com o mesmo prompt.

REPROVAÇÃO SÓ PELA LARGURA DO IC
    Com 0 falsos negativos e todos os outros critérios atendidos, a validação reprova só porque há poucos
    incluídos humanos (o limite inferior de Clopper-Pearson com x = n só chega a 0,90 com n ≥ 36). Nesse
    caso `motivo_reprovacao = largura_ic` e a próxima ação é ampliar a amostra enriquecida (ou a elusão),
    não revisar critérios; se a rodada não tem incluídos pela IA suficientes, o limiar é inalcançável e o
    caminho é o remédio 5 (dupla humana) com aprovação humana do G4 (`--forcar --motivo`).

DESENHOS ANTIGOS E BUSCAS SUBSTITUÍDAS
    Reexecutar `amostrar`, `elusao` ou `estabilidade` com os mesmos parâmetros reaproveita o desenho; se ele
    tem IDs que hoje são `busca_inativa`, o comando avisa (`n_inativos_no_desenho`) e orienta sortear amostra
    nova com outra `--semente`. `calcular` também avisa.

POR QUE ASSIM
    - A planilha é cega: não tem decisão, justificativa, trecho, estrato nem peso. O
      desenho (estratos, pesos, origem de cada ID) fica num JSON separado que não vai
      para os codificadores; a ordem das linhas é embaralhada.
    - O enriquecimento sobre-representa o que a IA incluiu (proxy dos incluídos humanos,
      desconhecidos antes da codificação). Por isso as métricas populacionais são
      reponderadas por pós-estratificação: peso = N do estrato / nº codificado no estrato.
      Sem reponderar, o recall ficaria otimista. Os IC ponderados usam Clopper-Pearson
      com tamanho efetivo de Kish (aproximação de Korn-Graubard); sem pesos, são exatos.
    - "Não revisado" (célula vazia ou nao_revisado) é NA e sai de todas as contas, em vez
      de virar exclusão, erro que enviesou o κ nos scripts de exemplo do REFIS.
    - A decisão avaliada é a "só IA" (consolidação sem overrides humanos): medir depois
      que humanos corrigiram a IA inflaria o desempenho.
    - κ e PABAK são calculados sobre a decisão binária que importa para o fluxo (seguir
      ao texto completo = incluir ou incerto; excluir) e, entre humanos, também sobre as
      três categorias. Estatística em Python puro (sem numpy/scipy), testada contra
      scikit-learn e scipy quando instalados.
"""

import json
import math
import random
import sys
from collections import Counter
from pathlib import Path

from . import esquema, estado, normalizar, planilhas
from . import triagem_lotes as tl

ESTRATO_POS = "ia_positivo"
ESTRATO_NEG = "ia_negativo"
ESTRATO_SEM_IA = "sem_ia"
LIMIARES = {
    "kappa_humanos_min": 0.60,
    "concordancia_humanos_min": 0.75,
    "recall_min": 0.95,
    "recall_ic_inferior_min": 0.90,
    "incluidos_humanos_alvo": 60,
    "elusao_n_alvo": 300,
    "calibracao_n_min": 100,
    "calibracao_incluidos_min": 10,
}
FINALIDADE_PADRAO = {"amostra": "validacao", "elusao": "elusao", "estabilidade": "estabilidade"}
FINALIDADES_COMPATIVEIS = {
    "amostra": {"calibracao", "desenvolvimento", "validacao"},
    "elusao": {"elusao", "desenvolvimento"},
    "estabilidade": {"estabilidade", "desenvolvimento"},
}
PENDENCIA_REPROVADA = {"validacao": "validacao_triagem_reprovada", "calibracao": "calibracao_reprovada"}


def conferir_finalidade(tipo, finalidade):
    """Finalidade efetiva (padrão do tipo quando None); ErroUso se não combina com o tipo de desenho."""
    finalidade = finalidade or FINALIDADE_PADRAO[tipo]
    if finalidade not in esquema.FINALIDADES_VALIDACAO:
        raise tl.ErroUso(f"finalidade desconhecida: {finalidade} (use {', '.join(esquema.FINALIDADES_VALIDACAO)})")
    if finalidade not in FINALIDADES_COMPATIVEIS[tipo]:
        raise tl.ErroUso(f"finalidade {finalidade} não combina com desenho do tipo {tipo} "
                         f"(aceitas: {', '.join(sorted(FINALIDADES_COMPATIVEIS[tipo]))})")
    return finalidade
COLUNAS_REGISTRO_PLANILHA = ["ordem", "id_rs", "titulo", "resumo", "palavras_chave", "ano", "veiculo",
                             "idioma", "tipo_publicacao"]
VALORES_CODIGO = ["incluir", "excluir", "incerto", "nao_revisado"]
_CODIGOS = {
    "incluir": "incluir", "inclui": "incluir", "incluido": "incluir", "include": "incluir", "sim": "incluir",
    "excluir": "excluir", "exclui": "excluir", "excluido": "excluir", "exclude": "excluir",
    "incerto": "incerto", "talvez": "incerto", "duvida": "incerto", "uncertain": "incerto", "?": "incerto",
}
_NA = {"", "na", "n/a", "nr", "nao revisado", "nao_revisado", "nao-revisado", "-"}


# ---------------------------------------------------------------------------
# Estatística (Python puro)
# ---------------------------------------------------------------------------
def _betacf(a, b, x, max_iter=10000, eps=3e-16):
    """Fração continuada da beta incompleta (algoritmo de Lentz modificado)."""
    minimo = 1e-300
    qab, qap, qam = a + b, a + 1.0, a - 1.0
    c, d = 1.0, 1.0 - qab * x / qap
    d = 1.0 / (d if abs(d) > minimo else minimo)
    h = d
    for m in range(1, max_iter + 1):
        m2 = 2 * m
        aa = m * (b - m) * x / ((qam + m2) * (a + m2))
        d = 1.0 + aa * d
        d = 1.0 / (d if abs(d) > minimo else minimo)
        c = 1.0 + aa / c
        c = c if abs(c) > minimo else minimo
        h *= d * c
        aa = -(a + m) * (qab + m) * x / ((a + m2) * (qap + m2))
        d = 1.0 + aa * d
        d = 1.0 / (d if abs(d) > minimo else minimo)
        c = 1.0 + aa / c
        c = c if abs(c) > minimo else minimo
        delta = d * c
        h *= delta
        if abs(delta - 1.0) < eps:
            break
    return h


def beta_incompleta(a, b, x):
    """Função beta incompleta regularizada I_x(a, b)."""
    if x <= 0.0:
        return 0.0
    if x >= 1.0:
        return 1.0
    log_bt = math.lgamma(a + b) - math.lgamma(a) - math.lgamma(b) + a * math.log(x) + b * math.log1p(-x)
    bt = math.exp(log_bt)
    if x < (a + 1.0) / (a + b + 2.0):
        return bt * _betacf(a, b, x) / a
    return 1.0 - bt * _betacf(b, a, 1.0 - x) / b


def beta_quantil(q, a, b):
    """Quantil da distribuição Beta(a, b) por bisseção (monótona, robusta, precisão ~1e-12)."""
    lo, hi = 0.0, 1.0
    for _ in range(200):
        meio = (lo + hi) / 2.0
        if beta_incompleta(a, b, meio) < q:
            lo = meio
        else:
            hi = meio
        if hi - lo < 1e-13:
            break
    return (lo + hi) / 2.0


def clopper_pearson(x, n, alfa=0.05):
    """IC bilateral exato (Clopper-Pearson) para x sucessos em n; aceita x e n não inteiros."""
    if n is None or n <= 0:
        return None, None
    x = min(max(x, 0.0), n)
    inferior = 0.0 if x <= 0 else beta_quantil(alfa / 2.0, x, n - x + 1.0)
    superior = 1.0 if x >= n else beta_quantil(1.0 - alfa / 2.0, x + 1.0, n - x)
    return inferior, superior


def _pares_validos(a, b):
    return [(x, y) for x, y in zip(a, b) if x is not None and y is not None]


def concordancia(a, b):
    pares = _pares_validos(a, b)
    if not pares:
        return None
    return sum(x == y for x, y in pares) / len(pares)


def kappa_cohen(a, b, categorias=None):
    """κ de Cohen com NA (None) excluído par a par; None quando indefinido (p_e = 1)."""
    pares = _pares_validos(a, b)
    n = len(pares)
    if n == 0:
        return None
    cats = sorted(set(categorias or []) | {v for par in pares for v in par})
    p_o = sum(x == y for x, y in pares) / n
    cont_a = Counter(x for x, _ in pares)
    cont_b = Counter(y for _, y in pares)
    p_e = sum((cont_a[c] / n) * (cont_b[c] / n) for c in cats)
    if abs(1.0 - p_e) < 1e-15:
        return None
    return (p_o - p_e) / (1.0 - p_e)


def pabak(a, b, k=2):
    """Kappa ajustado por prevalência e viés: (k·p_o − 1)/(k − 1); binário = 2·p_o − 1."""
    p_o = concordancia(a, b)
    if p_o is None:
        return None
    return (k * p_o - 1.0) / (k - 1.0)


def kappa_fleiss(itens, categorias=None):
    """κ de Fleiss. `itens` = lista de listas de rótulos (None = NA).

    Usa só itens com o número máximo de avaliações não-NA (Fleiss exige m constante).
    """
    limpos = [[v for v in item if v is not None] for item in itens]
    if not limpos:
        return None
    m = max(len(i) for i in limpos)
    limpos = [i for i in limpos if len(i) == m]
    if m < 2 or not limpos:
        return None
    cats = sorted(set(categorias or []) | {v for item in limpos for v in item})
    n_itens = len(limpos)
    totais = Counter()
    soma_p = 0.0
    for item in limpos:
        cont = Counter(item)
        totais.update(cont)
        soma_p += (sum(c * c for c in cont.values()) - m) / (m * (m - 1))
    p_barra = soma_p / n_itens
    p_e = sum((totais[c] / (n_itens * m)) ** 2 for c in cats)
    if abs(1.0 - p_e) < 1e-15:
        return None
    return (p_barra - p_e) / (1.0 - p_e)


def proporcao(pares, alfa=0.05):
    """Proporção ponderada com IC. `pares` = [(peso, sucesso_bool)].

    Com pesos iguais o IC é o Clopper-Pearson exato; com pesos desiguais usa o tamanho
    efetivo de Kish, n_ef = (Σw)²/Σw², e x_ef = p·n_ef.
    """
    if not pares:
        return {"valor": None, "ic_inferior": None, "ic_superior": None, "n": 0, "x": 0, "n_efetivo": 0}
    soma_w = sum(w for w, _ in pares)
    soma_w2 = sum(w * w for w, _ in pares)
    p = sum(w for w, s in pares if s) / soma_w
    n_ef = soma_w * soma_w / soma_w2
    x_ef = p * n_ef
    if all(abs(w - pares[0][0]) < 1e-12 for w, _ in pares):
        n_ef, x_ef = float(len(pares)), float(sum(1 for _, s in pares if s))
    inf, sup = clopper_pearson(x_ef, n_ef, alfa)
    return {"valor": p, "ic_inferior": inf, "ic_superior": sup, "n": len(pares),
            "x": sum(1 for _, s in pares if s), "n_efetivo": round(n_ef, 3)}


def metricas_classificador(itens, alfa=0.05):
    """Métricas de um classificador binário contra a referência humana.

    `itens` = [(pred_positivo, ref_positivo, peso, id_rs)]. Contagens brutas + proporções
    ponderadas (os pesos são 1 quando não há reponderação).
    """
    brutos = Counter()
    ponderados = Counter()
    for pred, ref, w, _ in itens:
        chave = ("tp" if pred else "fn") if ref else ("fp" if pred else "tn")
        brutos[chave] += 1
        ponderados[chave] += w
    sens = proporcao([(w, pred) for pred, ref, w, _ in itens if ref], alfa)
    esp = proporcao([(w, not pred) for pred, ref, w, _ in itens if not ref], alfa)
    prec = proporcao([(w, ref) for pred, ref, w, _ in itens if pred], alfa)
    vpn = proporcao([(w, not ref) for pred, ref, w, _ in itens if not pred], alfa)
    conc = proporcao([(w, pred == ref) for pred, ref, w, _ in itens], alfa)
    preds = [pred for pred, _, _, _ in itens]
    refs = [ref for _, ref, _, _ in itens]
    total_w = sum(ponderados.values())
    wss = wss95 = None
    if total_w > 0 and sens["valor"] is not None:
        fracao_nao_lida = (ponderados["tn"] + ponderados["fn"]) / total_w
        wss = fracao_nao_lida - (1.0 - sens["valor"])
        wss95 = fracao_nao_lida - 0.05 if sens["valor"] >= 0.95 else None
    return {
        "n": len(itens),
        "matriz": {k: brutos.get(k, 0) for k in ("tp", "fp", "fn", "tn")},
        "matriz_ponderada": {k: round(ponderados.get(k, 0.0), 3) for k in ("tp", "fp", "fn", "tn")},
        "sensibilidade": sens, "especificidade": esp, "precisao": prec, "vpn": vpn,
        "concordancia": conc,
        "kappa": kappa_cohen(preds, refs, [True, False]),
        "pabak": pabak(preds, refs, 2),
        "wss": wss, "wss_95": wss95,
        "falsos_negativos": sorted(i for pred, ref, _, i in itens if ref and not pred),
    }


# ---------------------------------------------------------------------------
# Códigos humanos e planilha
# ---------------------------------------------------------------------------
def normalizar_codigo(valor):
    """Código humano → incluir|excluir|incerto; None para não revisado. ValueError se ilegível."""
    if valor is None:
        return None
    s = normalizar.ascii_fold(str(valor)).strip().lower()
    if s in _NA:
        return None
    if s in _CODIGOS:
        return _CODIGOS[s]
    raise ValueError(f"código não reconhecido: {valor!r} (use {', '.join(VALORES_CODIGO)})")


def pasta_validacao(raiz, rodada):
    return Path(raiz) / esquema.DIR_VALIDACAO / rodada


def _desenhos_existentes(raiz, rodada):
    pasta = pasta_validacao(raiz, rodada)
    saida = []
    if pasta.exists():
        for arq in sorted(pasta.glob("*_desenho.json")):
            with open(arq, encoding="utf-8") as f:
                desenho = json.load(f)
            if desenho.get("tipo") in ("amostra", "elusao"):
                saida.append((arq, desenho))
    return saida


def _decisoes_ia(raiz, etapa, rodada, regra, unicos):
    vig = tl.decisoes_vigentes(tl.ler_decisoes(raiz), etapa, [rodada])
    if not vig:
        raise tl.ErroUso(f"nenhuma decisão da etapa {etapa} na rodada {rodada}")
    return tl.consolidar_decisoes(vig, regra, usar_overrides=False, sem_resumo=tl.ids_sem_resumo(unicos))


def _criterios_da_rodada(raiz, rodada):
    for manifesto in tl.manifestos_da_rodada(raiz, rodada):
        return manifesto.get("criterios_arquivo"), manifesto.get("criterios_sha")
    return None, None


def _escrever_planilha(caminho, linhas, codificadores, meta, instrucoes):
    """Planilha cega em xlsx: só dados bibliográficos + colunas vazias para os humanos."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as e:  # dependência obrigatória, mas falhar com instrução clara
        raise tl.ErroUso("openpyxl ausente: pip install openpyxl") from e
    colunas = COLUNAS_REGISTRO_PLANILHA[:]
    for k in range(1, codificadores + 1):
        colunas += [f"decisao_h{k}", f"criterio_h{k}"]
    colunas += ["decisao_consenso", "observacoes"]
    wb = Workbook()
    ws = wb.active
    ws.title = "codificacao"
    ws.append(colunas)
    for c in ws[1]:
        c.font = Font(bold=True)
    for linha in linhas:
        ws.append([linha.get(c) for c in colunas])  # colunas de codificação ficam realmente vazias
    validacao = DataValidation(type="list", formula1='"' + ",".join(VALORES_CODIGO) + '"', allow_blank=True)
    ws.add_data_validation(validacao)
    ultima = max(2, len(linhas) + 1)
    for i, nome in enumerate(colunas, start=1):
        letra = ws.cell(row=1, column=i).column_letter
        if nome.startswith("decisao_"):
            validacao.add(f"{letra}2:{letra}{ultima}")
        largura = {"titulo": 50, "resumo": 90, "palavras_chave": 30, "observacoes": 30}.get(nome, 14)
        ws.column_dimensions[letra].width = largura
    for linha_celulas in ws.iter_rows(min_row=2):
        for c in linha_celulas:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws_inst = wb.create_sheet("instrucoes")
    for texto in instrucoes:
        ws_inst.append([texto])
    ws_inst.column_dimensions["A"].width = 120
    ws_meta = wb.create_sheet("_meta")
    for chave, valor in meta.items():
        ws_meta.append([chave, valor])
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


def ler_planilha(caminho):
    """Lê planilha codificada (xlsx ou csv). Devolve (meta, linhas com códigos normalizados, erros).

    A leitura passa pelo leitor único de planilhas.py: CSV com `,`, `;` ou tabulação, BOM, UTF-8, UTF-16 ou
    cp1252, ou xlsx (aba `codificacao`; metadados da aba `_meta`). Sem coluna id_rs ou sem nenhuma coluna decisao_* o comando para com erro, em
    vez de tratar todas as linhas como não revisadas.
    """
    caminho = Path(caminho)
    meta = {}
    _, brutas, _ = tl.ler_tabela_humana(caminho, ["id_rs"], aba="codificacao")
    brutas = [{k.strip().lower(): v for k, v in l.items() if k} for l in brutas]
    colunas = {k for l in brutas for k in l}
    if brutas and not any(c.startswith("decisao_") for c in colunas):
        raise tl.ErroUso(f"{caminho.name}: sem colunas decisao_hN nem decisao_consenso; mantenha o cabeçalho da planilha "
                         "cega gerada por `validar amostrar`")
    meta.update(planilhas.ler_meta_xlsx(caminho))
    erros, saida = [], []
    for n, bruta in enumerate(brutas, start=2):
        id_rs = str(bruta.get("id_rs") or "").strip().upper()
        if not id_rs:
            continue
        codigos, consenso = {}, None
        for coluna, valor in bruta.items():
            if not coluna.startswith("decisao_"):
                continue
            try:
                codigo = normalizar_codigo(valor)
            except ValueError as e:
                erros.append(f"{caminho.name} linha {n} ({id_rs}), {coluna}: {e}")
                continue
            if coluna == "decisao_consenso":
                consenso = codigo
            else:
                codigos[coluna[len("decisao_"):]] = codigo
        saida.append({"id_rs": id_rs, "codigos": codigos, "consenso": consenso})
    return meta, saida, erros


def referencia_humana(codigos, consenso):
    """Decisão humana de referência: consenso; senão codificadores concordes; senão NA."""
    if consenso is not None:
        return consenso, "consenso"
    valores = [v for v in codigos.values() if v is not None]
    if not valores:
        return None, "nao_revisado"
    if len({tl.positiva(v) for v in valores}) > 1:
        return None, "discordancia_sem_consenso"
    rotulo = valores[0] if len(set(valores)) == 1 else "incerto"
    return rotulo, ("unico" if len(valores) == 1 else "concordancia")


# ---------------------------------------------------------------------------
# Amostragem
# ---------------------------------------------------------------------------
def _sha_ids(ids):
    return estado.sha256_texto("\n".join(sorted(ids))) if ids else None


def ids_do_desenho(desenho):
    """id_rs sorteados num desenho de amostra/elusão (`amostra`) ou de estabilidade (`ids`)."""
    return [a["id_rs"] for a in desenho.get("amostra") or []] or list(desenho.get("ids") or [])


def ids_inativos_no_desenho(raiz, desenho, unicos=None):
    """IDs do desenho que hoje são clusters `busca_inativa` (desenho sorteado antes de uma substituição de busca)."""
    unicos = unicos if unicos is not None else tl.ler_unicos(raiz, obrigatorio=False)
    return sorted(set(ids_do_desenho(desenho)) & tl.ids_inativos(unicos))


def avisos_desenho_com_inativos(raiz, desenho, unicos=None):
    inativos = ids_inativos_no_desenho(raiz, desenho, unicos)
    if not inativos:
        return []
    nome = desenho.get("amostra_id") or "de estabilidade"
    if desenho.get("amostra") is None and desenho.get("ids") is not None:
        orientacao = ("a concordância ainda mede a estabilidade do modelo, mas relate; para sortear só registros ativos, "
                      f"renomeie {esquema.ARQ_DESENHO_ESTABILIDADE} (arquivo antigo) e rode com outra --semente")
    else:
        orientacao = ("os resultados medem registros fora do fluxo: sorteie amostra nova com outra --semente (os IDs "
                      "inativos ficam fora do quadro sozinhos) antes de usar esta amostra no G4")
    return [f"o desenho {nome} foi sorteado antes de uma substituição de busca: {len(inativos)} IDs são hoje "
            f"{esquema.FLAG_BUSCA_INATIVA} ({inativos[:10]}); {orientacao}"]


# ---------------------------------------------------------------------------
# Atalho da variante rápida (esquema.CAMPOS_ATALHO_RAPIDA)
# ---------------------------------------------------------------------------
def variante_do_projeto(raiz):
    try:
        return ((estado.carregar_estado(raiz).get("projeto") or {}).get("variante")) or None
    except (estado.ErroProjeto, OSError, ValueError):
        return None


def avisos_atalho_amostra(raiz, n_amostra, n_populacao):
    avisos = []
    if variante_do_projeto(raiz) != "rapida":
        avisos.append("--atalho-rapida num projeto que não é variante rápida (projeto.variante): `validar calcular` "
                      "não marca o atalho e o G4 segue a regra geral (references/tipos-de-revisao.md, seção 6)")
    minimo = math.ceil(esquema.FRACAO_MINIMA_DUPLA_RAPIDA * n_populacao)
    if n_populacao and n_amostra < minimo:
        avisos.append(f"atalho da variante rápida: a amostra tem {n_amostra} registros e a dupla humana precisa cobrir "
                      f">= {esquema.FRACAO_MINIMA_DUPLA_RAPIDA:.0%} dos {n_populacao} da rodada (>= {minimo}); "
                      "amplie --n para não reprovar no G4")
    return avisos


def ids_excluidos_pela_ia(raiz, etapa, rodada, regra="consenso", unicos=None, res=None):
    """id_rs que a IA excluiu na rodada (decisão só IA, sem overrides), fora clusters inativos e o funil formal."""
    unicos = unicos if unicos is not None else tl.ler_unicos(raiz, obrigatorio=False)
    res = res if res is not None else _decisoes_ia(raiz, etapa, rodada, regra, unicos)
    fora = tl.ids_inativos(unicos) | tl.ids_excluidos_funil(raiz)
    return sorted(i for i, r in res.items() if r["decisao_final"] == "excluir" and i not in fora)


def caminho_segunda_leitura(raiz, rodada):
    return pasta_validacao(raiz, rodada) / esquema.ARQ_SEGUNDA_LEITURA


def ler_segunda_leitura(raiz, rodada):
    caminho = caminho_segunda_leitura(raiz, rodada)
    if not caminho.is_file():
        return None
    try:
        with open(caminho, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return None


def bloco_atalho_rapida(raiz, desenho, linhas_amostra, par_principal, res, unicos, regra):
    """Campos do atalho da variante rápida para validacao_calculada (esquema.CAMPOS_ATALHO_RAPIDA) e o que falta."""
    populacao = desenho.get("populacao_rodada") or desenho.get("populacao") or {}
    n_populacao = int(sum(populacao.values()))
    n_dupla = sum(1 for l in linhas_amostra if sum(1 for v in l["codigos"].values() if v is not None) >= 2)
    fracao = (n_dupla / n_populacao) if n_populacao else None
    kappa = (par_principal or {}).get("kappa_binario")
    param = desenho["parametros"]
    excluidos = set(ids_excluidos_pela_ia(raiz, param["etapa"], param["rodada"], regra, unicos, res))
    relidos_amostra = {l["id_rs"] for l in linhas_amostra if any(v is not None for v in l["codigos"].values())}
    segunda = ler_segunda_leitura(raiz, param["rodada"]) or {}
    relidos_segunda = set(segunda.get("ids_relidos") or [])
    if segunda and segunda.get("regra", regra) != regra:
        relidos_segunda = set()
    relidos = excluidos & (relidos_amostra | relidos_segunda)
    faltam = sorted(excluidos - relidos)
    bloco = {
        "atalho_rapida": True,
        "fracao_dupla_humana": None if fracao is None else round(fracao, 4),
        "n_dupla_humana": n_dupla, "n_populacao": n_populacao,
        "kappa_humanos": kappa,
        "segunda_leitura_excluidos": not faltam,
        "n_excluidos_ia": len(excluidos), "n_excluidos_relidos": len(relidos),
    }
    criterios = {
        "fracao_dupla_humana": {"valor": bloco["fracao_dupla_humana"], "limiar": esquema.FRACAO_MINIMA_DUPLA_RAPIDA,
                                "atende": fracao is not None and fracao >= esquema.FRACAO_MINIMA_DUPLA_RAPIDA},
        "kappa_humanos_calculado": {"valor": kappa, "limiar": "calculado", "atende": kappa is not None},
        "segunda_leitura_excluidos": {"valor": f"{len(relidos)}/{len(excluidos)}", "limiar": "todos",
                                      "atende": not faltam},
    }
    extras = {"criterios": criterios, "atende": all(c["atende"] for c in criterios.values()),
              "ids_faltam_reler": faltam[:50], "n_faltam_reler": len(faltam),
              "n_relidos_na_amostra": len(excluidos & relidos_amostra),
              "ids_relidos_na_amostra": sorted(excluidos & relidos_amostra),
              "segunda_leitura": (tl.relativo(raiz, caminho_segunda_leitura(raiz, param["rodada"]))
                                  if segunda else None),
              "segunda_leitura_sha256": (estado.sha256_arquivo(caminho_segunda_leitura(raiz, param["rodada"]))
                                         if segunda else None)}
    return bloco, extras


def gerar_amostra(raiz, etapa, rodada, n, semente, enriquecer=None, elusao=0, regra="consenso",
                  codificadores=2, tipo="amostra", finalidade=None, excluir_ids=None, ids=None,
                  sem_ia=False, criterios=None, atalho_rapida=False):
    """Sorteia a amostra cega e grava planilha + desenho. Idempotente para os mesmos parâmetros.

    `excluir_ids` sai do quadro amostral e da população dos pesos; `ids` restringe o quadro. Clusters com a
    flag `busca_inativa` (buscas substituídas) nunca entram no quadro nem na população (`n_inativos_ignorados`).
    `sem_ia` (só calibração) sorteia de registros_unicos.csv sem estratificar por decisão de IA; a
    calibração também cai nesse modo quando a rodada ainda não tem decisões de IA.
    """
    raiz = Path(raiz)
    tl._validar_nome(rodada, "rodada")
    finalidade = conferir_finalidade(tipo, finalidade)
    if n < 0 or (elusao or 0) < 0 or (enriquecer or 0) < 0:
        raise tl.ErroUso("tamanhos de amostra não podem ser negativos")
    if n == 0 and not elusao:
        raise tl.ErroUso("amostra vazia: informe --n e/ou --elusao")
    if codificadores < 1:
        raise tl.ErroUso("--codificadores deve ser ≥ 1")
    if sem_ia and finalidade != "calibracao":
        raise tl.ErroUso("--sem-ia só vale para --finalidade calibracao (as demais medem a IA)")
    if atalho_rapida and (tipo != "amostra" or finalidade != "validacao" or sem_ia):
        raise tl.ErroUso("--atalho-rapida marca a amostra de validação (finalidade validacao, com decisões de IA) "
                         "da variante rápida")
    if atalho_rapida and codificadores < 2:
        raise tl.ErroUso("--atalho-rapida exige dupla codificação humana (--codificadores 2 ou mais)")
    excluir = set(excluir_ids or [])
    restritos = set(ids) if ids is not None else None
    parametros = {"etapa": etapa, "rodada": rodada, "n": n, "semente": semente, "enriquecer_incluidos": enriquecer,
                  "elusao": elusao, "regra": regra, "codificadores": codificadores, "tipo": tipo}
    # Chaves novas só quando diferem do padrão: desenhos antigos continuam idempotentes.
    if finalidade != FINALIDADE_PADRAO[tipo]:
        parametros["finalidade"] = finalidade
    if excluir:
        parametros["excluir_ids_sha"] = _sha_ids(excluir)
    if restritos is not None:
        parametros["ids_sha"] = _sha_ids(restritos)
    if sem_ia:
        parametros["sem_ia"] = True
    if atalho_rapida:
        parametros["atalho_rapida"] = True
    existentes = _desenhos_existentes(raiz, rodada)
    for arq, desenho in existentes:
        if desenho.get("parametros") == parametros:
            return arq, desenho, True, avisos_desenho_com_inativos(raiz, desenho)
    unicos = tl.ler_unicos(raiz, obrigatorio=False)
    avisos = []
    # Clusters busca_inativa (só registros de buscas substituídas) saem do quadro e da população dos pesos,
    # como no PRISMA e na triagem: validar sobre eles mediria registros que não estão no fluxo.
    inativos = tl.ids_inativos(unicos)
    vigentes = tl.decisoes_vigentes(tl.ler_decisoes(raiz), etapa, [rodada])
    modo_sem_ia = sem_ia or (finalidade == "calibracao" and not vigentes)
    if modo_sem_ia:
        if tipo != "amostra" or enriquecer or elusao:
            raise tl.ErroUso("amostra sem decisões de IA não tem enriquecimento nem elusão (estratos da IA)")
        if restritos is None and not unicos:
            raise tl.ErroUso(f"{esquema.ARQ_UNICOS} ausente: rode `dedup` ou informe --ids")
        if restritos is not None and unicos:
            desconhecidos = sorted(restritos - set(unicos))
            if desconhecidos:
                raise tl.ErroUso(f"IDs ausentes de {esquema.ARQ_UNICOS}: {desconhecidos[:5]}")
        quadro_rodada = set(restritos) if restritos is not None else set(unicos) - tl.ids_excluidos_funil(raiz)
        n_inativos_ignorados = len(quadro_rodada & inativos)
        quadro_rodada -= inativos
        if not sem_ia:
            avisos.append(f"rodada {rodada} sem decisões de IA: calibração só entre humanos")
        positivos, negativos = [], []
        quadro = sorted(quadro_rodada - excluir)
        estrato = {i: ESTRATO_SEM_IA for i in quadro}
        populacao = {ESTRATO_SEM_IA: len(quadro)}
        populacao_rodada = {ESTRATO_SEM_IA: len(quadro_rodada)}
    else:
        res = _decisoes_ia(raiz, etapa, rodada, regra, unicos)
        quadro_rodada = set(res) if restritos is None else set(res) & restritos
        n_inativos_ignorados = len(quadro_rodada & inativos)
        quadro_rodada -= inativos
        if restritos is not None and len(quadro_rodada) < len(restritos):
            avisos.append(f"{len(restritos) - len(quadro_rodada)} IDs de --ids sem decisão da IA na rodada ficaram de fora")
        positivos_rodada = sorted(i for i in quadro_rodada if tl.positiva(res[i]["decisao_final"]))
        negativos_rodada = sorted(i for i in quadro_rodada if not tl.positiva(res[i]["decisao_final"]))
        positivos = [i for i in positivos_rodada if i not in excluir]
        negativos = [i for i in negativos_rodada if i not in excluir]
        quadro = sorted(set(positivos) | set(negativos))
        estrato = {i: ESTRATO_POS for i in positivos}
        estrato.update({i: ESTRATO_NEG for i in negativos})
        populacao = {ESTRATO_POS: len(positivos), ESTRATO_NEG: len(negativos)}
        populacao_rodada = {ESTRATO_POS: len(positivos_rodada), ESTRATO_NEG: len(negativos_rodada)}
    if n_inativos_ignorados:
        avisos.append(f"{n_inativos_ignorados} registros de buscas substituídas (flag {esquema.FLAG_BUSCA_INATIVA}) "
                      "saíram do quadro amostral e da população dos pesos")
    n_excluidos_quadro = len(excluir & quadro_rodada)
    if excluir:
        avisos.append(f"{n_excluidos_quadro} IDs de --excluir-ids saíram do quadro amostral e da população dos pesos")
    usados = {item["id_rs"] for _, d in existentes for item in d.get("amostra", [])}
    if usados:
        avisos.append(f"{len(usados)} IDs de amostras anteriores desta rodada ficaram de fora (amostra nova)")
    rng = random.Random(f"{semente}:{rodada}:{tipo}")
    disponiveis = sorted(set(quadro) - usados)
    principal = rng.sample(disponiveis, min(n, len(disponiveis)))
    if n > len(disponiveis):
        avisos.append(f"pedidos {n}, disponíveis {len(disponiveis)}: todos entraram")
    origem = {i: "aleatoria" for i in principal}
    escolhidos = set(principal)
    if enriquecer:
        faltam = enriquecer - sum(1 for i in principal if estrato[i] == ESTRATO_POS)
        resto = sorted(set(positivos) - usados - escolhidos)
        extra = rng.sample(resto, min(max(faltam, 0), len(resto)))
        origem.update({i: "enriquecimento" for i in extra})
        escolhidos |= set(extra)
        total_pos = sum(1 for i in escolhidos if estrato[i] == ESTRATO_POS)
        if total_pos < enriquecer:
            avisos.append(f"só {total_pos} incluídos/incertos pela IA disponíveis para enriquecer (alvo {enriquecer}); "
                          "relate o IC resultante")
    if elusao:
        faltam = elusao - sum(1 for i in escolhidos if estrato[i] == ESTRATO_NEG)
        resto = sorted(set(negativos) - usados - escolhidos)
        extra = rng.sample(resto, min(max(faltam, 0), len(resto)))
        origem.update({i: "elusao" for i in extra})
        escolhidos |= set(extra)
        total_neg = sum(1 for i in escolhidos if estrato[i] == ESTRATO_NEG)
        if total_neg < elusao:
            avisos.append(f"só {total_neg} excluídos pela IA disponíveis para a elusão (alvo {elusao}): todos entraram")
    if not escolhidos:
        raise tl.ErroUso("nenhum registro disponível para amostrar nesta rodada")
    if atalho_rapida:
        avisos.extend(avisos_atalho_amostra(raiz, len(escolhidos), sum(populacao_rodada.values())))

    numero = 1 + sum(1 for _, d in existentes if d.get("tipo") == tipo)
    amostra_id = f"{tipo}{numero:02d}"
    pasta = pasta_validacao(raiz, rodada)
    caminho_planilha = pasta / f"{amostra_id}{esquema.SUFIXO_PLANILHA_CEGA}"
    caminho_desenho = pasta / f"{amostra_id}{esquema.SUFIXO_DESENHO}"
    ordem = sorted(escolhidos)
    random.Random(f"{semente}:{rodada}:{tipo}:planilha").shuffle(ordem)
    linhas = []
    for k, id_rs in enumerate(ordem, start=1):
        reg = unicos.get(id_rs, {})
        linha = {"ordem": k, "id_rs": id_rs}
        for c in COLUNAS_REGISTRO_PLANILHA[2:]:
            linha[c] = normalizar.texto(reg.get(c))
        linhas.append(linha)
    criterios_arquivo, criterios_sha = _criterios_da_rodada(raiz, rodada)
    if criterios:
        caminho_criterios = tl.resolver(raiz, criterios)
        if not caminho_criterios.exists():
            raise tl.ErroUso(f"arquivo de critérios não existe: {criterios}")
        sha_informado = estado.sha256_arquivo(caminho_criterios)
        if criterios_sha and sha_informado != criterios_sha:
            avisos.append("--criterios difere dos critérios congelados nos lotes da rodada; vale o dos lotes")
        else:
            criterios_arquivo, criterios_sha = tl.relativo(raiz, caminho_criterios), sha_informado
    meta = {"amostra_id": amostra_id, "rodada": rodada, "etapa": etapa, "finalidade": finalidade,
            "desenho": tl.relativo(raiz, caminho_desenho), "criterios_sha": criterios_sha or ""}
    rotulo = {"calibracao": "calibração", "desenvolvimento": "desenvolvimento", "validacao": "validação",
              "elusao": "elusão"}.get(finalidade, finalidade)
    instrucoes = [
        f"Amostra {amostra_id} da rodada {rodada} (etapa {etapa}, finalidade: {rotulo}). "
        "Planilha cega: não há decisões da IA.",
        f"Critérios: {criterios_arquivo or 'arquivo de critérios da rodada'} (sha256 {criterios_sha or 'ver log'}).",
        "Cada codificador preenche só a própria coluna decisao_hN (e criterio_hN nas exclusões), sem ver a do outro "
        "(oculte as demais colunas enquanto codifica).",
        "Valores: incluir, excluir, incerto. Célula vazia ou nao_revisado = não revisado (fica fora das métricas).",
        "Depois da codificação independente, preencha decisao_consenso nas discordâncias.",
        "Registro sem resumo: codifique pelo título; na dúvida, incerto.",
        f"Depois: rs.py validar calcular --planilha {tl.relativo(raiz, caminho_planilha)}",
    ]
    _escrever_planilha(caminho_planilha, linhas, codificadores, meta, instrucoes)
    contagens = dict(Counter(origem[i] for i in escolhidos))
    por_estrato = Counter(estrato[i] for i in escolhidos)
    desenho = {
        "versao": 1, "amostra_id": amostra_id, "tipo": tipo, "finalidade": finalidade, "sem_ia": modo_sem_ia,
        "atalho_rapida": bool(atalho_rapida),
        "parametros": parametros, "criado_em": estado.agora(),
        "planilha": tl.relativo(raiz, caminho_planilha), "criterios_arquivo": criterios_arquivo,
        "criterios_sha": criterios_sha,
        "populacao": populacao, "populacao_rodada": populacao_rodada, "n_excluidos_quadro": n_excluidos_quadro,
        "n_inativos_ignorados": n_inativos_ignorados,
        "n_amostra": len(escolhidos), "contagens_origem": contagens,
        "amostra_por_estrato": dict(por_estrato),
        "pesos_nominais": {s: (populacao[s] / por_estrato[s]) if por_estrato.get(s) else None for s in populacao},
        "amostra": [{"id_rs": i, "estrato": estrato[i], "origem": origem[i]} for i in sorted(escolhidos)],
        "avisos": avisos,
    }
    tl.escrever_json(caminho_desenho, desenho)
    return caminho_desenho, desenho, False, avisos


def finalidade_do_desenho(desenho):
    return desenho.get("finalidade") or (desenho.get("parametros") or {}).get("finalidade") \
        or FINALIDADE_PADRAO.get(desenho.get("tipo"), "validacao")


def _registrar_amostra(raiz, arq_desenho, desenho, reexecucao, avisos, comando):
    rel_planilha = desenho["planilha"]
    etapa = desenho["parametros"]["etapa"]
    etapa_projeto = tl.ETAPA_PROJETO[etapa]
    finalidade = finalidade_do_desenho(desenho)
    if not reexecucao:
        estado.registrar_evento(
            raiz, "artefato_versionado", etapa_projeto, "script", "validacao",
            dados={"tipo": f"desenho_{desenho['tipo']}", "amostra_id": desenho["amostra_id"],
                   "rodada": desenho["parametros"]["rodada"], "finalidade": finalidade,
                   "sem_ia": bool(desenho.get("sem_ia")), "parametros": desenho["parametros"],
                   "populacao": desenho["populacao"], "n_amostra": desenho["n_amostra"],
                   "n_excluidos_quadro": desenho.get("n_excluidos_quadro", 0),
                   "n_inativos_ignorados": desenho.get("n_inativos_ignorados", 0),
                   "contagens_origem": desenho["contagens_origem"], "pesos_nominais": desenho["pesos_nominais"]},
            artefatos=[tl.relativo(raiz, arq_desenho), rel_planilha])
    pendencia = tl.sincronizar_pendencia(
        raiz, "validacao_humana", etapa_projeto, tl.PORTAO_ETAPA[etapa],
        f"codificar em dupla, às cegas, a amostra {desenho['amostra_id']} ({desenho['n_amostra']} registros)",
        desenho["n_amostra"], rel_planilha, "validacao") if not reexecucao else None
    for aviso in avisos:
        print(f"aviso: {aviso}", file=sys.stderr)
    avisos_resumo = list(desenho.get("avisos", [])) + [a for a in avisos if a not in desenho.get("avisos", [])]
    estado.resumo({
        "comando": comando, "ok": True, "reexecucao": reexecucao, "amostra_id": desenho["amostra_id"],
        "rodada": desenho["parametros"]["rodada"], "finalidade": finalidade, "sem_ia": bool(desenho.get("sem_ia")),
        "n_amostra": desenho["n_amostra"], "n_inativos_ignorados": desenho.get("n_inativos_ignorados", 0),
        "n_inativos_no_desenho": len(ids_inativos_no_desenho(raiz, desenho)),
        "populacao": desenho["populacao"], "amostra_por_estrato": desenho["amostra_por_estrato"],
        "contagens_origem": desenho["contagens_origem"], "planilha": rel_planilha,
        "desenho": tl.relativo(raiz, arq_desenho), "pendencia": pendencia, "avisos": avisos_resumo,
        "proxima_acao": f"codificação humana dupla e cega; depois `validar calcular --planilha {rel_planilha}`",
    })
    return 0


def _ids_opcionais(raiz, caminho):
    return tl.ler_ids_arquivo(tl.resolver(raiz, caminho)) if caminho else None


def cmd_amostrar(args):
    raiz = estado.exigir_projeto(args.dir)
    arq, desenho, reexec, avisos = gerar_amostra(
        raiz, args.etapa, args.rodada, args.n, args.semente, args.enriquecer_incluidos, args.elusao,
        args.regra, args.codificadores, tipo="amostra", finalidade=args.finalidade,
        excluir_ids=_ids_opcionais(raiz, args.excluir_ids), ids=_ids_opcionais(raiz, args.ids),
        sem_ia=args.sem_ia, criterios=args.criterios, atalho_rapida=getattr(args, "atalho_rapida", False))
    return _registrar_amostra(raiz, arq, desenho, reexec, avisos, "validar amostrar")


def cmd_elusao(args):
    raiz = estado.exigir_projeto(args.dir)
    arq, desenho, reexec, avisos = gerar_amostra(
        raiz, args.etapa, args.rodada, 0, args.semente, None, args.n, args.regra, args.codificadores, tipo="elusao",
        finalidade=args.finalidade, excluir_ids=_ids_opcionais(raiz, args.excluir_ids))
    return _registrar_amostra(raiz, arq, desenho, reexec, avisos, "validar elusao")


# ---------------------------------------------------------------------------
# Cálculo
# ---------------------------------------------------------------------------
def _itens_ponderados(ids, pred, ref, estrato, populacao):
    """Itens (pred, ref, peso, id) com pós-estratificação; sem unidade codificada num estrato → pesos 1."""
    usados = [i for i in ids if pred.get(i) is not None and ref.get(i) is not None and i in estrato]
    n_estrato = Counter(estrato[i] for i in usados)
    ponderavel = bool(usados) and all(n_estrato.get(s, 0) > 0 for s, total in populacao.items() if total > 0)
    itens = []
    for i in usados:
        peso = populacao[estrato[i]] / n_estrato[estrato[i]] if ponderavel else 1.0
        itens.append((tl.positiva(pred[i]), tl.positiva(ref[i]), peso, i))
    return itens, ponderavel


def _bloco_humanos(linhas):
    codificadores = sorted({c for l in linhas for c in l["codigos"]})
    bloco = {"codificadores": codificadores, "pares": {}, "fleiss_binario": None, "fleiss_3cat": None}
    for x in range(len(codificadores)):
        for y in range(x + 1, len(codificadores)):
            a = [l["codigos"].get(codificadores[x]) for l in linhas]
            b = [l["codigos"].get(codificadores[y]) for l in linhas]
            a_bin = [None if v is None else tl.positiva(v) for v in a]
            b_bin = [None if v is None else tl.positiva(v) for v in b]
            bloco["pares"][f"{codificadores[x]}~{codificadores[y]}"] = {
                "n": len(_pares_validos(a, b)),
                "concordancia_binaria": concordancia(a_bin, b_bin),
                "kappa_binario": kappa_cohen(a_bin, b_bin, [True, False]),
                "pabak_binario": pabak(a_bin, b_bin, 2),
                "concordancia_3cat": concordancia(a, b),
                "kappa_3cat": kappa_cohen(a, b, ["incluir", "excluir", "incerto"]),
                "pabak_3cat": pabak(a, b, 3),
            }
    if len(codificadores) >= 2:
        itens = [[l["codigos"].get(c) for c in codificadores] for l in linhas]
        bloco["fleiss_3cat"] = kappa_fleiss(itens, ["incluir", "excluir", "incerto"])
        bloco["fleiss_binario"] = kappa_fleiss(
            [[None if v is None else tl.positiva(v) for v in item] for item in itens], [True, False])
    return bloco


def calcular_metricas(raiz, planilhas, caminho_desenho=None, regra=None, alfa=0.05, finalidade=None,
                      atalho_rapida=False):
    """Calcula todas as métricas de uma amostra codificada. Devolve (metricas, linhas_fn).

    `finalidade` None herda a do desenho; informar outra diferente da registrada no desenho é erro
    (uma amostra de calibração não vira a validação que decide o G4).
    Atalho da variante rápida: com o desenho marcado (`amostrar --atalho-rapida`) ou `atalho_rapida=True`,
    finalidade validacao e projeto.variante = rapida, `metricas["atalho_rapida"]` traz os campos de
    esquema.CAMPOS_ATALHO_RAPIDA (dupla humana sobre a população da rodada, κ, segunda leitura dos excluídos
    pela IA, somando a amostra e `validar segunda-leitura`) e o que falta.
    """
    raiz = Path(raiz)
    metas, linhas, erros, fontes = [], {}, [], []
    for p in planilhas:
        caminho = tl.resolver(raiz, p)
        if not caminho.exists():
            raise tl.ErroUso(f"planilha não existe: {p}")
        meta, lidas, errs = ler_planilha(caminho)
        metas.append(meta)
        erros.extend(errs)
        fontes.append({"caminho": tl.relativo(raiz, caminho), "sha256": estado.sha256_arquivo(caminho)})
        for l in lidas:
            atual = linhas.setdefault(l["id_rs"], {"id_rs": l["id_rs"], "codigos": {}, "consenso": None})
            for c, v in l["codigos"].items():
                if v is not None and atual["codigos"].get(c) not in (None, v):
                    erros.append(f"{l['id_rs']}: codificador {c} com valores conflitantes entre planilhas")
                if v is not None or c not in atual["codigos"]:
                    atual["codigos"][c] = v
            if l["consenso"] is not None:
                atual["consenso"] = l["consenso"]
    if erros:
        raise tl.ErroUso("planilha com códigos inválidos:\n  " + "\n  ".join(erros[:30]))
    if caminho_desenho is None:
        caminho_desenho = next((m.get("desenho") for m in metas if m.get("desenho")), None)
    if not caminho_desenho:
        raise tl.ErroUso("não achei o desenho da amostra (aba _meta); informe --desenho")
    caminho_desenho = tl.resolver(raiz, caminho_desenho)
    with open(caminho_desenho, encoding="utf-8") as f:
        desenho = json.load(f)
    param = desenho["parametros"]
    etapa, rodada = param["etapa"], param["rodada"]
    regra = regra or param.get("regra", "consenso")
    registrada = desenho.get("finalidade") or param.get("finalidade")
    if finalidade and registrada and finalidade != registrada:
        raise tl.ErroUso(f"o desenho {desenho['amostra_id']} foi sorteado com finalidade {registrada}, não {finalidade}; "
                         "sorteie amostra nova para outra finalidade")
    finalidade = conferir_finalidade(desenho["tipo"], finalidade or registrada)
    sem_ia = bool(desenho.get("sem_ia"))
    avisos = []
    estrato = {a["id_rs"]: a["estrato"] for a in desenho["amostra"]}
    populacao = desenho["populacao"]
    fora = sorted(set(linhas) - set(estrato))
    if fora:
        avisos.append(f"{len(fora)} IDs da planilha não pertencem ao desenho e foram ignorados")
    ids = sorted(estrato)

    unicos = tl.ler_unicos(raiz, obrigatorio=False)
    avisos.extend(avisos_desenho_com_inativos(raiz, desenho, unicos))
    res = {} if sem_ia else _decisoes_ia(raiz, etapa, rodada, regra, unicos)
    mudaram = [i for i in ids if i in res and
               (ESTRATO_POS if tl.positiva(res[i]["decisao_final"]) else ESTRATO_NEG) != estrato[i]]
    if mudaram:
        avisos.append(f"{len(mudaram)} decisões da IA mudaram depois da amostragem (estratos mantidos pelo desenho)")

    ref, origem_ref = {}, Counter()
    for i in ids:
        l = linhas.get(i, {"codigos": {}, "consenso": None})
        ref[i], origem = referencia_humana(l["codigos"], l["consenso"])
        origem_ref[origem] += 1
    if origem_ref.get("discordancia_sem_consenso"):
        avisos.append(f"{origem_ref['discordancia_sem_consenso']} discordâncias humanas sem decisao_consenso ficaram fora")
    codificados = [i for i in ids if ref[i] is not None]
    n_incl_humanos = sum(1 for i in codificados if tl.positiva(ref[i]))
    linhas_amostra = [linhas.get(i, {"id_rs": i, "codigos": {}, "consenso": None}) for i in ids]
    humanos = _bloco_humanos(linhas_amostra)

    metricas = {
        "amostra_id": desenho["amostra_id"], "tipo": desenho["tipo"], "finalidade": finalidade, "sem_ia": sem_ia,
        "rodada": rodada, "etapa": etapa, "regra": regra, "planilhas": fontes, "desenho": tl.relativo(raiz, caminho_desenho),
        "n_amostra": len(ids), "n_codificados": len(codificados), "n_na": len(ids) - len(codificados),
        "origem_referencia": dict(origem_ref), "n_incluidos_humanos": n_incl_humanos,
        "populacao": populacao, "humanos": humanos, "limiares": LIMIARES,
    }

    # Elusão: proporção de incluídos humanos entre os excluídos pela IA (estrato do desenho).
    neg = [i for i in codificados if estrato[i] == ESTRATO_NEG]
    elusao = proporcao([(1.0, tl.positiva(ref[i])) for i in neg], alfa)
    n_neg_pop = populacao.get(ESTRATO_NEG, 0)
    metricas["elusao"] = {
        "n_codificados": len(neg), "taxa": elusao,
        "perdidos_estimados": None if elusao["valor"] is None else elusao["valor"] * n_neg_pop,
        "perdidos_ic": None if elusao["valor"] is None else [elusao["ic_inferior"] * n_neg_pop,
                                                             elusao["ic_superior"] * n_neg_pop],
        "excluidos_ia_populacao": n_neg_pop,
    }
    if desenho["tipo"] == "elusao":
        if len(neg) < LIMIARES["elusao_n_alvo"] and len(neg) < n_neg_pop:
            avisos.append(f"elusão com {len(neg)} codificados (< {LIMIARES['elusao_n_alvo']} e < todos os excluídos)")
        metricas.update({"ia_final": None, "por_revisor": {}, "criterios": {}, "atende_limiares": None,
                         "falsos_negativos": sorted(i for i in neg if tl.positiva(ref[i])), "avisos": avisos,
                         "nota": "elusão não tem limiar vinculante (references/ia-validacao.md, seção 4 F); relatar "
                                 "taxa e perdidos estimados"})
        linhas_fn = _linhas_falsos_negativos(metricas["falsos_negativos"], res, ref, estrato, unicos)
        return metricas, linhas_fn

    par_principal = next(iter(humanos["pares"].values()), None)
    criterios_humanos = {
        "kappa_humanos": _criterio(par_principal and par_principal["kappa_binario"], LIMIARES["kappa_humanos_min"]),
        "concordancia_humanos": _criterio(par_principal and par_principal["concordancia_binaria"],
                                          LIMIARES["concordancia_humanos_min"]),
    }
    if par_principal is None:
        avisos.append("sem dupla codificação humana: a validação exige dois codificadores (references/ia-validacao.md, "
                      "seção 4 C; critério não atendido)")
    elif par_principal["kappa_binario"] is None:
        avisos.append("κ humano indefinido (sem variação nas decisões): critério não atendido")
    if finalidade == "calibracao":
        n_dupla = par_principal["n"] if par_principal else 0
        criterios_humanos["tamanho_calibracao"] = {
            "valor": {"n_dupla_codificacao": n_dupla, "incluidos_humanos": n_incl_humanos},
            "limiar": f">= {LIMIARES['calibracao_n_min']} registros ou >= {LIMIARES['calibracao_incluidos_min']} "
                      "incluídos humanos",
            "atende": bool(n_dupla >= LIMIARES["calibracao_n_min"]
                           or n_incl_humanos >= LIMIARES["calibracao_incluidos_min"]),
        }
        if not criterios_humanos["tamanho_calibracao"]["atende"]:
            avisos.append(f"calibração pequena: {n_dupla} registros em dupla e {n_incl_humanos} incluídos humanos "
                          f"(references/ia-validacao.md, seção 4 A: >= {LIMIARES['calibracao_n_min']} ou >= "
                          f"{LIMIARES['calibracao_incluidos_min']} incluídos)")
    if sem_ia:
        metricas.update({"ponderado": False, "ia_final": None, "por_revisor": {}, "criterios": criterios_humanos,
                         "atende_limiares": all(c["atende"] is True for c in criterios_humanos.values()),
                         "falsos_negativos": [], "avisos": avisos,
                         "nota": "calibração sem decisões de IA: só concordância entre humanos "
                                 "(references/ia-validacao.md, seção 4 A)"})
        return metricas, []

    pred_final = {i: res[i]["decisao_final"] for i in ids if i in res}
    itens, ponderavel = _itens_ponderados(ids, pred_final, ref, estrato, populacao)
    if not ponderavel:
        avisos.append("sem codificação em algum estrato: métricas não reponderadas (valem só para a amostra)")
    metricas["ponderado"] = ponderavel
    metricas["ia_final"] = metricas_classificador(itens, alfa)
    revisores = sorted({r for i in ids if i in res for r in res[i]["primarias"]})
    por_revisor = {}
    for rev in revisores + ([tl.REVISOR_ARBITRO] if any(res.get(i, {}).get("arbitro") for i in ids) else []):
        if rev == tl.REVISOR_ARBITRO:
            pred = {i: res[i]["arbitro"]["decisao"] for i in ids if i in res and res[i]["arbitro"]}
        else:
            pred = {i: res[i]["primarias"][rev]["decisao"] for i in ids if i in res and rev in res[i]["primarias"]}
        itens_rev, pond_rev = _itens_ponderados(ids, pred, ref, estrato, populacao)
        if rev == tl.REVISOR_ARBITRO:
            itens_rev = [(p, r, 1.0, i) for p, r, _, i in itens_rev]  # árbitro só vê divergências: sem reponderar
            pond_rev = False
        por_revisor[rev] = {"ponderado": pond_rev, **metricas_classificador(itens_rev, alfa)}
    metricas["por_revisor"] = por_revisor

    sens = metricas["ia_final"]["sensibilidade"]
    if finalidade == "calibracao":
        # Calibração (A) mede os critérios pela dupla humana; o desempenho da IA fica só como informação.
        criterios = criterios_humanos
    else:
        criterios = {
            "recall": _criterio(sens["valor"], LIMIARES["recall_min"]),
            "recall_ic_inferior": _criterio(sens["ic_inferior"], LIMIARES["recall_ic_inferior_min"]),
            **criterios_humanos,
        }
    if finalidade != "calibracao" and n_incl_humanos < LIMIARES["incluidos_humanos_alvo"]:
        avisos.append(f"{n_incl_humanos} incluídos humanos na amostra (alvo ≥ {LIMIARES['incluidos_humanos_alvo']}); "
                      "em revisões pequenas, relate que todos os disponíveis entraram e o IC resultante")
    metricas["criterios"] = criterios
    metricas["atende_limiares"] = all(c["atende"] is True for c in criterios.values())
    metricas["falsos_negativos"] = metricas["ia_final"]["falsos_negativos"]
    pedido_atalho = bool(atalho_rapida or desenho.get("atalho_rapida") or param.get("atalho_rapida"))
    if pedido_atalho:
        if finalidade != "validacao":
            avisos.append(f"atalho da variante rápida ignorado: a amostra tem finalidade {finalidade} (só a validação "
                          "decide o G4)")
        elif variante_do_projeto(raiz) != "rapida":
            avisos.append("atalho da variante rápida ignorado: o projeto não é variante rápida (projeto.variante); "
                          "vale a regra geral do G4 (references/tipos-de-revisao.md, seção 6)")
        else:
            campos, extras = bloco_atalho_rapida(raiz, desenho, linhas_amostra, par_principal, res, unicos, regra)
            metricas["atalho_rapida"] = {**campos, **extras}
    metricas["avisos"] = avisos
    linhas_fn = _linhas_falsos_negativos(metricas["falsos_negativos"], res, ref, estrato, unicos)
    return metricas, linhas_fn


def _criterio(valor, limiar):
    return {"valor": valor, "limiar": limiar, "atende": None if valor is None else bool(valor >= limiar)}


COLUNAS_FALSOS_NEGATIVOS = ["id_rs", "titulo", "estrato", "decisao_ia", "decidido_por", "pareceres",
                            "referencia_humana"]


def _linhas_falsos_negativos(ids, res, ref, estrato, unicos):
    saida = []
    for i in ids:
        r = res.get(i, {})
        pareceres = [tl._descrever_parecer(p) for p in r.get("primarias", {}).values()]
        if r.get("arbitro"):
            pareceres.append(tl._descrever_parecer(r["arbitro"]))
        saida.append({"id_rs": i, "titulo": unicos.get(i, {}).get("titulo", ""), "estrato": estrato.get(i),
                      "decisao_ia": r.get("decisao_final"), "decidido_por": r.get("decidido_por"),
                      "pareceres": " || ".join(pareceres), "referencia_humana": ref.get(i)})
    return saida


def _sem_carimbo(objeto):
    return {k: v for k, v in objeto.items() if k != "calculado_em"}


CRITERIOS_HUMANOS = ("kappa_humanos", "concordancia_humanos", "tamanho_calibracao")


def incluidos_minimos_para_ic(limiar=None, alfa=0.05):
    """Menor n com limite inferior de Clopper-Pearson de n acertos em n ≥ limiar (0,90 → 36)."""
    limiar = LIMIARES["recall_ic_inferior_min"] if limiar is None else limiar
    n = max(1, math.ceil(math.log(alfa / 2.0) / math.log(limiar)) - 1)
    while clopper_pearson(n, n, alfa)[0] < limiar:
        n += 1
    return n


def motivo_reprovacao(metricas):
    """Por que a amostra reprovou: None (não reprovou), largura_ic, humanos ou desempenho_ia.

    `largura_ic`: 0 falsos negativos e o único critério não atendido é o limite inferior do IC do recall, isto é,
    faltam incluídos humanos para o IC ficar estreito; revisar critérios não resolve.
    """
    if metricas.get("atende_limiares") is not False:
        return None
    falhos = sorted(k for k, c in (metricas.get("criterios") or {}).items() if c.get("atende") is not True)
    if falhos == ["recall_ic_inferior"] and not metricas.get("falsos_negativos"):
        return "largura_ic"
    if falhos and all(k in CRITERIOS_HUMANOS for k in falhos):
        return "humanos"
    return "desempenho_ia"


def incluidos_ia_nao_sorteados(raiz, metricas):
    """Registros que a IA mandou seguir (estrato ia_positivo) e que nenhuma amostra da rodada sorteou ainda."""
    populacao = (metricas.get("populacao") or {}).get(ESTRATO_POS, 0)
    usados = {a["id_rs"] for _, d in _desenhos_existentes(raiz, metricas["rodada"])
              for a in d.get("amostra", []) if a.get("estrato") == ESTRATO_POS}
    return max(0, populacao - len(usados))


def plano_reprovacao(raiz, metricas, motivo):
    """proxima_acao (e números de apoio) para uma amostra reprovada; `resumo` curto para a pendência."""
    finalidade = metricas["finalidade"]
    if motivo is None:
        return {"resumo": ""}
    if finalidade == "calibracao":
        texto = ("revisar os critérios (vN+1) com a dupla humana e recalibrar em amostra nova "
                 "(references/ia-validacao.md, seção 4 A)")
        return {"proxima_acao": texto, "resumo": texto}
    if finalidade == "desenvolvimento" and motivo != "largura_ic":
        texto = ("ajustar o prompt no conjunto de desenvolvimento; congelar e validar em amostra nova "
                 "(references/ia-validacao.md, seção 4 B e C)")
        return {"proxima_acao": texto, "resumo": texto}
    if motivo == "humanos":
        texto = ("a IA não é o problema: a concordância entre os humanos ficou abaixo do limiar. Discuta as "
                 "discordâncias e preencha decisao_consenso; se a concordância seguir baixa, recalibre os critérios "
                 "com a dupla humana (references/ia-validacao.md, seção 4 A e C)")
        return {"proxima_acao": texto, "resumo": texto}
    if motivo == "desempenho_ia":
        texto = ("analisar falsos negativos → critérios vN+1 em amostra nova → regra liberal → IA só prioriza → "
                 "dupla humana (references/ia-validacao.md, seção 4 E)")
        return {"proxima_acao": texto, "resumo": texto}
    necessarios = incluidos_minimos_para_ic()
    disponiveis = incluidos_ia_nao_sorteados(raiz, metricas)
    n_incl = metricas["n_incluidos_humanos"]
    regra = f" --regra {metricas['regra']}" if metricas.get("regra") not in (None, "consenso") else ""
    base = {"incluidos_humanos_necessarios": necessarios, "incluidos_ia_nao_sorteados": disponiveis}
    if disponiveis >= necessarios:
        alvo = min(max(necessarios, LIMIARES["incluidos_humanos_alvo"]), disponiveis)
        resumo = (f"não revise os critérios; amplie a amostra enriquecida (≥ {necessarios} incluídos humanos) ou a elusão")
        texto = (f"reprovada só pela largura do IC do recall (0 falsos negativos; {n_incl} incluídos humanos, são "
                 f"precisos ≥ {necessarios} para o limite inferior chegar a {LIMIARES['recall_ic_inferior_min']}): não "
                 "revise os critérios. Amplie a validação com amostra nova da mesma rodada, enriquecida: "
                 f"`validar amostrar --etapa {metricas['etapa']} --rodada {metricas['rodada']} --n <n> --semente <outra> "
                 f"--enriquecer-incluidos {alvo}{regra}` ({disponiveis} incluídos pela IA ainda não sorteados; os IDs já "
                 f"sorteados ficam de fora sozinhos), ou meça os perdidos com `validar elusao --etapa {metricas['etapa']} "
                 f"--rodada {metricas['rodada']} --n {LIMIARES['elusao_n_alvo']} --semente <outra>{regra}` "
                 "(references/ia-validacao.md, seção 4 C, D e F)")
    else:
        resumo = ("limiar inalcançável nesta rodada; remédio 5 (dupla humana) e aprovação humana do G4 com --forcar")
        texto = (f"reprovada só pela largura do IC do recall (0 falsos negativos; {n_incl} incluídos humanos) e a rodada "
                 f"só tem {disponiveis} incluídos pela IA ainda não sorteados (< {necessarios}): o limiar é inalcançável; "
                 "não revise os critérios. Siga o remédio 5 (dupla humana completa, IA como terceira leitura) e, com "
                 "as decisões humanas registradas, peça a aprovação humana do G4 com `portao G4 --aprovar --por "
                 f"{esquema.PAPEL_HUMANO_PADRAO} --forcar --motivo \"...\"` (references/ia-validacao.md, seção 4 D e E)")
    return {**base, "proxima_acao": texto, "resumo": resumo}


def plano_atalho(raiz, metricas, planilhas_usadas):
    """proxima_acao do atalho da variante rápida: aprovar o G4 ou completar dupla, κ ou segunda leitura."""
    atalho = metricas["atalho_rapida"]
    etapa, rodada = metricas["etapa"], metricas["rodada"]
    if atalho["atende"]:
        texto = (f"atalho da variante rápida completo (dupla humana em {atalho['fracao_dupla_humana']:.0%}, κ = "
                 f"{atalho['kappa_humanos']:.2f}, {atalho['n_excluidos_relidos']}/{atalho['n_excluidos_ia']} excluídos "
                 f"pela IA relidos): peça a aprovação humana do G4 e declare o atalho como limitação (PRISMA 2020 "
                 "item 23c; references/tipos-de-revisao.md, seção 6)")
        return {"proxima_acao": texto, "resumo": ""}
    faltas, resumo = [], []
    crit = atalho["criterios"]
    if not crit["fracao_dupla_humana"]["atende"]:
        minimo = math.ceil(esquema.FRACAO_MINIMA_DUPLA_RAPIDA * atalho["n_populacao"])
        faltas.append(f"dupla humana em {atalho['n_dupla_humana']} de {atalho['n_populacao']} registros (< "
                      f"{esquema.FRACAO_MINIMA_DUPLA_RAPIDA:.0%}): sorteie amostra maior, com >= {minimo} registros, "
                      f"`validar amostrar --etapa {etapa} --rodada {rodada} --n {minimo} --semente <outra> --atalho-rapida`")
        resumo.append("dupla humana abaixo de 20%")
    if not crit["kappa_humanos_calculado"]["atende"]:
        faltas.append("κ da dupla humana indefinido: os dois codificadores precisam preencher decisao_h1 e decisao_h2")
        resumo.append("κ indefinido")
    if not crit["segunda_leitura_excluidos"]["atende"]:
        planilhas_txt = " ".join(f"--planilha {p}" for p in planilhas_usadas)
        faltas.append(f"{atalho['n_faltam_reler']} excluídos pela IA sem leitura humana ({atalho['ids_faltam_reler'][:5]}"
                      f"...): sorteie todos com `validar elusao --etapa {etapa} --rodada {rodada} --n "
                      f"{atalho['n_excluidos_ia']} --semente <s>`, faça a segunda leitura, registre com `validar "
                      f"segunda-leitura --rodada {rodada} --planilha <planilha codificada>` e rode de novo `validar "
                      f"calcular {planilhas_txt}`")
        resumo.append(f"{atalho['n_faltam_reler']} excluídos pela IA sem segunda leitura")
    return {"proxima_acao": "; ".join(faltas) + " (references/tipos-de-revisao.md, seção 6)",
            "resumo": "; ".join(resumo)}


def cmd_calcular(args):
    raiz = estado.exigir_projeto(args.dir)
    metricas, linhas_fn = calcular_metricas(raiz, args.planilha, args.desenho, args.regra,
                                            finalidade=getattr(args, "finalidade", None),
                                            atalho_rapida=getattr(args, "atalho_rapida", False))
    pasta = pasta_validacao(raiz, metricas["rodada"])
    caminho_metricas = pasta / f"{metricas['amostra_id']}{esquema.SUFIXO_METRICAS}"
    caminho_fn = pasta / f"{metricas['amostra_id']}_falsos_negativos.csv"
    reexecucao = False
    if caminho_metricas.exists():
        with open(caminho_metricas, encoding="utf-8") as f:
            reexecucao = _sem_carimbo(json.load(f)) == _sem_carimbo(json.loads(json.dumps(metricas)))
    metricas["calculado_em"] = estado.agora()
    if not reexecucao:
        tl.escrever_json(caminho_metricas, metricas)
    tl.escrever_csv(caminho_fn, COLUNAS_FALSOS_NEGATIVOS, linhas_fn)
    rel_metricas, rel_fn = tl.relativo(raiz, caminho_metricas), tl.relativo(raiz, caminho_fn)
    etapa_projeto = tl.ETAPA_PROJETO[metricas["etapa"]]
    ia = metricas.get("ia_final") or {}
    par = next(iter(metricas["humanos"]["pares"].values()), {}) or {}
    compacto = {
        "tipo": metricas["tipo"], "finalidade": metricas["finalidade"], "sem_ia": metricas["sem_ia"],
        "amostra_id": metricas["amostra_id"], "rodada": metricas["rodada"],
        "etapa": metricas["etapa"], "regra": metricas["regra"], "n_amostra": metricas["n_amostra"],
        "n_codificados": metricas["n_codificados"], "n_incluidos_humanos": metricas["n_incluidos_humanos"],
        "ponderado": metricas.get("ponderado"),
        "sensibilidade": (ia.get("sensibilidade") or {}).get("valor"),
        "sensibilidade_ic": [(ia.get("sensibilidade") or {}).get("ic_inferior"),
                             (ia.get("sensibilidade") or {}).get("ic_superior")] if ia else None,
        "especificidade": (ia.get("especificidade") or {}).get("valor"),
        "precisao": (ia.get("precisao") or {}).get("valor"), "vpn": (ia.get("vpn") or {}).get("valor"),
        "kappa_ia": ia.get("kappa"), "pabak_ia": ia.get("pabak"), "wss": ia.get("wss"), "wss_95": ia.get("wss_95"),
        "kappa_humanos": par.get("kappa_binario"), "pabak_humanos": par.get("pabak_binario"),
        "concordancia_humanos": par.get("concordancia_binaria"),
        "elusao": metricas["elusao"]["taxa"]["valor"], "perdidos_estimados": metricas["elusao"]["perdidos_estimados"],
        "limiares": LIMIARES, "atende_limiares": metricas["atende_limiares"],
        "n_falsos_negativos": len(metricas["falsos_negativos"]),
        "motivo_reprovacao": motivo_reprovacao(metricas),
    }
    atalho = metricas.get("atalho_rapida")
    if atalho:
        # Campos lidos pelo G4 no caminho do atalho (esquema.CAMPOS_ATALHO_RAPIDA); kappa_humanos já está no compacto.
        compacto.update({c: atalho[c] for c in esquema.CAMPOS_ATALHO_RAPIDA})
        compacto["atalho_rapida_atende"] = atalho["atende"]
    dados_evento = dict(compacto)
    if dados_evento["atende_limiares"] is None:
        # Elusão não tem limiar vinculante: sem a chave, `status` continua lendo a última validação de desempenho.
        dados_evento.pop("atende_limiares")
    if not reexecucao:
        estado.registrar_evento(raiz, "validacao_calculada", etapa_projeto, "script", "validacao",
                                dados=dados_evento, artefatos=[rel_metricas, rel_fn]
                                + [p["caminho"] for p in metricas["planilhas"] if not Path(p["caminho"]).is_absolute()])
    with open(tl.resolver(raiz, metricas["desenho"]), encoding="utf-8") as f:
        desenho = json.load(f)
    portao = tl.PORTAO_ETAPA[metricas["etapa"]]
    finalidade = metricas["finalidade"]
    # No atalho da variante rápida o G4 não exige recall: reprova só se faltar dupla >= 20%, κ ou segunda leitura.
    reprovada = (not atalho["atende"]) if atalho else metricas["atende_limiares"] is False
    tl.sincronizar_pendencia(raiz, "validacao_humana", etapa_projeto, portao, "", 0, desenho["planilha"], "validacao")
    motivo = compacto["motivo_reprovacao"]
    plano = plano_atalho(raiz, metricas, args.planilha) if atalho else plano_reprovacao(raiz, metricas, motivo)
    # Só a validação e a calibração reprovadas pedem ação humana; desenvolvimento reprovado é o esperado.
    if atalho:
        descricao_validacao = f"validação {metricas['amostra_id']} (atalho da variante rápida) incompleta: {plano['resumo']}"
    elif motivo == "largura_ic":
        descricao_validacao = (f"validação {metricas['amostra_id']} reprovada só pela largura do IC (0 falsos negativos, "
                               f"{metricas['n_incluidos_humanos']} incluídos humanos): {plano['resumo']}")
    else:
        descricao_validacao = (f"validação {metricas['amostra_id']} abaixo dos limiares: analisar "
                               f"{len(metricas['falsos_negativos'])} falsos negativos ({rel_fn}) e seguir os remédios "
                               "de references/ia-validacao.md, seção 4 E")
    tl.sincronizar_pendencia(
        raiz, PENDENCIA_REPROVADA["validacao"], etapa_projeto, portao, descricao_validacao,
        max(1, len(metricas["falsos_negativos"])) if reprovada and finalidade == "validacao" else 0,
        rel_metricas, "validacao")
    tl.sincronizar_pendencia(
        raiz, PENDENCIA_REPROVADA["calibracao"], etapa_projeto, portao,
        f"calibração {metricas['amostra_id']} abaixo dos limiares (κ >= 0,6, >= 75% e tamanho): revisar os critérios "
        "(vN+1) e recalibrar em amostra nova (references/ia-validacao.md, seção 4 A)",
        1 if reprovada and finalidade == "calibracao" else 0, rel_metricas, "validacao")
    for aviso in metricas["avisos"]:
        print(f"aviso: {aviso}", file=sys.stderr)
    saida = {"comando": "validar calcular", "ok": not reprovada if atalho else metricas["atende_limiares"] is not False,
             "reexecucao": reexecucao, **compacto, "criterios": metricas["criterios"],
             "metricas": rel_metricas, "falsos_negativos": rel_fn, "avisos": metricas["avisos"]}
    if metricas["atende_limiares"] is False:
        saida["ids_falsos_negativos"] = metricas["falsos_negativos"]
    if atalho:
        saida["criterios_atalho_rapida"] = atalho["criterios"]
        saida["n_faltam_reler"] = atalho["n_faltam_reler"]
        saida.update({k: v for k, v in plano.items() if k != "resumo"})
    elif metricas["atende_limiares"] is False:
        saida.update({k: v for k, v in plano.items() if k != "resumo"})
    estado.resumo(saida)
    return 2 if reprovada else 0


# ---------------------------------------------------------------------------
# Segunda leitura humana dos excluídos pela IA (atalho da variante rápida)
# ---------------------------------------------------------------------------
COLUNAS_DECISAO_SEGUNDA_LEITURA = ("decisao_consenso", "decisao_humana", "decisao_segunda_leitura", "decisao")
MOTIVO_FILA_SEGUNDA_LEITURA = "segunda_leitura"


def ler_planilhas_segunda_leitura(raiz, planilhas_humanas):
    """{id_rs: {"relido": bool, "decisao": incluir|excluir|incerto|None}} e as fontes (caminho, sha256).

    Aceita a planilha cega de `validar elusao`/`amostrar` codificada (decisao_hN, decisao_consenso) ou qualquer
    planilha com id_rs e decisao_humana (ou decisao_segunda_leitura, decisao). Linha com algum código humano
    conta como relida; a decisão é o consenso, a decisao_humana ou a concordância dos codificadores.
    """
    leituras, fontes, erros, avisos = {}, [], [], []
    for p in planilhas_humanas:
        caminho = tl.resolver(raiz, p)
        colunas, linhas, info = tl.ler_tabela_humana(caminho, ["id_rs"], aba="codificacao")
        avisos.extend(info["avisos"])
        colunas_dec = [c for c in (c.strip().lower() for c in colunas)
                       if c.startswith("decisao_h") or c in COLUNAS_DECISAO_SEGUNDA_LEITURA]
        if not colunas_dec:
            raise tl.ErroUso(f"{caminho.name}: sem coluna de decisão humana (decisao_humana, decisao_h1/decisao_h2 ou "
                             "decisao_consenso); o arquivo não foi alterado")
        fontes.append({"caminho": tl.relativo(raiz, caminho), "sha256": estado.sha256_arquivo(caminho)})
        for numero, bruta in zip(info["numeros"], linhas):
            linha = {k.strip().lower(): v for k, v in bruta.items() if k}
            id_rs = str(linha.get("id_rs") or "").strip().upper()
            if not id_rs:
                continue
            codigos, diretas = {}, {}
            for coluna in colunas_dec:
                try:
                    codigo = normalizar_codigo(linha.get(coluna))
                except ValueError as e:
                    erros.append(f"{caminho.name} linha {numero} ({id_rs}), {coluna}: {e}")
                    continue
                if coluna in COLUNAS_DECISAO_SEGUNDA_LEITURA:
                    diretas[coluna] = codigo
                else:
                    codigos[coluna[len("decisao_"):]] = codigo
            # consenso > decisao_humana > decisao_segunda_leitura > decisao (ordem da tupla)
            direta = next((diretas[c] for c in COLUNAS_DECISAO_SEGUNDA_LEITURA if diretas.get(c) is not None), None)
            relido = direta is not None or any(v is not None for v in codigos.values())
            if not relido:
                continue
            decisao = direta if direta is not None else referencia_humana(codigos, None)[0]
            leituras[id_rs] = {"relido": True, "decisao": decisao}
    if erros:
        raise tl.ErroUso("planilha da segunda leitura com códigos inválidos:\n  " + "\n  ".join(erros[:30]))
    return leituras, fontes, avisos


def cmd_segunda_leitura(args):
    raiz = estado.exigir_projeto(args.dir)
    rodada = args.rodada or tl.rodada_ativa(raiz, args.etapa)
    if not rodada:
        raise tl.ErroUso("sem rodada ativa: informe --rodada (a rodada consolidada da triagem por IA)")
    tl._validar_nome(rodada, "rodada")
    if tl.eh_rodada_estabilidade(rodada, tl.rodadas_estabilidade(raiz)):
        raise tl.ErroUso(f"{rodada} é rodada de estabilidade; a segunda leitura é da rodada consolidada")
    unicos = tl.ler_unicos(raiz, obrigatorio=False)
    res = _decisoes_ia(raiz, args.etapa, rodada, args.regra, unicos)
    excluidos = ids_excluidos_pela_ia(raiz, args.etapa, rodada, args.regra, unicos, res)
    leituras, fontes, avisos = ler_planilhas_segunda_leitura(raiz, args.planilha)
    caminho = caminho_segunda_leitura(raiz, rodada)
    anterior = ler_segunda_leitura(raiz, rodada) or {}
    decisoes = {}
    if anterior.get("regra", args.regra) == args.regra and anterior.get("etapa", args.etapa) == args.etapa:
        decisoes.update(anterior.get("decisoes") or {})  # leituras em partes acumulam; a mais nova vale
        fontes = [f for f in anterior.get("planilhas") or [] if f.get("caminho") not in {x["caminho"] for x in fontes}] \
            + fontes
    decisoes.update({i: l["decisao"] for i, l in leituras.items()})
    conjunto = set(excluidos)
    fora = sorted(i for i in leituras if i not in conjunto)
    if fora:
        avisos.append(f"{len(fora)} IDs da planilha não são excluídos pela IA nesta rodada (ignorados na contagem): "
                      f"{fora[:10]}")
    relidos = sorted(i for i in decisoes if i in conjunto)
    # Excluídos já codificados por humanos na amostra do atalho (último `calcular` com atalho) também contam como lidos.
    ultima = next((ev for ev in reversed(estado.ler_log(raiz)) if ev.get("evento") == "validacao_calculada"
                   and (ev.get("dados") or {}).get("rodada") == rodada
                   and (ev.get("dados") or {}).get(esquema.CRITERIO_ATALHO_RAPIDA) is True), None)
    na_amostra = set()
    for artefato in (ultima or {}).get("artefatos") or []:
        if str(artefato.get("caminho", "")).endswith(esquema.SUFIXO_METRICAS):
            try:
                with open(tl.resolver(raiz, artefato["caminho"]), encoding="utf-8") as f:
                    na_amostra = set((json.load(f).get("atalho_rapida") or {}).get("ids_relidos_na_amostra") or [])
            except (OSError, json.JSONDecodeError):
                na_amostra = set()
    na_amostra &= conjunto
    faltam = sorted(conjunto - set(relidos) - na_amostra)
    resgatados = sorted(i for i in relidos if decisoes[i] is not None and tl.positiva(decisoes[i]))
    sem_decisao = sorted(i for i in relidos if decisoes[i] is None)
    resultado = {
        "tipo": esquema.TIPO_SEGUNDA_LEITURA, "rodada": rodada, "etapa": args.etapa, "regra": args.regra,
        "planilhas": fontes, "n_excluidos_ia": len(excluidos), "n_excluidos_relidos": len(conjunto) - len(faltam),
        "n_relidos_nas_planilhas": len(relidos), "n_relidos_na_amostra_do_atalho": len(na_amostra - set(relidos)),
        "segunda_leitura_excluidos": not faltam, "n_faltam": len(faltam), "ids_faltam": faltam[:200],
        "n_resgatados": len(resgatados), "ids_resgatados": resgatados, "n_sem_decisao": len(sem_decisao),
        "ids_sem_decisao": sem_decisao, "n_fora_dos_excluidos": len(fora), "ids_relidos": relidos,
        "decisoes": {i: decisoes[i] for i in sorted(decisoes)},
    }
    reexecucao = bool(anterior) and _sem_carimbo(anterior) == _sem_carimbo(json.loads(json.dumps(resultado)))
    resultado["calculado_em"] = estado.agora()
    rel = tl.relativo(raiz, caminho)
    rel_resgatados = tl.relativo(raiz, pasta_validacao(raiz, rodada) / "segunda_leitura_resgatados.csv")
    etapa_projeto = tl.ETAPA_PROJETO[args.etapa]
    if not reexecucao:
        tl.escrever_json(caminho, resultado)
        fila = []
        for i in resgatados + sem_decisao:
            reg, r = unicos.get(i, {}), res.get(i, {})
            pareceres = " || ".join(tl._descrever_parecer(x) for x in (r.get("primarias") or {}).values())
            fila.append({"id_rs": i, "titulo": reg.get("titulo", ""), "resumo": reg.get("resumo", ""),
                         "ano": reg.get("ano", ""), "veiculo": reg.get("veiculo", ""),
                         "motivo_fila": MOTIVO_FILA_SEGUNDA_LEITURA, "pareceres": pareceres,
                         "decisao_humana": decisoes[i] or "", "criterio_humano": "",
                         "motivo_humano": "segunda leitura humana dos excluídos pela IA" if decisoes[i] else ""})
        tl.escrever_csv(tl.resolver(raiz, rel_resgatados), tl.COLUNAS_FILA_HUMANA, fila)
        taxa = (len(resgatados) / len(relidos)) if relidos else None
        estado.registrar_evento(
            raiz, "validacao_calculada", etapa_projeto, "script", "validacao",
            dados={"tipo": esquema.TIPO_SEGUNDA_LEITURA, "finalidade": "elusao", "etapa": args.etapa, "rodada": rodada,
                   "regra": args.regra, "amostra_id": esquema.TIPO_SEGUNDA_LEITURA,
                   "n_excluidos_ia": len(excluidos), "n_excluidos_relidos": len(conjunto) - len(faltam),
                   "segunda_leitura_excluidos": not faltam, "n_resgatados": len(resgatados),
                   "n_sem_decisao": len(sem_decisao), "n_codificados": len(relidos),
                   "elusao": taxa, "perdidos_estimados": len(resgatados)},
            artefatos=[rel, rel_resgatados] + [f["caminho"] for f in fontes
                                               if not Path(f["caminho"]).is_absolute()
                                               and tl.resolver(raiz, f["caminho"]).is_file()])
    for aviso in avisos:
        print(f"aviso: {aviso}", file=sys.stderr)
    passos = []
    if resgatados or sem_decisao:
        passos.append(f"{len(resgatados)} excluídos pela IA que a segunda leitura manda seguir"
                      + (f" e {len(sem_decisao)} sem decisão (discordância sem consenso)" if sem_decisao else "")
                      + f": confira {rel_resgatados}, aplique com `triagem override --fila {rel_resgatados} --rodada "
                        f"{rodada}` e rode `triagem consolidar --rodada {rodada}`")
    if faltam:
        passos.append(f"faltam {len(faltam)} excluídos pela IA sem leitura humana ({faltam[:5]}): complete a planilha e "
                      "rode este comando de novo (as leituras acumulam)")
    if ultima:
        planilhas_validacao = [a.get("caminho") for a in ultima.get("artefatos") or []
                               if str(a.get("caminho", "")).endswith((".xlsx", ".xlsm", ".csv", ".tsv"))
                               and "_falsos_negativos" not in str(a.get("caminho"))]
        passos.append("grave a segunda leitura na validação do atalho: `validar calcular "
                      + " ".join(f"--planilha {c}" for c in planilhas_validacao) + "`"
                      + (" (com --desenho, se a planilha for CSV)" if any(c.endswith((".csv", ".tsv"))
                                                                          for c in planilhas_validacao) else ""))
    else:
        passos.append("a validação do atalho (`validar amostrar --atalho-rapida` e `validar calcular`) lê este "
                      "arquivo ao ser calculada")
    estado.resumo({
        "comando": "validar segunda-leitura", "ok": True, "reexecucao": reexecucao, "rodada": rodada,
        "etapa": args.etapa, "regra": args.regra, "n_excluidos_ia": len(excluidos),
        "n_excluidos_relidos": len(conjunto) - len(faltam), "n_relidos_nas_planilhas": len(relidos),
        "n_relidos_na_amostra_do_atalho": len(na_amostra - set(relidos)),
        "segunda_leitura_excluidos": not faltam, "n_faltam": len(faltam),
        "n_resgatados": len(resgatados), "n_sem_decisao": len(sem_decisao), "n_fora_dos_excluidos": len(fora),
        "arquivo": rel, "resgatados": rel_resgatados, "avisos": avisos, "proxima_acao": "; ".join(passos),
    })
    return 0


# ---------------------------------------------------------------------------
# Estabilidade (reexecução de 5–10%)
# ---------------------------------------------------------------------------
def _revisores_ia(vigentes):
    return sorted({l["revisor"] for l in vigentes
                   if str(l.get("tipo_ator", "")).startswith("ia_") and l["revisor"] != tl.REVISOR_ARBITRO
                   and not tl.eh_override(l)})


def cmd_estabilidade(args):
    raiz = estado.exigir_projeto(args.dir)
    tl._validar_nome(args.rodada, "rodada")
    if not 0 < args.amostra <= 1:
        raise tl.ErroUso("--amostra é uma fração entre 0 e 1 (ex.: 0.1)")
    rodada_re = args.rodada_reexecucao or f"{args.rodada}{tl.SUFIXO_ESTABILIDADE}"
    tl._validar_nome(rodada_re, "--rodada-reexecucao")
    if tl.eh_rodada_estabilidade(args.rodada):
        raise tl.ErroUso(f"{args.rodada} já é uma rodada de reexecução; informe a rodada original")
    finalidade = conferir_finalidade("estabilidade", getattr(args, "finalidade", None))
    avisos = []
    if not 0.05 <= args.amostra <= 0.10:
        avisos.append("references/ia-validacao.md (seção 4 G) recomenda reexecutar 5–10% dos registros")
    todas = tl.ler_decisoes(raiz)
    vig_orig = tl.decisoes_vigentes(todas, args.etapa, [args.rodada])
    revisores = _revisores_ia(vig_orig)
    if args.revisores:
        revisores = [r for r in revisores if r in set(args.revisores.split(","))]
    if not revisores:
        raise tl.ErroUso(f"nenhum revisor de IA com decisões na rodada {args.rodada}")
    pasta = pasta_validacao(raiz, args.rodada)
    arq_desenho = pasta / esquema.ARQ_DESENHO_ESTABILIDADE
    parametros = {"rodada": args.rodada, "rodada_reexecucao": rodada_re, "amostra": args.amostra,
                  "semente": args.semente, "revisores": revisores, "etapa": args.etapa}
    if finalidade != FINALIDADE_PADRAO["estabilidade"]:
        parametros["finalidade"] = finalidade  # só fora do padrão: desenhos antigos seguem válidos
    desenho = None
    if arq_desenho.exists():
        with open(arq_desenho, encoding="utf-8") as f:
            desenho = json.load(f)
        if desenho["parametros"] != parametros:
            raise tl.ErroUso(f"já existe desenho de estabilidade com outros parâmetros em {tl.relativo(raiz, arq_desenho)}")
        avisos.extend(avisos_desenho_com_inativos(raiz, desenho))
    etapa_projeto = tl.ETAPA_PROJETO[args.etapa]
    if desenho is None:
        inativos = tl.ids_inativos(tl.ler_unicos(raiz, obrigatorio=False))  # buscas substituídas: fora do fluxo
        pool_rodada = {l["id_rs"] for l in vig_orig if l["revisor"] in revisores}
        pool = sorted(pool_rodada - inativos)
        if not pool:
            raise tl.ErroUso(f"nenhum registro ativo com decisão de IA na rodada {args.rodada}")
        k = max(1, math.ceil(args.amostra * len(pool)))
        ids = sorted(random.Random(f"{args.semente}:{args.rodada}:estabilidade").sample(pool, min(k, len(pool))))
        desenho = {"versao": 1, "parametros": parametros, "criado_em": estado.agora(), "ids": ids,
                   "n_inativos_ignorados": len(pool_rodada & inativos)}
        tl.escrever_json(arq_desenho, desenho)
        tl.escrever_csv(pasta / "estabilidade_ids.csv", ["id_rs"], [{"id_rs": i} for i in ids])
    ids = desenho["ids"]
    vig_re = tl.decisoes_vigentes(todas, args.etapa, [rodada_re])

    if args.acao == "calcular" and not vig_re:
        raise tl.ErroUso(f"a rodada {rodada_re} ainda não tem decisões; prepare e mescle a reexecução antes")
    if args.acao == "preparar" or (args.acao == "auto" and not vig_re):
        preparados, instrucoes_api = [], []
        for rev in revisores:
            manifesto = tl.ler_manifesto(raiz, args.rodada, rev)
            criterios = args.criterios or (manifesto or {}).get("criterios_arquivo")
            if not criterios:
                instrucoes_api.append(f"revisor {rev}: rodada sem lotes (modo API); reexecute os IDs de "
                                      f"{tl.relativo(raiz, pasta / 'estabilidade_ids.csv')} como rodada {rodada_re}")
                continue
            if manifesto and tl.sha_opcional(tl.resolver(raiz, criterios)) != manifesto["criterios_sha"]:
                raise tl.ErroUso(f"critérios de {rev} mudaram desde a rodada {args.rodada}; estabilidade exige o mesmo prompt")
            r = tl.preparar_lotes(raiz, args.etapa, rodada_re, rev, criterios, args.tamanho, args.semente, ids=ids)
            if r["lotes_novos"]:
                estado.registrar_evento(
                    raiz, "lote_preparado", etapa_projeto, "script", "validacao",
                    dados={"rodada": rodada_re, "etapa": args.etapa, "revisor": rev, "lotes_novos": r["lotes_novos"],
                           "registros_novos": r["registros_novos"], "semente": args.semente,
                           "criterios_sha": r["manifesto"]["criterios_sha"], "finalidade": finalidade,
                           "rodada_original": args.rodada},
                    artefatos=[tl.relativo(raiz, arq_desenho), tl.relativo(raiz, tl.caminho_manifesto(raiz, rodada_re, rev))])
            preparados.append({"revisor": rev, "lotes_novos": len(r["lotes_novos"]), "lotes_pendentes": [
                {"lote": str(tl.resolver(raiz, l["arquivo"])), "resposta": str(tl.resolver(raiz, l["resposta"]))}
                for l in r["pendentes"]]})
        for aviso in avisos + instrucoes_api:
            print(f"aviso: {aviso}", file=sys.stderr)
        estado.resumo({"comando": "validar estabilidade", "ok": True, "acao": "preparar", "rodada": args.rodada,
                       "rodada_reexecucao": rodada_re, "n_ids": len(ids), "revisores": preparados,
                       "instrucoes_api": instrucoes_api, "avisos": avisos,
                       "proxima_acao": f"despachar os lotes, `triagem mesclar --rodada {rodada_re} --revisor <R>` e "
                                       f"rodar `validar estabilidade --rodada {args.rodada}` de novo"})
        return 0

    resultado = {"rodada": args.rodada, "rodada_reexecucao": rodada_re, "etapa": args.etapa, "n_ids": len(ids),
                 "finalidade": finalidade, "revisores": {}}
    for rev in revisores:
        orig = {l["id_rs"]: l["decisao"] for l in vig_orig if l["revisor"] == rev and not tl.eh_override(l)}
        reex = {l["id_rs"]: l["decisao"] for l in vig_re if l["revisor"] == rev and not tl.eh_override(l)}
        comuns = [i for i in ids if i in orig and i in reex]
        faltam = [i for i in ids if i in orig and i not in reex]
        if faltam:
            avisos.append(f"revisor {rev}: {len(faltam)} IDs ainda sem reexecução")
        a = [orig[i] for i in comuns]
        b = [reex[i] for i in comuns]
        a_bin = [tl.positiva(v) for v in a]
        b_bin = [tl.positiva(v) for v in b]
        mudancas = [{"id_rs": i, "original": orig[i], "reexecucao": reex[i]} for i in comuns if orig[i] != reex[i]]
        resultado["revisores"][rev] = {
            "n": len(comuns), "faltando": len(faltam),
            "concordancia_3cat": concordancia(a, b), "kappa_3cat": kappa_cohen(a, b, ["incluir", "excluir", "incerto"]),
            "concordancia_binaria": concordancia(a_bin, b_bin), "kappa_binario": kappa_cohen(a_bin, b_bin, [True, False]),
            "pabak_binario": pabak(a_bin, b_bin, 2),
            "mudancas_de_seguimento": sum(1 for x, y in zip(a_bin, b_bin) if x != y),
            "mudancas": mudancas,
        }
    resultado["atende_limiares"] = None
    resultado["nota"] = ("não há limiar vinculante para estabilidade (references/ia-validacao.md, seção 4 G); "
                         "relate concordância e mudanças")
    resultado["avisos"] = avisos
    arq = pasta / "estabilidade.json"
    reexecucao = False
    if arq.exists():
        with open(arq, encoding="utf-8") as f:
            reexecucao = _sem_carimbo(json.load(f)) == _sem_carimbo(json.loads(json.dumps(resultado)))
    resultado["calculado_em"] = estado.agora()
    if not reexecucao:
        tl.escrever_json(arq, resultado)
        estado.registrar_evento(
            raiz, "validacao_calculada", etapa_projeto, "script", "validacao",
            dados={"tipo": "estabilidade", "finalidade": finalidade, "etapa": args.etapa, "rodada": args.rodada,
                   "rodada_reexecucao": rodada_re, "n_ids": len(ids),
                   "revisores": {r: {k: v for k, v in m.items() if k != "mudancas"}
                                 for r, m in resultado["revisores"].items()}},
            artefatos=[tl.relativo(raiz, arq)])
    for aviso in avisos:
        print(f"aviso: {aviso}", file=sys.stderr)
    estado.resumo({"comando": "validar estabilidade", "ok": True, "acao": "calcular", "reexecucao": reexecucao,
                   "rodada": args.rodada, "rodada_reexecucao": rodada_re, "finalidade": finalidade,
                   "revisores": {r: {k: v for k, v in m.items() if k != "mudancas"}
                                 for r, m in resultado["revisores"].items()},
                   "arquivo": tl.relativo(raiz, arq), "avisos": avisos})
    return 0


# ---------------------------------------------------------------------------
# Registro no dispatcher
# ---------------------------------------------------------------------------
def _protegido(funcao):
    def executar(args):
        try:
            return funcao(args)
        except tl.ErroDependencia as e:
            print(f"erro: {e}", file=sys.stderr)
            estado.resumo({"comando": f"validar {getattr(args, 'subcomando', '')}".strip(), "ok": False, "erro": str(e),
                           "dependencia_ausente": True})
            return 3
        except (tl.ErroUso, estado.ErroProjeto, ValueError) as e:
            print(f"erro: {e}", file=sys.stderr)
            estado.resumo({"comando": f"validar {getattr(args, 'subcomando', '')}".strip(), "ok": False, "erro": str(e)})
            return 1

    return executar


def registrar(subparsers):
    sub = tl.subparsers_do_comando(subparsers, "validar", "validação humana da triagem por IA (amostra cega, métricas, elusão)")

    p = sub.add_parser("amostrar", help="amostra cega estratificada para dupla codificação humana")
    p.add_argument("--etapa", default="ta", choices=esquema.ETAPAS_DECISAO)
    p.add_argument("--rodada", required=True)
    p.add_argument("--n", type=int, default=100, help="tamanho da amostra aleatória simples (padrão 100)")
    p.add_argument("--semente", type=int, required=True)
    p.add_argument("--enriquecer-incluidos", type=int, default=None,
                   help="completa a amostra até N incluídos/incertos pela IA (alvo 60; references/ia-validacao.md, 4 C)")
    p.add_argument("--elusao", type=int, default=0,
                   help="completa até N excluídos pela IA (alvo 300; references/ia-validacao.md, 4 F)")
    p.add_argument("--regra", default="consenso", choices=["consenso", "liberal"])
    p.add_argument("--codificadores", type=int, default=2)
    p.add_argument("--finalidade", default=None, choices=esquema.FINALIDADES_VALIDACAO,
                   help="calibracao | desenvolvimento | validacao (padrão; só ela decide o G4)")
    p.add_argument("--excluir-ids", default=None,
                   help="CSV com id_rs (ou um por linha) a tirar do quadro amostral: IDs usados em outra finalidade")
    p.add_argument("--ids", default=None, help="restringe o quadro amostral a estes id_rs")
    p.add_argument("--sem-ia", action="store_true",
                   help="calibração humana sem decisões de IA (planilha para dois humanos; métricas só entre humanos)")
    p.add_argument("--criterios", default=None, help="arquivo de critérios, quando a rodada não tem lotes")
    p.add_argument("--atalho-rapida", action="store_true",
                   help="amostra do atalho da variante rápida: dupla humana em >= 20%% da rodada; `calcular` grava os "
                        "campos do atalho (references/tipos-de-revisao.md, seção 6)")
    p.set_defaults(func=_protegido(cmd_amostrar))

    p = sub.add_parser("elusao", help="amostra cega só dos excluídos pela IA, após a rodada completa")
    p.add_argument("--etapa", default="ta", choices=esquema.ETAPAS_DECISAO)
    p.add_argument("--rodada", required=True)
    p.add_argument("--n", type=int, default=300)
    p.add_argument("--semente", type=int, required=True)
    p.add_argument("--regra", default="consenso", choices=["consenso", "liberal"])
    p.add_argument("--codificadores", type=int, default=1)
    p.add_argument("--finalidade", default=None, choices=esquema.FINALIDADES_VALIDACAO,
                   help="elusao (padrão) ou desenvolvimento")
    p.add_argument("--excluir-ids", default=None, help="id_rs a tirar do quadro amostral")
    p.set_defaults(func=_protegido(cmd_elusao))

    p = sub.add_parser("calcular", help="κ, PABAK, sensibilidade/especificidade com IC, WSS, elusão, limiares")
    p.add_argument("--planilha", required=True, action="append", help="xlsx/csv codificado (repita para juntar)")
    p.add_argument("--desenho", default=None, help="desenho da amostra (padrão: lido da aba _meta)")
    p.add_argument("--regra", default=None, choices=["consenso", "liberal"], help="padrão: a do desenho")
    p.add_argument("--finalidade", default=None, choices=esquema.FINALIDADES_VALIDACAO,
                   help="padrão: a do desenho (validacao em desenhos antigos); divergir do desenho é erro")
    p.add_argument("--atalho-rapida", action="store_true",
                   help="grava os campos do atalho da variante rápida mesmo com desenho sorteado sem --atalho-rapida "
                        "(só em projeto.variante = rapida e finalidade validacao)")
    p.set_defaults(func=_protegido(cmd_calcular))

    p = sub.add_parser("segunda-leitura", help="registra a segunda leitura humana dos excluídos pela IA (atalho da "
                                               "variante rápida)")
    p.add_argument("--planilha", required=True, action="append",
                   help="planilha codificada (xlsx ou CSV com , ; ou tab): a de `validar elusao` ou uma com id_rs e "
                        "decisao_humana; repita para juntar partes")
    p.add_argument("--rodada", default=None, help="padrão: rodada ativa (versoes_ativas)")
    p.add_argument("--etapa", default="ta", choices=esquema.ETAPAS_DECISAO)
    p.add_argument("--regra", default="consenso", choices=["consenso", "liberal"],
                   help="regra da decisão só IA que define os excluídos (padrão consenso)")
    p.set_defaults(func=_protegido(cmd_segunda_leitura))

    p = sub.add_parser("estabilidade", help="reexecuta 5–10%% dos registros e mede a concordância consigo mesmo")
    p.add_argument("--etapa", default="ta", choices=esquema.ETAPAS_DECISAO)
    p.add_argument("--rodada", required=True)
    p.add_argument("--amostra", type=float, default=0.1, help="fração dos registros (padrão 0.1)")
    p.add_argument("--semente", type=int, default=7)
    p.add_argument("--revisores", default=None, help="lista separada por vírgula (padrão: todos os de IA)")
    p.add_argument("--rodada-reexecucao", default=None, help="padrão: <rodada>_estab")
    p.add_argument("--criterios", default=None, help="necessário se a rodada original não tem lotes (modo API)")
    p.add_argument("--tamanho", type=int, default=None)
    p.add_argument("--acao", default="auto", choices=["auto", "preparar", "calcular"])
    p.add_argument("--finalidade", default=None, choices=esquema.FINALIDADES_VALIDACAO,
                   help="estabilidade (padrão) ou desenvolvimento")
    p.set_defaults(func=_protegido(cmd_estabilidade))
