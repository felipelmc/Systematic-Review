"""Testes da validação da triagem: estatística (κ, Fleiss, PABAK, Clopper-Pearson, pesos),
amostra cega, cálculo com limiares do Apêndice B, elusão e estabilidade."""

import json
import random

import pytest

from test_triagem_lotes import (criar_registros, eventos, marcar_inativos, responder_todos, rodar, triar_rodada)

from rslib import esquema, estado
from rslib import triagem_lotes as tl
from rslib import validacao as va


# ---------------------------------------------------------------------------
# Estatística
# ---------------------------------------------------------------------------
def test_kappa_cohen_igual_scikit_learn():
    sklearn_metrics = pytest.importorskip("sklearn.metrics")
    rng = random.Random(3)
    for categorias in (["incluir", "excluir"], ["incluir", "excluir", "incerto"]):
        for _ in range(20):
            n = rng.randint(10, 80)
            a = [rng.choice(categorias) for _ in range(n)]
            b = [x if rng.random() < 0.7 else rng.choice(categorias) for x in a]
            assert va.kappa_cohen(a, b) == pytest.approx(sklearn_metrics.cohen_kappa_score(a, b), abs=1e-12)


def test_kappa_exclui_na_par_a_par():
    a = ["incluir", "excluir", None, "incluir", "excluir", "incluir", None]
    b = ["incluir", "incluir", "excluir", "incluir", "excluir", None, "incluir"]
    filtrados = [(x, y) for x, y in zip(a, b) if x is not None and y is not None]
    esperado = va.kappa_cohen([x for x, _ in filtrados], [y for _, y in filtrados])
    assert va.kappa_cohen(a, b) == pytest.approx(esperado)
    # a == b em 3 de 4 pares válidos; manualmente: p_o = 0,75, p_e = 0,5·0,75 + 0,5·0,25 = 0,5
    assert va.kappa_cohen(a, b) == pytest.approx(0.5)
    assert va.concordancia(a, b) == pytest.approx(0.75)
    assert va.kappa_cohen(["incluir"] * 5, ["incluir"] * 5) is None, "κ indefinido sem variação"
    assert va.kappa_cohen([None], ["incluir"]) is None


def test_pabak():
    a = [True, True, False, False, True]
    b = [True, False, False, False, True]
    assert va.pabak(a, b) == pytest.approx(0.6)
    assert va.pabak(a, b, k=3) == pytest.approx((3 * 0.8 - 1) / 2)
    assert va.pabak([None, True], [False, True]) == pytest.approx(1.0)


def test_fleiss_exemplo_classico():
    # Exemplo de Fleiss (1971) reproduzido na literatura: 10 itens, 14 avaliadores, 5 categorias, κ ≈ 0,210.
    contagens = [
        [0, 0, 0, 0, 14], [0, 2, 6, 4, 2], [0, 0, 3, 5, 6], [0, 3, 9, 2, 0], [2, 2, 8, 1, 1],
        [7, 7, 0, 0, 0], [3, 2, 6, 3, 0], [2, 5, 3, 2, 2], [6, 5, 2, 1, 0], [0, 2, 2, 3, 7],
    ]
    itens = [[cat for cat, n in enumerate(linha) for _ in range(n)] for linha in contagens]
    assert va.kappa_fleiss(itens) == pytest.approx(0.2099, abs=5e-4)
    # com 2 avaliadores e sem NA, Fleiss difere de Cohen só pela p_e combinada
    assert va.kappa_fleiss([[1, 1], [0, 0], [1, 0], [0, 0]]) == pytest.approx(7 / 15)
    assert va.kappa_fleiss([[1, None], [None, None]]) is None


def test_clopper_pearson_valores_conhecidos_e_scipy():
    inf, sup = va.clopper_pearson(0, 10)
    assert inf == 0.0 and sup == pytest.approx(1 - 0.025 ** (1 / 10), abs=1e-9)
    inf, sup = va.clopper_pearson(10, 10)
    assert sup == 1.0 and inf == pytest.approx(0.025 ** (1 / 10), abs=1e-9)
    inf, sup = va.clopper_pearson(5, 10)
    assert (inf, sup) == (pytest.approx(0.187086, abs=1e-6), pytest.approx(0.812914, abs=1e-6))
    assert va.clopper_pearson(3, 0) == (None, None)
    stats = pytest.importorskip("scipy.stats")
    for x, n in [(1, 7), (57, 60), (60, 60), (143, 150), (12.4, 30.7), (980, 1000), (3, 2500)]:
        inf, sup = va.clopper_pearson(x, n)
        esp_inf = 0.0 if x == 0 else stats.beta.ppf(0.025, x, n - x + 1)
        esp_sup = 1.0 if x >= n else stats.beta.ppf(0.975, x + 1, n - x)
        assert inf == pytest.approx(esp_inf, abs=1e-8) and sup == pytest.approx(esp_sup, abs=1e-8)


def test_limite_inferior_de_recall_perfeito():
    # 36 de 36 detectados já passa de 0,90; 35 de 35, não (motivo para enriquecer a amostra).
    assert va.clopper_pearson(36, 36)[0] >= 0.90 > va.clopper_pearson(35, 35)[0]


def test_proporcao_e_metricas_ponderadas():
    iguais = va.proporcao([(2.0, True)] * 9 + [(2.0, False)])
    assert iguais["valor"] == pytest.approx(0.9)
    assert (iguais["ic_inferior"], iguais["ic_superior"]) == va.clopper_pearson(9, 10)
    assert va.proporcao([])["valor"] is None

    # Estrato IA+ (N=100): 50 amostrados, todos incluídos por humanos. Estrato IA− (N=900): 100
    # amostrados, 2 incluídos por humanos (falsos negativos). Sem pesos o recall seria 50/52.
    estrato = {}
    pred, ref = {}, {}
    for k in range(50):
        i = f"RS{k + 1:04d}"
        estrato[i], pred[i], ref[i] = va.ESTRATO_POS, "incluir", "incluir"
    for k in range(100):
        i = f"RS{k + 101:04d}"
        estrato[i], pred[i] = va.ESTRATO_NEG, "excluir"
        ref[i] = "incluir" if k < 2 else "excluir"
    ref["RS0150"] = None  # não revisado: fora
    populacao = {va.ESTRATO_POS: 100, va.ESTRATO_NEG: 900}
    itens, ponderavel = va._itens_ponderados(sorted(estrato), pred, ref, estrato, populacao)
    assert ponderavel and len(itens) == 149
    m = va.metricas_classificador(itens)
    assert m["matriz"] == {"tp": 50, "fp": 0, "fn": 2, "tn": 97}
    peso_neg = 900 / 99
    assert m["sensibilidade"]["valor"] == pytest.approx(100 / (100 + 2 * peso_neg))
    assert m["especificidade"]["valor"] == pytest.approx(1.0)
    assert m["vpn"]["valor"] == pytest.approx(97 / 99)
    assert m["falsos_negativos"] == ["RS0101", "RS0102"]
    fracao_nao_lida = 900 / 1000
    assert m["wss"] == pytest.approx(fracao_nao_lida - (1 - m["sensibilidade"]["valor"]))
    assert m["wss_95"] is None
    assert m["sensibilidade"]["ic_inferior"] < m["sensibilidade"]["valor"] < m["sensibilidade"]["ic_superior"]

    sem_amostra_neg = {i: e for i, e in estrato.items() if e == va.ESTRATO_POS}
    _, ponderavel = va._itens_ponderados(sorted(sem_amostra_neg), pred, ref, sem_amostra_neg, populacao)
    assert not ponderavel


def test_normalizar_codigo():
    assert va.normalizar_codigo("Incluir ") == "incluir"
    assert va.normalizar_codigo("EXCLUÍDO") == "excluir"
    assert va.normalizar_codigo("?") == "incerto"
    for vazio in (None, "", "NA", "não revisado", "nao_revisado"):
        assert va.normalizar_codigo(vazio) is None
    with pytest.raises(ValueError):
        va.normalizar_codigo("talvez sim")


def test_referencia_humana():
    assert va.referencia_humana({"h1": "incluir", "h2": "excluir"}, "excluir") == ("excluir", "consenso")
    assert va.referencia_humana({"h1": "incluir", "h2": "incerto"}, None) == ("incerto", "concordancia")
    assert va.referencia_humana({"h1": None, "h2": "excluir"}, None) == ("excluir", "unico")
    assert va.referencia_humana({"h1": "incluir", "h2": "excluir"}, None) == (None, "discordancia_sem_consenso")
    assert va.referencia_humana({"h1": None}, None) == (None, "nao_revisado")


# ---------------------------------------------------------------------------
# Projeto sintético: 120 registros; IA inclui RS0001–RS0050 e exclui RS0051–RS0120
# ---------------------------------------------------------------------------
N_POS, N_TOTAL = 50, 120


def _projeto_triado(raiz):
    criar_registros(raiz, n=N_TOTAL)
    linhas = []
    for k in range(1, N_TOTAL + 1):
        id_rs = f"RS{k:04d}"
        dec = "incluir" if k <= N_POS else "excluir"
        for rev in ("A", "B"):
            linhas.append(tl.nova_decisao(id_rs, "ta", "ta_v2", rev, "ia_subagente", dec, modelo="m",
                                          criterio_falhou="C2" if dec == "excluir" else None))
    tl.registrar_decisoes(raiz, linhas)


def _preencher(caminho, codigos_por_id, colunas=("decisao_h1", "decisao_h2")):
    from openpyxl import load_workbook

    wb = load_workbook(caminho)
    ws = wb["codificacao"]
    cab = [c.value for c in ws[1]]
    col_id = cab.index("id_rs") + 1
    for linha in range(2, ws.max_row + 1):
        id_rs = ws.cell(row=linha, column=col_id).value
        valores = codigos_por_id.get(id_rs)
        if valores is None:
            continue
        for nome, valor in zip(colunas, valores):
            ws.cell(row=linha, column=cab.index(nome) + 1).value = valor
    wb.save(caminho)


def _gabarito_ia(id_rs):
    return "incluir" if int(id_rs[2:]) <= N_POS else "excluir"


def test_amostrar_planilha_cega_desenho_e_idempotencia(projeto_vazio, capsys):
    from openpyxl import load_workbook

    raiz = projeto_vazio
    _projeto_triado(raiz)
    est = estado.carregar_estado(raiz)
    est["modo"]["autonomia"] = "autopiloto"
    estado.salvar_estado(raiz, est)
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "40", "--semente", "11",
                        "--enriquecer-incluidos", "30", "--elusao", "35")
    assert codigo == 0 and res["amostra_id"] == "amostra01" and res["pendencia"]
    desenho = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    assert desenho["populacao"] == {va.ESTRATO_POS: N_POS, va.ESTRATO_NEG: N_TOTAL - N_POS}
    por_estrato = desenho["amostra_por_estrato"]
    assert por_estrato[va.ESTRATO_POS] >= 30 and por_estrato[va.ESTRATO_NEG] >= 35
    assert sum(desenho["contagens_origem"].values()) == desenho["n_amostra"] == len(desenho["amostra"])
    assert desenho["contagens_origem"]["aleatoria"] == 40
    for item in desenho["amostra"]:
        assert item["estrato"] == (va.ESTRATO_POS if _gabarito_ia(item["id_rs"]) == "incluir" else va.ESTRATO_NEG)
    assert desenho["pesos_nominais"][va.ESTRATO_NEG] == pytest.approx((N_TOTAL - N_POS) / por_estrato[va.ESTRATO_NEG])

    wb = load_workbook(raiz / res["planilha"], read_only=True)
    assert wb.sheetnames == ["codificacao", "instrucoes", "_meta"]
    linhas = list(wb["codificacao"].iter_rows(values_only=True))
    cabecalho = list(linhas[0])
    assert cabecalho == va.COLUNAS_REGISTRO_PLANILHA + ["decisao_h1", "criterio_h1", "decisao_h2", "criterio_h2",
                                                        "decisao_consenso", "observacoes"]
    proibidas = {"decisao_ia", "decisao", "justificativa", "trecho", "criterio_falhou", "modelo", "estrato", "peso",
                 "origem", "revisor", "decidido_por", "divergente"}
    assert not proibidas & set(cabecalho)
    ids_planilha = [l[1] for l in linhas[1:]]
    assert len(ids_planilha) == desenho["n_amostra"] and ids_planilha != sorted(ids_planilha)
    tudo = json.dumps([list(l) for l in wb["_meta"].iter_rows(values_only=True)]
                      + [list(l) for l in wb["instrucoes"].iter_rows(values_only=True)]
                      + [list(l) for l in linhas], ensure_ascii=False, default=str)
    assert "ia_positivo" not in tudo and "ia_negativo" not in tudo and "enriquecimento" not in tudo
    assert all(v is None for l in linhas[1:] for v in l[9:])
    wb.close()

    sha = estado.sha256_arquivo(raiz / res["planilha"])
    n_ev = len(eventos(raiz, "artefato_versionado"))
    codigo, res2 = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "40", "--semente", "11",
                         "--enriquecer-incluidos", "30", "--elusao", "35")
    assert codigo == 0 and res2["reexecucao"] and res2["amostra_id"] == "amostra01"
    assert estado.sha256_arquivo(raiz / res["planilha"]) == sha, "planilha possivelmente codificada não é reescrita"
    assert len(eventos(raiz, "artefato_versionado")) == n_ev
    pend = estado.pendencias_abertas(estado.carregar_estado(raiz))
    assert [p["tipo"] for p in pend] == ["validacao_humana"]

    codigo, res3 = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "10", "--semente", "12")
    assert res3["amostra_id"] == "amostra02"
    d3 = json.loads((raiz / res3["desenho"]).read_text(encoding="utf-8"))
    assert not {a["id_rs"] for a in d3["amostra"]} & {a["id_rs"] for a in desenho["amostra"]}, "amostra nova"


def _amostra_padrao(capsys, raiz):
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "60", "--semente", "11",
                        "--enriquecer-incluidos", "40")
    assert codigo == 0
    desenho = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    return res, desenho


def test_calcular_atende_limiares_com_na(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    est = estado.carregar_estado(raiz)
    est["modo"]["autonomia"] = "autopiloto"
    estado.salvar_estado(raiz, est)
    res, desenho = _amostra_padrao(capsys, raiz)
    negativos = [a["id_rs"] for a in desenho["amostra"] if a["estrato"] == va.ESTRATO_NEG]
    codigos = {a["id_rs"]: (_gabarito_ia(a["id_rs"]), _gabarito_ia(a["id_rs"])) for a in desenho["amostra"]}
    for i in negativos[:3]:
        codigos[i] = (None, "nao_revisado")  # não revisado pelos dois
    codigos[negativos[3]] = (None, "excluir")  # só um codificador
    _preencher(raiz / res["planilha"], codigos)
    assert estado.pendencias_abertas(estado.carregar_estado(raiz))

    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 0 and out["ok"] and out["atende_limiares"] is True, out
    metricas = json.loads((raiz / out["metricas"]).read_text(encoding="utf-8"))
    assert metricas["n_na"] == 3 and metricas["n_codificados"] == desenho["n_amostra"] - 3
    assert metricas["origem_referencia"]["unico"] == 1
    assert metricas["ponderado"] is True
    ia = metricas["ia_final"]
    assert ia["sensibilidade"]["valor"] == 1.0 and ia["sensibilidade"]["ic_inferior"] >= 0.90
    assert ia["especificidade"]["valor"] == 1.0 and ia["kappa"] == pytest.approx(1.0)
    assert ia["pabak"] == pytest.approx(1.0)
    assert ia["wss_95"] == pytest.approx((N_TOTAL - N_POS) / N_TOTAL - 0.05)
    par = metricas["humanos"]["pares"]["h1~h2"]
    assert par["kappa_binario"] == pytest.approx(1.0) and par["n"] == desenho["n_amostra"] - 4
    assert set(metricas["por_revisor"]) == {"A", "B"}
    assert all(c["atende"] for c in metricas["criterios"].values())
    assert any("incluídos humanos" in a for a in metricas["avisos"])  # < 60 incluídos
    ev = eventos(raiz, "validacao_calculada")
    assert len(ev) == 1 and ev[0]["dados"]["atende_limiares"] is True
    assert ev[0]["dados"]["finalidade"] == "validacao" and out["finalidade"] == "validacao"  # padrão (v1.1)
    assert ev[0]["dados"]["sensibilidade_ic"][0] >= 0.90
    assert not estado.pendencias_abertas(estado.carregar_estado(raiz)), "codificação feita fecha a pendência"

    codigo, out2 = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 0 and out2["reexecucao"] and len(eventos(raiz, "validacao_calculada")) == 1


def test_calcular_falha_lista_falsos_negativos_e_exit_2(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    est = estado.carregar_estado(raiz)
    est["modo"]["autonomia"] = "autopiloto"
    estado.salvar_estado(raiz, est)
    res, desenho = _amostra_padrao(capsys, raiz)
    negativos = sorted(a["id_rs"] for a in desenho["amostra"] if a["estrato"] == va.ESTRATO_NEG)
    perdido, perdido2 = negativos[0], negativos[1]
    codigos = {a["id_rs"]: (_gabarito_ia(a["id_rs"]),) * 2 for a in desenho["amostra"]}
    codigos[perdido] = ("incluir", "excluir")  # discordância humana resolvida no consenso
    codigos[perdido2] = ("incluir", "incerto")  # os dois mandam adiante
    _preencher(raiz / res["planilha"], codigos)
    _preencher(raiz / res["planilha"], {perdido: ("incluir",)}, colunas=("decisao_consenso",))

    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 2 and out["atende_limiares"] is False and not out["ok"]
    assert out["ids_falsos_negativos"] == [perdido, perdido2]
    n_neg = desenho["amostra_por_estrato"][va.ESTRATO_NEG]
    fn_ponderado = 2 * (N_TOTAL - N_POS) / n_neg
    assert out["sensibilidade"] == pytest.approx(N_POS / (N_POS + fn_ponderado))
    assert out["sensibilidade"] < 0.95  # sem reponderar seria alto: todos os IA+ amostrados são acertos
    assert out["criterios"]["recall"]["atende"] is False
    assert out["elusao"] == pytest.approx(2 / n_neg)
    assert out["perdidos_estimados"] == pytest.approx(2 * (N_TOTAL - N_POS) / n_neg)
    assert out["kappa_humanos"] < 1.0
    fn_csv = (raiz / out["falsos_negativos"]).read_text(encoding="utf-8")
    assert perdido in fn_csv and "A=excluir [C2]" in fn_csv
    pend = estado.pendencias_abertas(estado.carregar_estado(raiz))
    assert [p["tipo"] for p in pend] == ["validacao_triagem_reprovada"]


def test_calcular_sem_dupla_codificacao_nao_atende(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "60", "--semente", "11",
                        "--enriquecer-incluidos", "40", "--codificadores", "1")
    desenho = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    _preencher(raiz / res["planilha"], {a["id_rs"]: (_gabarito_ia(a["id_rs"]),) for a in desenho["amostra"]},
               colunas=("decisao_h1",))
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 2 and out["criterios"]["recall"]["atende"] is True
    assert out["criterios"]["kappa_humanos"]["atende"] is None
    assert any("dupla codificação" in a for a in out["avisos"])


def test_calcular_codigo_invalido_e_overrides_ignorados(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    res, desenho = _amostra_padrao(capsys, raiz)
    negativo = next(a["id_rs"] for a in desenho["amostra"] if a["estrato"] == va.ESTRATO_NEG)
    _preencher(raiz / res["planilha"], {negativo: ("sim?", "excluir")})
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 1 and "sim?" in out["erro"]

    # override humano não pode melhorar a métrica da IA
    tl.registrar_decisoes(raiz, [tl.nova_decisao(negativo, "ta", "ta_v2", "humano_1", "humano", "incluir",
                                                 motivo_override="humano corrigiu")])
    codigos = {a["id_rs"]: (_gabarito_ia(a["id_rs"]),) * 2 for a in desenho["amostra"]}
    codigos[negativo] = ("incluir", "incluir")
    _preencher(raiz / res["planilha"], codigos)
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 2 and out["ids_falsos_negativos"] == [negativo]


def test_elusao_apos_rodada(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    res1, desenho1 = _amostra_padrao(capsys, raiz)
    codigo, res = rodar(capsys, raiz, "validar", "elusao", "--rodada", "ta_v2", "--n", "20", "--semente", "5")
    assert codigo == 0 and res["amostra_id"] == "elusao01" and res["n_amostra"] == 20
    desenho = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    assert {a["estrato"] for a in desenho["amostra"]} == {va.ESTRATO_NEG}
    assert not {a["id_rs"] for a in desenho["amostra"]} & {a["id_rs"] for a in desenho1["amostra"]}
    ids = sorted(a["id_rs"] for a in desenho["amostra"])
    codigos = {i: ("excluir",) for i in ids}
    codigos[ids[0]] = ("incluir",)
    _preencher(raiz / res["planilha"], codigos, colunas=("decisao_h1",))
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 0 and out["atende_limiares"] is None and out["tipo"] == "elusao"
    assert out["elusao"] == pytest.approx(1 / 20)
    assert out["perdidos_estimados"] == pytest.approx((N_TOTAL - N_POS) / 20)
    metricas = json.loads((raiz / out["metricas"]).read_text(encoding="utf-8"))
    inf, sup = metricas["elusao"]["perdidos_ic"]
    assert inf == pytest.approx(va.clopper_pearson(1, 20)[0] * (N_TOTAL - N_POS))
    assert sup == pytest.approx(va.clopper_pearson(1, 20)[1] * (N_TOTAL - N_POS))
    assert any("elusão com 20" in a for a in out["avisos"])
    ev = eventos(raiz, "validacao_calculada")[-1]
    assert ev["dados"]["tipo"] == "elusao" and "atende_limiares" not in ev["dados"]
    assert ev["dados"]["finalidade"] == "elusao"


def test_amostrar_sem_decisoes_erro(projeto_vazio, capsys):
    raiz = projeto_vazio
    criar_registros(raiz, n=5)
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v9", "--n", "3", "--semente", "1")
    assert codigo == 1 and "nenhuma decisão" in res["erro"]


def test_estabilidade_prepara_e_calcula(projeto_vazio, capsys):
    raiz = projeto_vazio
    criar_registros(raiz, n=20)
    gabarito = {f"RS{k:04d}": ("excluir" if k % 3 == 0 else "incluir") for k in range(1, 21)}
    triar_rodada(capsys, raiz, "ta_v2", {"A": gabarito, "B": gabarito}, tamanho=10)

    codigo, res = rodar(capsys, raiz, "validar", "estabilidade", "--rodada", "ta_v2", "--amostra", "0.1")
    assert codigo == 0 and res["acao"] == "preparar" and res["n_ids"] == 2
    assert [r["revisor"] for r in res["revisores"]] == ["A", "B"]
    desenho = json.loads((raiz / "02-triagem/validacao/ta_v2/estabilidade_desenho.json").read_text(encoding="utf-8"))
    ids = desenho["ids"]
    man = tl.ler_manifesto(raiz, "ta_v2_estab", "A")
    assert sorted(i for l in man["lotes"] for i in l["ids"]) == ids
    assert man["criterios_sha"] == tl.ler_manifesto(raiz, "ta_v2", "A")["criterios_sha"]

    responder_todos(raiz, "ta_v2_estab", "A", gabarito)
    virado = dict(gabarito)
    virado[ids[0]] = "incluir" if gabarito[ids[0]] == "excluir" else "excluir"
    responder_todos(raiz, "ta_v2_estab", "B", virado)
    for rev in ("A", "B"):
        codigo, _ = rodar(capsys, raiz, "triagem", "mesclar", "--rodada", "ta_v2_estab", "--revisor", rev)
        assert codigo == 0

    codigo, res = rodar(capsys, raiz, "validar", "estabilidade", "--rodada", "ta_v2", "--amostra", "0.1")
    assert codigo == 0 and res["acao"] == "calcular"
    assert res["revisores"]["A"]["concordancia_3cat"] == 1.0 and res["revisores"]["A"]["n"] == 2
    assert res["revisores"]["B"]["concordancia_binaria"] == 0.5
    assert res["revisores"]["B"]["mudancas_de_seguimento"] == 1
    arquivo = json.loads((raiz / res["arquivo"]).read_text(encoding="utf-8"))
    assert arquivo["revisores"]["B"]["mudancas"][0]["id_rs"] == ids[0]
    assert eventos(raiz, "validacao_calculada")[-1]["dados"]["tipo"] == "estabilidade"
    assert eventos(raiz, "validacao_calculada")[-1]["dados"]["finalidade"] == "estabilidade"
    n_ev = len(eventos(raiz, "validacao_calculada"))
    codigo, res = rodar(capsys, raiz, "validar", "estabilidade", "--rodada", "ta_v2", "--amostra", "0.1")
    assert res["reexecucao"] and len(eventos(raiz, "validacao_calculada")) == n_ev

    codigo, res = rodar(capsys, raiz, "validar", "estabilidade", "--rodada", "ta_v2", "--amostra", "0.2")
    assert codigo == 1 and "outros parâmetros" in res["erro"]


# ---------------------------------------------------------------------------
# Finalidade, --excluir-ids e calibração sem IA (v1.1)
# ---------------------------------------------------------------------------
def _autopiloto(raiz):
    est = estado.carregar_estado(raiz)
    est["modo"]["autonomia"] = "autopiloto"
    estado.salvar_estado(raiz, est)


def test_finalidade_desenvolvimento_reprovado_sem_pendencia_e_conflito(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    _autopiloto(raiz)
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "60", "--semente", "11",
                        "--enriquecer-incluidos", "40", "--finalidade", "desenvolvimento")
    assert codigo == 0 and res["finalidade"] == "desenvolvimento"
    desenho = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    assert desenho["finalidade"] == "desenvolvimento" and desenho["parametros"]["finalidade"] == "desenvolvimento"
    negativos = sorted(a["id_rs"] for a in desenho["amostra"] if a["estrato"] == va.ESTRATO_NEG)
    codigos = {a["id_rs"]: (_gabarito_ia(a["id_rs"]),) * 2 for a in desenho["amostra"]}
    codigos[negativos[0]] = ("incluir", "incluir")  # falso negativo: reprova o recall
    _preencher(raiz / res["planilha"], codigos)

    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"], "--finalidade", "validacao")
    assert codigo == 1 and "desenvolvimento" in out["erro"], "amostra de desenvolvimento não vira validação"

    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 2 and out["finalidade"] == "desenvolvimento" and out["atende_limiares"] is False
    assert "desenvolvimento" in out["proxima_acao"]
    assert eventos(raiz, "validacao_calculada")[-1]["dados"]["finalidade"] == "desenvolvimento"
    tipos = [p["tipo"] for p in estado.pendencias_abertas(estado.carregar_estado(raiz))]
    assert "validacao_triagem_reprovada" not in tipos, "desenvolvimento reprovado não abre pendência"

    # finalidade incompatível com o tipo de desenho
    codigo, out = rodar(capsys, raiz, "validar", "elusao", "--rodada", "ta_v2", "--n", "5", "--semente", "1",
                        "--finalidade", "validacao")
    assert codigo == 1 and "não combina" in out["erro"]
    codigo, out = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "5", "--semente", "1",
                        "--sem-ia")
    assert codigo == 1 and "--sem-ia" in out["erro"]


def test_amostrar_excluir_ids_tira_do_quadro_e_da_populacao(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    usados = [f"RS{k:04d}" for k in (1, 2, 3, 60, 61)]  # 3 incluídos e 2 excluídos pela IA
    (raiz / "usados.csv").write_text("id_rs\n" + "\n".join(usados) + "\n", encoding="utf-8")
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", str(N_TOTAL), "--semente", "4",
                        "--excluir-ids", "usados.csv")
    assert codigo == 0 and res["n_amostra"] == N_TOTAL - len(usados)
    desenho = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    assert not {a["id_rs"] for a in desenho["amostra"]} & set(usados)
    assert desenho["populacao"] == {va.ESTRATO_POS: N_POS - 3, va.ESTRATO_NEG: N_TOTAL - N_POS - 2}
    assert desenho["populacao_rodada"] == {va.ESTRATO_POS: N_POS, va.ESTRATO_NEG: N_TOTAL - N_POS}
    assert desenho["n_excluidos_quadro"] == 5 and "excluir_ids_sha" in desenho["parametros"]
    codigo, res2 = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", str(N_TOTAL), "--semente", "4",
                         "--excluir-ids", "usados.csv")
    assert res2["reexecucao"] and res2["amostra_id"] == res["amostra_id"]


def test_calibracao_humana_sem_decisoes_de_ia(projeto_vazio, capsys):
    from openpyxl import load_workbook

    raiz = projeto_vazio
    criar_registros(raiz, n=40)
    _autopiloto(raiz)
    criterios = raiz / "02-triagem/prompts/ta_v1.md"
    criterios.write_text("- C1 população\n- C2 desenho\n", encoding="utf-8")
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "calib_v1", "--finalidade", "calibracao",
                        "--n", "30", "--semente", "3", "--criterios", "02-triagem/prompts/ta_v1.md")
    assert codigo == 0 and res["sem_ia"] is True and res["n_amostra"] == 30 and res["finalidade"] == "calibracao"
    desenho = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    assert {a["estrato"] for a in desenho["amostra"]} == {va.ESTRATO_SEM_IA}
    assert desenho["criterios_sha"] == estado.sha256_arquivo(criterios)
    wb = load_workbook(raiz / res["planilha"], read_only=True)
    cab = list(next(wb["codificacao"].iter_rows(values_only=True)))
    assert "decisao_h1" in cab and "decisao_h2" in cab
    wb.close()

    ids = sorted(a["id_rs"] for a in desenho["amostra"])
    codigos = {i: ("incluir", "incluir") if k < 12 else ("excluir", "excluir") for k, i in enumerate(ids)}
    codigos[ids[20]] = ("excluir", "incluir")  # uma discordância, sem consenso
    _preencher(raiz / res["planilha"], codigos)
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 0 and out["atende_limiares"] is True and out["sem_ia"] is True, out
    assert out["sensibilidade"] is None and out["kappa_ia"] is None
    assert out["concordancia_humanos"] == pytest.approx(29 / 30)
    assert set(out["criterios"]) == {"kappa_humanos", "concordancia_humanos", "tamanho_calibracao"}
    assert out["criterios"]["tamanho_calibracao"]["atende"] is True  # < 100 registros, mas >= 10 incluídos
    ev = eventos(raiz, "validacao_calculada")[-1]
    assert ev["dados"]["finalidade"] == "calibracao" and ev["dados"]["sem_ia"] is True

    # calibração reprovada: pendência própria, não a da validação
    codigo, res2 = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "calib_v1", "--finalidade", "calibracao",
                         "--n", "8", "--semente", "9", "--sem-ia")
    desenho2 = json.loads((raiz / res2["desenho"]).read_text(encoding="utf-8"))
    assert not {a["id_rs"] for a in desenho2["amostra"]} & set(ids)
    ids2 = sorted(a["id_rs"] for a in desenho2["amostra"])
    _preencher(raiz / res2["planilha"], {i: ("incluir", "excluir") if k % 2 else ("excluir", "excluir")
                                         for k, i in enumerate(ids2)})
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res2["planilha"])
    assert codigo == 2 and out["atende_limiares"] is False and "critérios" in out["proxima_acao"]
    tipos = [p["tipo"] for p in estado.pendencias_abertas(estado.carregar_estado(raiz))]
    assert "calibracao_reprovada" in tipos and "validacao_triagem_reprovada" not in tipos


def test_calibracao_com_ia_mede_so_humanos(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "40", "--semente", "2",
                        "--finalidade", "calibracao")
    assert codigo == 0 and res["sem_ia"] is False
    desenho = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    # humanos incluem tudo: a IA erra todos os negativos, mas a calibração mede os critérios pela dupla humana
    _preencher(raiz / res["planilha"], {a["id_rs"]: ("incluir", "incluir") if k % 4 else ("excluir", "excluir")
                                        for k, a in enumerate(desenho["amostra"])})
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert out["finalidade"] == "calibracao" and "recall" not in out["criterios"]
    assert out["sensibilidade"] is not None and out["sensibilidade"] < 0.95
    assert codigo == 0 and out["atende_limiares"] is True


def test_estabilidade_finalidade_incompativel(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    codigo, res = rodar(capsys, raiz, "validar", "estabilidade", "--rodada", "ta_v2", "--finalidade", "validacao")
    assert codigo == 1 and "não combina" in res["erro"]
    codigo, res = rodar(capsys, raiz, "validar", "estabilidade", "--rodada", "ta_v2_estab")
    assert codigo == 1 and "reexecução" in res["erro"]


def test_regressao_validar_ignora_clusters_de_busca_substituida(projeto_vazio, capsys):
    """amostrar, elusao, calibração sem IA e estabilidade: busca_inativa fora do quadro e da população."""
    raiz = projeto_vazio
    _projeto_triado(raiz)
    inativos = {f"RS{k:04d}" for k in (1, 2, 60, 61, 62)}  # 2 incluídos e 3 excluídos pela IA
    marcar_inativos(raiz, inativos)
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", str(N_TOTAL),
                        "--semente", "4")
    assert codigo == 0 and res["n_amostra"] == N_TOTAL - len(inativos) and res["n_inativos_ignorados"] == 5
    desenho = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    assert not {a["id_rs"] for a in desenho["amostra"]} & inativos
    assert desenho["populacao"] == {va.ESTRATO_POS: N_POS - 2, va.ESTRATO_NEG: N_TOTAL - N_POS - 3}
    assert desenho["n_inativos_ignorados"] == 5
    assert eventos(raiz, "artefato_versionado")[-1]["dados"]["n_inativos_ignorados"] == 5

    codigo, res = rodar(capsys, raiz, "validar", "elusao", "--rodada", "ta_v3", "--n", "500", "--semente", "5")
    assert codigo == 1  # rodada sem decisões
    tl.registrar_decisoes(raiz, [tl.nova_decisao(f"RS{k:04d}", "ta", "ta_v3", rev, "ia_subagente", "excluir",
                                                 modelo="m", criterio_falhou="C2")
                                 for k in range(1, N_TOTAL + 1) for rev in ("A", "B")])
    codigo, res = rodar(capsys, raiz, "validar", "elusao", "--rodada", "ta_v3", "--n", "500", "--semente", "5")
    assert codigo == 0 and res["n_amostra"] == N_TOTAL - len(inativos)
    d_el = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    assert not {a["id_rs"] for a in d_el["amostra"]} & inativos

    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "calib_v1", "--finalidade", "calibracao",
                        "--n", "500", "--semente", "3", "--sem-ia")
    d_cal = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    assert codigo == 0 and res["sem_ia"] and not {a["id_rs"] for a in d_cal["amostra"]} & inativos
    assert d_cal["populacao"] == {va.ESTRATO_SEM_IA: N_TOTAL - len(inativos)}

    codigo, res = rodar(capsys, raiz, "validar", "estabilidade", "--rodada", "ta_v2", "--amostra", "1.0",
                        "--semente", "7", "--acao", "preparar")
    d_est = json.loads((raiz / "02-triagem/validacao/ta_v2/estabilidade_desenho.json").read_text(encoding="utf-8"))
    assert not set(d_est["ids"]) & inativos and d_est["n_inativos_ignorados"] == 5


# ---------------------------------------------------------------------------
# v1.2: reprovação só pela largura do IC, desenhos antigos com busca substituída, planilha CSV do Excel
# ---------------------------------------------------------------------------
def _amostra_pequena_codificada(capsys, raiz, n, enriquecer, semente):
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", str(n), "--semente",
                        str(semente), "--enriquecer-incluidos", str(enriquecer))
    assert codigo == 0, res
    desenho = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    _preencher(raiz / res["planilha"], {a["id_rs"]: (_gabarito_ia(a["id_rs"]),) * 2 for a in desenho["amostra"]})
    return res, desenho


def test_incluidos_minimos_para_ic():
    assert va.incluidos_minimos_para_ic() == 36
    assert va.clopper_pearson(36, 36)[0] >= 0.90 > va.clopper_pearson(35, 35)[0]


def test_regressao_largura_ic_sugere_ampliar_amostra_e_nao_revisar_criterios(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    est = estado.carregar_estado(raiz)
    est["modo"]["autonomia"] = "autopiloto"
    estado.salvar_estado(raiz, est)
    res, desenho = _amostra_pequena_codificada(capsys, raiz, n=12, enriquecer=10, semente=11)
    assert desenho["amostra_por_estrato"][va.ESTRATO_NEG] > 0 and desenho["amostra_por_estrato"][va.ESTRATO_POS] == 10
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 2 and out["n_falsos_negativos"] == 0 and out["motivo_reprovacao"] == "largura_ic", out
    assert [k for k, c in out["criterios"].items() if not c["atende"]] == ["recall_ic_inferior"]
    assert out["incluidos_humanos_necessarios"] == 36 and out["incluidos_ia_nao_sorteados"] == N_POS - 10
    assert "--enriquecer-incluidos 40" in out["proxima_acao"] and "validar elusao" in out["proxima_acao"]
    assert "vN+1" not in out["proxima_acao"] and "não revise os critérios" in out["proxima_acao"]
    assert "Apêndice" not in json.dumps(out, ensure_ascii=False)
    pend = estado.pendencias_abertas(estado.carregar_estado(raiz))
    assert [p["tipo"] for p in pend] == ["validacao_triagem_reprovada"] and "largura do IC" in pend[0]["descricao"]
    assert eventos(raiz, "validacao_calculada")[-1]["dados"]["motivo_reprovacao"] == "largura_ic"


def test_regressao_largura_ic_inalcancavel_aponta_remedio_5(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    res, desenho = _amostra_pequena_codificada(capsys, raiz, n=25, enriquecer=20, semente=11)
    assert desenho["amostra_por_estrato"][va.ESTRATO_POS] == 20
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 2 and out["motivo_reprovacao"] == "largura_ic" and out["incluidos_ia_nao_sorteados"] == 30
    assert "inalcançável" in out["proxima_acao"] and "remédio 5" in out["proxima_acao"]
    assert f"--por {esquema.PAPEL_HUMANO_PADRAO} --forcar" in out["proxima_acao"]
    assert "references/ia-validacao.md" in out["proxima_acao"]


def test_regressao_falha_por_falsos_negativos_mantem_remedios(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    res, desenho = _amostra_padrao(capsys, raiz)
    codigos = {a["id_rs"]: (_gabarito_ia(a["id_rs"]),) * 2 for a in desenho["amostra"]}
    perdido = sorted(a["id_rs"] for a in desenho["amostra"] if a["estrato"] == va.ESTRATO_NEG)[0]
    codigos[perdido] = ("incluir", "incluir")
    _preencher(raiz / res["planilha"], codigos)
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 2 and out["motivo_reprovacao"] == "desempenho_ia"
    assert "vN+1" in out["proxima_acao"] and "seção 4 E" in out["proxima_acao"]
    assert "incluidos_humanos_necessarios" not in out


def test_regressao_desenho_antigo_com_busca_substituida_avisa(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    res, desenho = _amostra_padrao(capsys, raiz)
    sorteados = sorted(a["id_rs"] for a in desenho["amostra"])[:3]
    marcar_inativos(raiz, set(sorteados))  # `importar --substituir` + dedup depois do sorteio

    codigo, res2 = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "60", "--semente", "11",
                         "--enriquecer-incluidos", "40")
    assert codigo == 0 and res2["reexecucao"] is True and res2["n_inativos_no_desenho"] == 3
    assert any("--semente" in a and esquema.FLAG_BUSCA_INATIVA in a for a in res2["avisos"])
    _preencher(raiz / res["planilha"], {a["id_rs"]: (_gabarito_ia(a["id_rs"]),) * 2 for a in desenho["amostra"]})
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert any("substituição de busca" in a for a in out["avisos"])

    codigo, _ = rodar(capsys, raiz, "validar", "estabilidade", "--rodada", "ta_v2", "--amostra", "0.1",
                      "--semente", "7", "--acao", "preparar")
    d_est = json.loads((raiz / "02-triagem/validacao/ta_v2/estabilidade_desenho.json").read_text(encoding="utf-8"))
    marcar_inativos(raiz, set(sorteados) | {d_est["ids"][0]})
    codigo, res3 = rodar(capsys, raiz, "validar", "estabilidade", "--rodada", "ta_v2", "--amostra", "0.1",
                         "--semente", "7", "--acao", "preparar")
    assert codigo == 0 and any(esquema.ARQ_DESENHO_ESTABILIDADE in a for a in res3["avisos"])


def test_regressao_planilha_csv_do_excel_e_cabecalho_sem_decisao(projeto_vazio, capsys):
    import csv

    raiz = projeto_vazio
    _projeto_triado(raiz)
    res, desenho = _amostra_padrao(capsys, raiz)
    csv_path = raiz / "02-triagem/validacao/ta_v2/amostra01_codificada.csv"
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";", lineterminator="\r\n")
        w.writerow(["ordem", "ID_RS", "titulo", "decisao_h1", "decisao_h2", "decisao_consenso"])
        for k, a in enumerate(desenho["amostra"], start=1):
            w.writerow([k, a["id_rs"].lower(), "Título; com ponto e vírgula", _gabarito_ia(a["id_rs"]).capitalize(),
                        _gabarito_ia(a["id_rs"]), ""])
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", str(csv_path), "--desenho", res["desenho"])
    assert codigo == 0 and out["atende_limiares"] is True and out["n_codificados"] == desenho["n_amostra"], out

    csv_path.write_text("id_rs;h1;h2\nRS0001;incluir;incluir\n", encoding="utf-8")
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", str(csv_path), "--desenho", res["desenho"])
    assert codigo == 1 and "decisao_h" in out["erro"]


# ---------------------------------------------------------------------------
# v1.3: atalho da variante rápida gravado por `validar calcular` e segunda leitura dos excluídos pela IA
# ---------------------------------------------------------------------------
def _variante_rapida(raiz):
    est = estado.carregar_estado(raiz)
    est["projeto"]["variante"] = "rapida"
    estado.salvar_estado(raiz, est)


def _csv_segunda_leitura(caminho, decisoes):
    """Planilha do Excel pt-BR: cp1252, ponto e vírgula, cabeçalho com caixa diferente."""
    linhas = ["ID_RS;decisao_humana;observação"] + [f"{i};{d};relido à mão" for i, d in decisoes.items()]
    caminho.write_bytes(("\r\n".join(linhas) + "\r\n").encode("cp1252"))
    return str(caminho)


def test_atalho_rapida_calcular_grava_campos_e_segunda_leitura_completa(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    _variante_rapida(raiz)
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "10", "--semente", "4",
                        "--atalho-rapida", "--codificadores", "1")
    assert codigo == 1 and "dupla" in res["erro"]
    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "10", "--semente", "4",
                        "--atalho-rapida")
    assert codigo == 0 and any(">= 20%" in a and "24" in a for a in res["avisos"])
    (raiz / res["planilha"]).unlink()
    (raiz / res["desenho"]).unlink()

    codigo, res = rodar(capsys, raiz, "validar", "amostrar", "--rodada", "ta_v2", "--n", "30", "--semente", "5",
                        "--atalho-rapida")
    assert codigo == 0 and not any(">= 20%" in a for a in res["avisos"])
    desenho = json.loads((raiz / res["desenho"]).read_text(encoding="utf-8"))
    assert desenho["atalho_rapida"] is True and desenho["parametros"]["atalho_rapida"] is True
    _preencher(raiz / res["planilha"], {a["id_rs"]: (_gabarito_ia(a["id_rs"]),) * 2 for a in desenho["amostra"]})
    na_amostra = {a["id_rs"] for a in desenho["amostra"]}
    excluidos = [f"RS{k:04d}" for k in range(N_POS + 1, N_TOTAL + 1)]

    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 2 and out["ok"] is False and out["atalho_rapida"] is True
    assert (out["n_dupla_humana"], out["n_populacao"], out["fracao_dupla_humana"]) == (30, N_TOTAL, 0.25)
    assert out["kappa_humanos"] == pytest.approx(1.0) and out["n_excluidos_ia"] == N_TOTAL - N_POS
    assert out["n_excluidos_relidos"] == len(na_amostra & set(excluidos)) and out["segunda_leitura_excluidos"] is False
    assert "validar segunda-leitura" in out["proxima_acao"] and "vN+1" not in out["proxima_acao"]
    ev = eventos(raiz, "validacao_calculada")[-1]
    assert set(esquema.CAMPOS_ATALHO_RAPIDA) <= set(ev["dados"]) and "motivo_reprovacao" in ev["dados"]
    assert ev["dados"]["atalho_rapida"] is True and ev["dados"]["finalidade"] == "validacao"

    faltam = [i for i in excluidos if i not in na_amostra]
    parte1 = {i: "excluir" for i in faltam[:20]}
    parte2 = {i: "excluir" for i in faltam[20:]}
    resgatado = faltam[-1]
    parte2[resgatado] = "Incluir"
    codigo, sl = rodar(capsys, raiz, "validar", "segunda-leitura", "--rodada", "ta_v2", "--planilha",
                       _csv_segunda_leitura(raiz / "02-triagem/validacao/ta_v2/segunda_parte1.csv", parte1))
    assert codigo == 0 and sl["segunda_leitura_excluidos"] is False and sl["n_faltam"] == len(faltam) - 20, sl
    assert sl["n_relidos_na_amostra_do_atalho"] == len(na_amostra & set(excluidos))
    assert any("cp1252" in a for a in sl["avisos"]) and "validar calcular --planilha" in sl["proxima_acao"]
    codigo, sl = rodar(capsys, raiz, "validar", "segunda-leitura", "--rodada", "ta_v2", "--planilha",
                       _csv_segunda_leitura(raiz / "02-triagem/validacao/ta_v2/segunda_parte2.csv", parte2))
    assert codigo == 0 and sl["segunda_leitura_excluidos"] is True and sl["n_resgatados"] == 1, sl
    ev_sl = eventos(raiz, "validacao_calculada")[-1]
    assert ev_sl["dados"]["tipo"] == esquema.TIPO_SEGUNDA_LEITURA and ev_sl["dados"]["finalidade"] == "elusao"
    assert esquema.CRITERIO_ATALHO_RAPIDA not in ev_sl["dados"] and "atende_limiares" not in ev_sl["dados"]
    fila = tl.ler_tabela_humana(raiz / sl["resgatados"], ["id_rs"])[1]
    assert [(l["id_rs"], l["decisao_humana"]) for l in fila] == [(resgatado, "incluir")]
    codigo, rep = rodar(capsys, raiz, "validar", "segunda-leitura", "--rodada", "ta_v2", "--planilha",
                        str(raiz / "02-triagem/validacao/ta_v2/segunda_parte2.csv"))
    assert codigo == 0 and rep["reexecucao"] is True and eventos(raiz, "validacao_calculada")[-1]["seq"] == ev_sl["seq"]
    codigo, ov = rodar(capsys, raiz, "triagem", "override", "--fila", sl["resgatados"], "--rodada", "ta_v2")
    assert codigo == 0 and ov["registrados"] == 1

    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"])
    assert codigo == 0 and out["ok"] is True and out["reexecucao"] is False, out
    assert out["segunda_leitura_excluidos"] is True and out["n_excluidos_relidos"] == out["n_excluidos_ia"]
    assert "G4" in out["proxima_acao"]
    ev = eventos(raiz, "validacao_calculada")[-1]
    assert ev["dados"]["segunda_leitura_excluidos"] is True and ev["dados"]["atalho_rapida"] is True
    assert not [p for p in estado.pendencias_abertas(estado.carregar_estado(raiz))
                if p["tipo"] == "validacao_triagem_reprovada"]
    try:
        from rslib import projeto
    except ImportError:  # pragma: no cover
        return
    if hasattr(projeto, "checar_atalho_rapida"):
        assert projeto.checar_atalho_rapida(ev) == []
    if hasattr(projeto, "validacao_do_portao"):
        decisora, _, _ = projeto.validacao_do_portao(estado.carregar_estado(raiz), estado.ler_log(raiz))
        assert decisora["seq"] == ev["seq"], "a segunda leitura (finalidade elusao) nunca decide o G4"


def test_atalho_rapida_ignorado_fora_da_variante_rapida(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    res, desenho = _amostra_pequena_codificada(capsys, raiz, n=30, enriquecer=0, semente=5)
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"], "--atalho-rapida")
    assert "atalho_rapida" not in out and any("variante rápida" in a for a in out["avisos"])
    assert esquema.CRITERIO_ATALHO_RAPIDA not in eventos(raiz, "validacao_calculada")[-1]["dados"]
    # na variante rápida, --atalho-rapida no calcular marca um desenho sorteado sem a flag
    _variante_rapida(raiz)
    codigo, out = rodar(capsys, raiz, "validar", "calcular", "--planilha", res["planilha"], "--atalho-rapida")
    assert out["atalho_rapida"] is True and out["fracao_dupla_humana"] == 0.25


def test_segunda_leitura_exige_coluna_de_decisao_e_rodada(projeto_vazio, capsys):
    raiz = projeto_vazio
    _projeto_triado(raiz)
    arq = raiz / "sl.csv"
    arq.write_text("id_rs;obs\nRS0060;x\n", encoding="utf-8")
    codigo, res = rodar(capsys, raiz, "validar", "segunda-leitura", "--planilha", str(arq))
    assert codigo == 1 and "--rodada" in res["erro"]
    codigo, res = rodar(capsys, raiz, "validar", "segunda-leitura", "--rodada", "ta_v2", "--planilha", str(arq))
    assert codigo == 1 and "decisao_humana" in res["erro"]
