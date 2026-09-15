"""Testes de `rs.py filtrar`: fronteira de palavra, etiquetar por padrão, regras de exclusão e âncoras."""

import argparse
import csv
import json
import shutil

import pytest

from conftest import FIXTURES, SKILL, ler_jsonl

DIR = FIXTURES / "filtrar"
DICIONARIO = SKILL / "assets" / "dicionarios" / "metodo_pt_en_es.csv"


def rodar(raiz, *argv):
    from rslib import filtrar
    parser = argparse.ArgumentParser()
    parser.add_argument("--dir")
    sub = parser.add_subparsers()
    filtrar.registrar(sub)
    args = parser.parse_args(["--dir", str(raiz), "filtrar", *argv])
    return args.func(args)


def ler(caminho):
    with open(caminho, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def resultados(raiz):
    from rslib import esquema
    return {(l["id_rs"], l["filtro"]): l for l in ler(raiz / esquema.ARQ_FILTRO_FORMAL)}


def escrever_config(raiz, filtros, **topo):
    caminho = raiz / "01-busca" / "filtros_teste.json"
    caminho.write_text(json.dumps({"filtros": filtros, **topo}, ensure_ascii=False), encoding="utf-8")
    return str(caminho)


def ultima_linha_json(capsys):
    return json.loads(capsys.readouterr().out.strip().splitlines()[-1])


@pytest.fixture
def projeto(projeto_vazio):
    from rslib import esquema
    shutil.copy(DIR / "registros_unicos.csv", projeto_vazio / esquema.ARQ_UNICOS)
    return projeto_vazio


# ---------------------------------------------------------------------------
# Casamento de termos
# ---------------------------------------------------------------------------
def termo(t, sensivel=False, grupo="g"):
    from rslib.filtrar import Termo
    return Termo(t, "xx", grupo, sensivel)


def casa(t, texto, sensivel=False):
    from rslib.filtrar import casar
    return bool(casar([termo(t, sensivel)], [texto]))


def test_fronteira_de_palavra_e_siglas():
    # SEM (structural equation modeling) x "sem" (preposição)
    assert casa("SEM", "We estimate a SEM with survey data", sensivel=True)
    assert not casa("SEM", "Programa sem avaliação de impacto", sensivel=True)
    assert not casa("SEM", "AVALIAÇÃO SEM GRUPO DE COMPARAÇÃO EM MUNICÍPIOS", sensivel=True)  # campo todo em caixa alta
    assert casa("SEM", "Modelo SEM; análise", sensivel=True)
    # ols x bolsa
    assert not casa("OLS", "Bolsa Família e frequência escolar", sensivel=True)
    assert not casa("ols", "Bolsa Família e frequência escolar")
    assert casa("OLS", "estimated by OLS.", sensivel=True)
    assert not casa("OLS", "the ols estimator", sensivel=True)
    # refis x Refiscalizar
    assert not casa("refis", "Refiscalizar contribuintes")
    assert casa("refis", "adesão ao REFIS em 2000")
    # PCA/RCT dentro de palavras
    assert not casa("PCA", "UPCASE letters", sensivel=True)
    assert casa("PCA", "a PCA of fiscal indicators", sensivel=True)
    assert not casa("RCT", "RCTX trial", sensivel=True)
    assert casa("RCT", "an RCT in schools", sensivel=True)


def test_truncamento_hifen_e_acentos():
    assert casa("quasi-experiment*", "a quasi experimental design")
    assert casa("quasi-experiment*", "Quasi–Experiments in policy")  # travessão tipográfico
    assert not casa("quasi-experiment*", "quasiexperimental")
    assert casa("difference* in difference*", "Differences-in-Differences estimates")
    assert casa("análise de conteúdo", "uma ANALISE DE CONTEUDO dos discursos")
    assert casa("análise de conteúdo", "análise  de\nconteúdo")
    assert casa("entrevista*", "Entrevistas semiestruturadas")
    assert not casa("entrevista*", "preentrevista")
    assert casa("2SLS", "estimated with 2SLS and", sensivel=True)


def test_dicionario_de_metodos_bem_formado():
    from rslib.filtrar import COLUNAS_DICIONARIO, carregar_dicionario, casar
    with open(DICIONARIO, encoding="utf-8", newline="") as f:
        leitor = csv.reader(f)
        assert next(leitor) == COLUNAS_DICIONARIO
        linhas = [l for l in leitor if l]
    chaves = [(t, i, g) for t, i, g, _ in linhas]
    assert len(chaves) == len(set(chaves)), "termos duplicados"
    assert {i for _, i, _, _ in linhas} == {"pt", "en", "es"}
    sensiveis = {t for t, _, _, s in linhas if s == "1"}
    assert {"SEM", "OLS", "PCA", "RCT"} <= sensiveis
    assert all(s in {"0", "1"} for *_, s in linhas)
    termos = carregar_dicionario(DICIONARIO)  # compila tudo
    assert {t.grupo for t in termos} >= {"metodos", "causalidade", "experimental", "quantitativo", "qualitativo", "misto"}
    minusculos = {t.termo.lower() for t in termos if not t.sensivel}
    assert not {"sem", "ols", "pca", "rct", "causa", "causas", "cause", "cluster", "matching"} & minusculos
    # falsos acertos do funil de exemplo do REFIS
    for texto in ["Programa sem avaliação", "Bolsa Família e bolsas de estudo", "Por causa da crise fiscal",
                  "Refiscalizar contribuintes", "AVALIAÇÃO SEM GRUPO DE COMPARAÇÃO"]:
        assert casar(termos, [texto]) == [], texto
    # acertos esperados em PT/EN/ES
    for texto, grupo in [("Estimamos por MQO com dados em painel", "quantitativo"),
                         ("we run an RCT in 40 schools", "experimental"),
                         ("entrevistas em profundidade com gestores", "qualitativo"),
                         ("un diseño cuasi-experimental", "experimental"),
                         ("métodos mistos com triangulação", "misto")]:
        assert grupo in {g for g, _ in casar(termos, [texto])}, texto


# ---------------------------------------------------------------------------
# Comando: etiquetar por padrão
# ---------------------------------------------------------------------------
def test_etiquetar_e_padrao_e_ninguem_sai(projeto, capsys):
    from rslib import esquema, estado
    assert rodar(projeto, "--config", str(DIR / "filtros_etiquetar.json")) == 0
    r = resultados(projeto)
    with open(projeto / esquema.ARQ_FILTRO_FORMAL, encoding="utf-8") as f:
        assert next(csv.reader(f)) == esquema.COLUNAS_FILTRO_FORMAL
    assert len(r) == 13 * 4 and not any(l["resultado"] == "exclui" for l in r.values())
    assert r[("RS0005", "ano")]["resultado"] == "etiqueta"
    assert r[("RS0006", "ano")]["resultado"] == "sem_dado"      # campo ausente mantém
    assert r[("RS0006", "tipo")]["resultado"] == "sem_dado"
    assert r[("RS0006", "idioma")]["resultado"] == "sem_dado"
    assert r[("RS0007", "tipo")]["resultado"] == "etiqueta"
    assert r[("RS0008", "idioma")]["resultado"] == "etiqueta"
    assert r[("RS0001", "metodo")]["resultado"] == "passa" and "SEM" in r[("RS0001", "metodo")]["detalhe"]
    assert r[("RS0002", "metodo")]["resultado"] == "etiqueta"   # "sem" não é SEM, "bolsa" não é OLS
    assert r[("RS0011", "metodo")]["resultado"] == "etiqueta"   # caixa alta não vira sigla
    assert r[("RS0013", "metodo")]["resultado"] == "etiqueta"   # "por causa" não é causalidade
    assert r[("RS0005", "metodo")]["resultado"] == "sem_dado" and "sem_resumo" in r[("RS0005", "metodo")]["detalhe"]
    assert r[("RS0010", "metodo")]["resultado"] == "sem_dado" and "resumo_truncado" in r[("RS0010", "metodo")]["detalhe"]
    assert r[("RS0009", "metodo")]["resultado"] == "passa"      # truncado, mas o título casa

    resumo = ultima_linha_json(capsys)
    assert resumo["comando"] == "filtrar" and resumo["ok"] and resumo["n_saida"] == 13 and resumo["n_excluidos"] == 0
    contagens = json.loads((projeto / "02-triagem" / "filtro_formal_contagens.json").read_text(encoding="utf-8"))
    metodo = contagens["filtros"][3]
    assert (metodo["passa"], metodo["etiqueta"], metodo["exclui"], metodo["sem_dado"]) == (4, 7, 0, 2)
    assert estado.carregar_estado(projeto)["versoes_ativas"]["filtros"] == "filtros_v1"
    ev = [e for e in ler_jsonl(projeto / esquema.ARQ_LOG) if e["evento"] == "filtro_formal"][-1]
    assert ev["dados"]["n_saida"] == 13 and ev["dados"]["dicionarios"]
    assert {a["caminho"] for a in ev["artefatos"]} >= {esquema.ARQ_FILTRO_FORMAL, esquema.ARQ_UNICOS}


def test_reexecucao_idempotente(projeto):
    from rslib import esquema
    cfg = str(DIR / "filtros_etiquetar.json")
    assert rodar(projeto, "--config", cfg) == 0
    antes = (projeto / esquema.ARQ_FILTRO_FORMAL).read_bytes()
    assert rodar(projeto, "--config", cfg) == 0
    assert (projeto / esquema.ARQ_FILTRO_FORMAL).read_bytes() == antes
    eventos = [e for e in ler_jsonl(projeto / esquema.ARQ_LOG) if e["evento"] == "filtro_formal"]
    assert [e["dados"]["sem_mudancas"] for e in eventos] == [False, True]


def test_dicionario_personalizado_por_caminho(projeto):
    cfg = escrever_config(projeto, [{"nome": "programa", "tipo": "dicionario",
                                     "dicionario": str(DIR / "dicionario_substantivo.csv"),
                                     "campos": ["titulo", "resumo"]}])
    assert rodar(projeto, "--config", cfg) == 0
    r = resultados(projeto)
    assert r[("RS0003", "programa")]["resultado"] == "etiqueta"   # Refiscalizar não é REFIS
    assert r[("RS0004", "programa")]["resultado"] == "passa"
    assert r[("RS0005", "programa")]["resultado"] == "passa"      # "Tax amnesty" no título
    assert r[("RS0009", "programa")]["resultado"] == "passa"      # truncamento amnest*


# ---------------------------------------------------------------------------
# Regras de exclusão (exit 2 sem previsão/validação/justificativa)
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("filtro, topo", [
    ({"tipo": "ano", "min": 2001, "modo": "excluir"}, {}),
    ({"tipo": "dicionario", "modo": "excluir", "previsto_no_protocolo": True}, {}),
    ({"tipo": "idioma", "aceitar": ["pt", "en"], "modo": "excluir"}, {"previsto_no_protocolo": True}),
    ({"tipo": "tipo", "aceitar": ["artigo"], "modo": "excluir"}, {"previsto_no_protocolo": True}),
    ({"tipo": "tipo", "recusar": ["tese"], "modo": "excluir", "previsto_no_protocolo": True}, {}),
])
def test_exclusao_sem_salvaguarda_da_exit_2(projeto, filtro, topo):
    from rslib import esquema
    assert rodar(projeto, "--config", escrever_config(projeto, [filtro], **topo)) == 2
    assert not (projeto / esquema.ARQ_FILTRO_FORMAL).exists()


def test_exclusao_prevista_de_editorial_nao_exige_justificativa(projeto):
    cfg = escrever_config(projeto, [{"tipo": "tipo", "recusar": ["editorial", "errata"], "modo": "excluir"}],
                          previsto_no_protocolo=True)
    assert rodar(projeto, "--config", cfg) == 0
    assert resultados(projeto)[("RS0007", "tipo")]["resultado"] == "exclui"


def test_exclusao_sequencial_e_sem_resumo_nunca_sai(projeto, capsys):
    cfg = escrever_config(projeto, [
        {"nome": "ano", "tipo": "ano", "min": 2001, "modo": "excluir"},
        {"nome": "metodo", "tipo": "dicionario", "modo": "excluir",
         "validacao_elusao": "02-triagem/validacao/elusao_filtro_metodo.csv"},
    ], previsto_no_protocolo=True)
    assert rodar(projeto, "--config", cfg) == 0
    r = resultados(projeto)
    assert r[("RS0005", "ano")]["resultado"] == "exclui"
    assert ("RS0005", "metodo") not in r                           # já saiu do funil
    assert r[("RS0006", "ano")]["resultado"] == "sem_dado" and ("RS0006", "metodo") in r
    assert r[("RS0002", "metodo")]["resultado"] == "exclui"
    assert r[("RS0010", "metodo")]["resultado"] == "sem_dado"       # resumo truncado não é excluído
    resumo = ultima_linha_json(capsys)
    assert resumo["por_filtro"]["ano"] == {"passa": 11, "etiqueta": 0, "exclui": 1, "sem_dado": 1}
    assert resumo["n_entrada"] == 13 and resumo["n_saida"] == 13 - 1 - resumo["por_filtro"]["metodo"]["exclui"]
    contagens = json.loads((projeto / "02-triagem" / "filtro_formal_contagens.json").read_text(encoding="utf-8"))
    assert contagens["filtros"][0]["saida"] == contagens["filtros"][1]["entrada"] == 12


def test_sem_resumo_nunca_excluido_por_dicionario(projeto_vazio):
    from rslib import esquema
    colunas = esquema.COLUNAS_UNICOS
    with open(projeto_vazio / esquema.ARQ_UNICOS, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=colunas, lineterminator="\n")
        w.writeheader()
        w.writerow({**{c: "" for c in colunas}, "id_rs": "RS0001", "titulo": "Um título sem termos de método"})
    cfg = escrever_config(projeto_vazio, [{"tipo": "dicionario", "modo": "excluir", "campos": ["titulo"],
                                           "validacao_elusao": "amostra E1"}], previsto_no_protocolo=True)
    assert rodar(projeto_vazio, "--config", cfg) == 0
    assert resultados(projeto_vazio)[("RS0001", "dicionario")]["resultado"] == "sem_dado"


# ---------------------------------------------------------------------------
# Âncoras
# ---------------------------------------------------------------------------
def test_ancora_excluida_da_exit_2_mas_grava_auditoria(projeto, capsys):
    from rslib import esquema
    cfg = escrever_config(projeto, [{"nome": "ano", "tipo": "ano", "min": 2001, "modo": "excluir"}],
                          previsto_no_protocolo=True)
    assert rodar(projeto, "--config", cfg, "--ancoras", str(DIR / "ancoras.csv")) == 2
    saida = capsys.readouterr()
    assert "ANC2" in saida.err
    resumo = json.loads(saida.out.strip().splitlines()[-1])
    assert resumo["ok"] is False and resumo["ancoras"]["excluida"] == 1 and resumo["ancoras"]["nao_encontrada"] == 1
    assert (projeto / esquema.ARQ_FILTRO_FORMAL).exists()
    ev = [e for e in ler_jsonl(projeto / esquema.ARQ_LOG) if e["evento"] == "filtro_formal"][-1]
    assert ev["dados"]["checagem_ancoras"] == "falhou"
    assert ev["dados"]["ancoras_excluidas"][0]["ids_rs"] == ["RS0005"]
    assert ev["dados"]["ancoras_nao_encontradas"] == ["ANC3"]


def test_ancora_etiquetada_nao_bloqueia(projeto, capsys):
    assert rodar(projeto, "--config", str(DIR / "filtros_etiquetar.json"), "--ancoras", str(DIR / "ancoras.csv")) == 0
    contagens = json.loads((projeto / "02-triagem" / "filtro_formal_contagens.json").read_text(encoding="utf-8"))
    por_id = {a["ancora"]: a for a in contagens["ancoras"]}
    assert por_id["ANC1"]["situacao"] == "mantida" and por_id["ANC1"]["ids_rs"] == ["RS0004"]  # DOI com prefixo/caixa
    assert por_id["ANC2"]["situacao"] == "etiquetada" and por_id["ANC2"]["filtros"] == ["ano"]
    assert por_id["ANC3"]["situacao"] == "nao_encontrada"


# ---------------------------------------------------------------------------
# Erros de uso
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("filtros", [
    [{"tipo": "palavra"}],
    [{"tipo": "ano"}],
    [{"tipo": "ano", "min": "2001"}],
    [{"tipo": "tipo", "aceitar": ["artigo"], "recusar": ["tese"]}],
    [{"tipo": "tipo", "aceitar": ["paper"]}],
    [{"tipo": "dicionario", "grupos": ["inexistente"]}],
    [{"tipo": "dicionario", "dicionario": "nao_existe.csv"}],
    [{"tipo": "dicionario", "campos": ["autores"]}],
    [{"tipo": "ano", "min": 2000, "modo": "apagar"}],
    [{"tipo": "ano", "min": 2000}, {"tipo": "ano", "max": 2020}],
    [],
])
def test_config_invalida_da_exit_1(projeto, filtros):
    assert rodar(projeto, "--config", escrever_config(projeto, filtros)) == 1


def test_sem_registros_unicos_da_exit_1(projeto_vazio, tmp_path_factory):
    assert rodar(projeto_vazio, "--config", str(DIR / "filtros_etiquetar.json")) == 1
    assert rodar(tmp_path_factory.mktemp("sem_projeto"), "--config", str(DIR / "filtros_etiquetar.json")) == 1


def test_erros_de_arquivo(projeto):
    cfg = str(DIR / "filtros_etiquetar.json")
    assert rodar(projeto, "--config", "nao_existe.json") == 1
    (projeto / "ruim.json").write_text("{nao é json", encoding="utf-8")
    assert rodar(projeto, "--config", str(projeto / "ruim.json")) == 1
    assert rodar(projeto, "--config", cfg, "--ancoras", "nao_existe.csv") == 1


# ---------------------------------------------------------------------------
# Integração com dedup
# ---------------------------------------------------------------------------
def test_integra_com_dedup(projeto_vazio):
    from rslib import dedup, esquema
    shutil.copy(FIXTURES / "dedup" / "registros_plantados.csv", projeto_vazio / esquema.ARQ_REGISTROS)
    parser = argparse.ArgumentParser()
    parser.add_argument("--dir")
    sub = parser.add_subparsers()
    dedup.registrar(sub)
    args = parser.parse_args(["--dir", str(projeto_vazio), "dedup"])
    assert args.func(args) == 0
    assert rodar(projeto_vazio, "--config", str(DIR / "filtros_etiquetar.json")) == 0
    unicos = ler(projeto_vazio / esquema.ARQ_UNICOS)
    linhas = ler(projeto_vazio / esquema.ARQ_FILTRO_FORMAL)
    assert len(linhas) == 4 * len(unicos)
    assert {l["id_rs"] for l in linhas} == {u["id_rs"] for u in unicos}


# ---------------------------------------------------------------------------
# Conjunto ativo, recall relativo das âncoras e elusão do dicionário (v1.1)
# ---------------------------------------------------------------------------
def escrever_unicos(raiz, linhas):
    from rslib import esquema
    with open(raiz / esquema.ARQ_UNICOS, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=esquema.COLUNAS_UNICOS, lineterminator="\n")
        w.writeheader()
        for l in linhas:
            w.writerow({c: l.get(c, "") for c in esquema.COLUNAS_UNICOS})


def test_clusters_de_busca_inativa_ficam_fora_do_funil(projeto, capsys):
    from rslib import esquema
    linhas = ler(projeto / esquema.ARQ_UNICOS)
    linhas[2]["flags"] = "busca_inativa"  # RS0003
    escrever_unicos(projeto, linhas)
    assert rodar(projeto, "--config", str(DIR / "filtros_etiquetar.json")) == 0
    resumo = ultima_linha_json(capsys)
    assert resumo["n_entrada"] == 12 and resumo["n_inativos_ignorados"] == 1
    assert "RS0003" not in {l["id_rs"] for l in ler(projeto / esquema.ARQ_FILTRO_FORMAL)}
    ev = [e for e in ler_jsonl(projeto / esquema.ARQ_LOG) if e["evento"] == "filtro_formal"][-1]
    assert ev["dados"]["n_inativos_ignorados"] == 1


def projeto_recall(raiz):
    """Três buscas de bases (B01 scopus, B02 wos, B03 substituída), uma de citação (SN1)."""
    from rslib import esquema, estado
    base = {c: "" for c in esquema.COLUNAS_UNICOS}
    escrever_unicos(raiz, [
        dict(base, id_rs="RS0001", ids_registro="B01-00001|B02-00001", doi="10.5555/a1", titulo="Anchor one", ano="2019"),
        dict(base, id_rs="RS0002", ids_registro="B01-00002", doi="10.5555/a2", titulo="Anchor two", ano="2020"),
        dict(base, id_rs="RS0003", ids_registro="SN1-00001", doi="10.5555/a3", titulo="Anchor three", ano="2018"),
        dict(base, id_rs="RS0004", ids_registro="B02-00002", titulo="Not an anchor", ano="2017"),
        dict(base, id_rs="RS0005", ids_registro="B03-00001", doi="10.5555/a5", titulo="Only in replaced search",
             ano="2016", flags="busca_inativa"),
    ])
    reg = {c: "" for c in esquema.COLUNAS_REGISTROS}
    with open(raiz / esquema.ARQ_REGISTROS, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=esquema.COLUNAS_REGISTROS, lineterminator="\n")
        w.writeheader()
        for rid, fonte, metodo in [("B01-00001", "scopus", "base"), ("B01-00002", "scopus", "base"),
                                   ("B02-00001", "wos", "base"), ("B02-00002", "wos", "base"),
                                   ("SN1-00001", "openalex", "citacao"), ("B03-00001", "scopus", "base")]:
            w.writerow(dict(reg, id_registro=rid, busca_id=rid.split("-")[0], fonte=fonte, metodo_identificacao=metodo))
    est = estado.carregar_estado(raiz)
    est["buscas"] = [{"id": "B01", "fonte": "scopus", "metodo_identificacao": "base"},
                     {"id": "B02", "fonte": "wos", "metodo_identificacao": "base"},
                     {"id": "B03", "fonte": "scopus", "metodo_identificacao": "base", "ativa": False},
                     {"id": "SN1", "fonte": "openalex", "metodo_identificacao": "citacao"}]
    estado.salvar_estado(raiz, est)
    ancoras = raiz / "00-protocolo" / "ancoras_validacao.csv"
    ancoras.write_text(
        "id,doi,titulo,ano,indexada_em\n"
        "A1,10.5555/a1,,,B01|B02\n"
        "A2,10.5555/a2,,,scopus|wos\n"
        "A3,10.5555/a3,,,B02\n"
        "A4,,Missing anchor,2015,B01\n"
        "A5,10.5555/a5,,,nenhuma\n"
        "A6,,,,B01\n", encoding="utf-8")
    return ancoras


def test_recall_ancoras_por_busca_combinado_e_por_base(projeto_vazio, capsys):
    from rslib import esquema
    raiz = projeto_vazio
    ancoras = projeto_recall(raiz)
    cfg = escrever_config(raiz, [{"nome": "ano", "tipo": "ano", "min": 1900}])
    assert rodar(raiz, "--config", cfg, "--ancoras", str(ancoras)) == 0
    saida = capsys.readouterr()
    resumo = json.loads(saida.out.strip().splitlines()[-1])
    assert "Anchor" not in saida.out  # imprime ids, nunca títulos
    rec = json.loads((raiz / "01-busca" / "recall_ancoras.json").read_text(encoding="utf-8"))
    por = {a["ancora"]: a for a in rec["ancoras"]}
    assert por["A1"]["buscas"] == ["B01", "B02"] and por["A3"]["buscas"] == ["SN1"]
    assert por["A5"]["situacao"] == "nao_encontrada" and por["A5"]["buscas"] == []  # cluster inativo não conta
    assert rec["tem_indexada_em"] and rec["nao_indexadas"] == ["A5"] and rec["n_sem_chave_de_casamento"] == 1
    # combinado: só buscas de método base (B01, B02); A3 só achada por citação
    comb = rec["combinado"]
    assert (comb["n_ancoras"], comb["n_encontradas"], comb["perdidas"]) == (4, 2, ["A3", "A4"])
    assert comb["recall"] == 0.5 and comb["buscas"] == ["B01", "B02"] and len(comb["ic95"]) == 2
    assert rec["combinado_todos_metodos"]["n_encontradas"] == 3
    # por busca: denominador = âncoras indexadas naquela busca (id ou fonte); B03 inativa não aparece
    assert set(rec["por_busca"]) == {"B01", "B02", "SN1"}
    b01, b02 = rec["por_busca"]["B01"], rec["por_busca"]["B02"]
    assert (b01["n_ancoras"], b01["n_encontradas"], b01["perdidas"]) == (3, 2, ["A4"])  # A1, A2 (scopus), A4
    assert (b02["n_ancoras"], b02["n_encontradas"], b02["perdidas"]) == (3, 1, ["A2", "A3"])
    assert b01["fonte"] == "scopus" and rec["por_busca"]["SN1"]["metodo"] == "citacao"
    assert rec["por_base"]["scopus"]["buscas"] == ["B01"] and rec["por_base"]["scopus"]["recall"] == 1.0
    assert rec["por_base"]["B02"]["n_encontradas"] == 1
    assert any("A3" in a for a in rec["avisos"])  # achada fora do declarado em indexada_em
    assert resumo["recall_ancoras"]["combinado"] == 0.5 and "01-busca/recall_ancoras.json" in resumo["arquivos"]
    ev = [e for e in ler_jsonl(raiz / esquema.ARQ_LOG) if e["evento"] == "filtro_formal"][-1]
    assert ev["dados"]["recall_ancoras"]["por_busca"]["B02"] == round(1 / 3, 4)
    assert "01-busca/recall_ancoras.json" in {a["caminho"] for a in ev["artefatos"]}
    # idempotente
    antes = (raiz / "01-busca" / "recall_ancoras.json").read_bytes()
    assert rodar(raiz, "--config", cfg, "--ancoras", str(ancoras)) == 0
    assert (raiz / "01-busca" / "recall_ancoras.json").read_bytes() == antes


def test_recall_sem_indexada_em_usa_ids_registro(projeto, capsys):
    assert rodar(projeto, "--config", str(DIR / "filtros_etiquetar.json"), "--ancoras", str(DIR / "ancoras.csv")) == 0
    rec = json.loads((projeto / "01-busca" / "recall_ancoras.json").read_text(encoding="utf-8"))
    assert rec["tem_indexada_em"] is False and rec["por_base"] is None
    assert list(rec["por_busca"]) == ["B01"] and rec["por_busca"]["B01"]["metodo"] == "base"
    assert (rec["combinado"]["n_ancoras"], rec["combinado"]["n_encontradas"]) == (3, 2)
    assert rec["combinado"]["perdidas"] == ["ANC3"]


def preparar_elusao(projeto):
    cfg = projeto / "01-busca" / "filtros_v1.json"
    shutil.copy(DIR / "filtros_etiquetar.json", cfg)
    return cfg


def ler_planilha(caminho):
    from openpyxl import load_workbook
    wb = load_workbook(caminho)
    ws = wb["codificacao"]
    linhas = list(ws.iter_rows(values_only=True))
    return wb, ws, linhas


def test_amostra_elusao_cega_e_reprodutivel(projeto, capsys):
    pytest.importorskip("openpyxl")
    from rslib import esquema
    cfg = preparar_elusao(projeto)
    assert rodar(projeto, "--config", str(cfg), "--amostra-elusao", "3", "--semente", "7") == 0
    resumo = ultima_linha_json(capsys)
    desenho = json.loads((projeto / "02-triagem/validacao/elusao_filtros_v1_desenho.json").read_text(encoding="utf-8"))
    etiquetados = {l["id_rs"] for l in ler(projeto / esquema.ARQ_FILTRO_FORMAL)
                   if l["filtro"] == "metodo" and l["resultado"] == "etiqueta"}
    assert desenho["populacao_n"] == len(etiquetados) == 7 and desenho["n_amostra"] == 3
    assert set(desenho["ids"]) <= etiquetados and desenho["filtros"][0]["nome"] == "metodo"
    planilha = projeto / "02-triagem/validacao/elusao_filtros_v1_cega.xlsx"
    wb, ws, linhas = ler_planilha(planilha)
    from rslib.filtrar import COLUNAS_CODIGO_ELUSAO, COLUNAS_REGISTRO_ELUSAO
    assert list(linhas[0]) == COLUNAS_REGISTRO_ELUSAO + COLUNAS_CODIGO_ELUSAO + ["observacoes"]
    assert not {"filtro", "resultado", "detalhe", "flags", "termos"} & set(linhas[0])  # cega
    assert [r[1] for r in linhas[1:]] == desenho["ids"] and all(r[9] is None for r in linhas[1:])
    assert resumo["amostra_elusao"]["reutilizada"] is False and resumo["amostra_elusao"]["populacao_n"] == 7
    ev = [e for e in ler_jsonl(projeto / esquema.ARQ_LOG) if e["evento"] == "filtro_formal"][-1]
    assert "02-triagem/validacao/elusao_filtros_v1_cega.xlsx" in {a["caminho"] for a in ev["artefatos"]}
    # mesma semente: não sorteia de novo nem sobrescreve a planilha em codificação
    ws.cell(row=2, column=10, value="incluir")
    wb.save(planilha)
    conteudo = planilha.read_bytes()
    assert rodar(projeto, "--config", str(cfg), "--amostra-elusao", "3", "--semente", "7") == 0
    assert ultima_linha_json(capsys)["amostra_elusao"]["reutilizada"] is True
    assert planilha.read_bytes() == conteudo
    # outros parâmetros para a mesma versão: recusa
    assert rodar(projeto, "--config", str(cfg), "--amostra-elusao", "3", "--semente", "8") == 1
    assert rodar(projeto, "--config", str(cfg), "--amostra-elusao", "3") == 1           # sem semente
    assert rodar(projeto, "--config", str(cfg), "--semente", "3") == 1                  # semente sozinha
    sem_dic = escrever_config(projeto, [{"tipo": "ano", "min": 2001}])
    assert rodar(projeto, "--config", sem_dic, "--amostra-elusao", "3", "--semente", "7") == 1


def codificar(planilha, decisoes):
    wb, ws, linhas = ler_planilha(planilha)
    for i, d in enumerate(decisoes, start=2):
        for col, valor in zip((10, 11, 12), d):
            ws.cell(row=i, column=col, value=valor)
    wb.save(planilha)
    return [r[1] for r in linhas[1:]]


def test_calcular_elusao_ic_evento_e_habilita_exclusao(projeto, capsys):
    pytest.importorskip("openpyxl")
    from rslib import esquema
    from rslib.validacao import clopper_pearson
    cfg = preparar_elusao(projeto)
    assert rodar(projeto, "--config", str(cfg), "--amostra-elusao", "4", "--semente", "11") == 0
    planilha = projeto / "02-triagem/validacao/elusao_filtros_v1_cega.xlsx"
    # h1 incluir; h1 excluir; h1 incerto x h2 excluir sem consenso (fora); vazio (não revisado)
    ids = codificar(planilha, [("incluir", None, None), ("excluir", "excluir", None), ("incerto", "excluir", None),
                               (None, None, None)])
    capsys.readouterr()
    assert rodar(projeto, "--calcular-elusao", str(planilha)) == 0
    resumo = ultima_linha_json(capsys)
    lo, hi = clopper_pearson(1, 2)
    assert resumo["n_revisados"] == 2 and resumo["n_nao_revisados"] == 2 and resumo["taxa_elusao"] == 0.5
    assert resumo["ic95"] == [round(lo, 4), round(hi, 4)] and resumo["perdidos_estimados"] == 3.5
    metricas = json.loads((projeto / "02-triagem/validacao/elusao_filtros_v1_metricas.json").read_text(encoding="utf-8"))
    assert metricas["n_discordancias_sem_consenso"] == 1 and metricas["ids_relevantes"] == [ids[0]]
    ev = [e for e in ler_jsonl(projeto / esquema.ARQ_LOG) if e["evento"] == "validacao_calculada"][-1]
    assert ev["etapa"] == "05_organizacao" and ev["dados"]["finalidade"] == "elusao"
    assert ev["dados"]["finalidade"] in esquema.FINALIDADES_VALIDACAO and "atende_limiares" not in ev["dados"]
    # habilita validacao_elusao na configuração (sem mudar o modo)
    config = json.loads(cfg.read_text(encoding="utf-8"))
    metodo = next(f for f in config["filtros"] if f["nome"] == "metodo")
    assert metodo["validacao_elusao"] == "02-triagem/validacao/elusao_filtros_v1_metricas.json"
    assert metodo.get("modo", "etiquetar") == "etiquetar" and resumo["config_habilitada"]["atualizada"]
    # com previsão no protocolo, a exclusão por dicionário passa e registra a validação verificada
    metodo.update(modo="excluir", previsto_no_protocolo=True)
    cfg.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    assert rodar(projeto, "--config", str(cfg)) == 0
    ev = [e for e in ler_jsonl(projeto / esquema.ARQ_LOG) if e["evento"] == "filtro_formal"][-1]
    assert ev["dados"]["validacoes_elusao"]["metodo"]["taxa_elusao"] == 0.5
    # definição do filtro mudou depois da amostra: a validação não vale mais (exit 2)
    metodo["campos"] = ["titulo"]
    cfg.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    assert rodar(projeto, "--config", str(cfg)) == 2
    # métrica de outro filtro também não vale
    metodo["campos"] = ["titulo", "resumo", "palavras_chave"]
    metodo["nome"] = "outro_dicionario"
    cfg.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    assert rodar(projeto, "--config", str(cfg)) == 2


def test_calcular_elusao_nao_altera_config_congelada_nem_fora_do_projeto(projeto, capsys):
    pytest.importorskip("openpyxl")
    from rslib import estado
    cfg = preparar_elusao(projeto)
    assert rodar(projeto, "--config", str(cfg), "--amostra-elusao", "2", "--semente", "1") == 0
    planilha = projeto / "02-triagem/validacao/elusao_filtros_v1_cega.xlsx"
    codificar(planilha, [("excluir", None, None), ("excluir", None, None)])
    est = estado.carregar_estado(projeto)
    est["artefatos"]["01-busca/filtros_v1.json"] = {"caminho": "01-busca/filtros_v1.json", "sha256": "x",
                                                    "congelado_em": "2026-01-01T00:00:00Z"}
    estado.salvar_estado(projeto, est)
    antes = cfg.read_bytes()
    capsys.readouterr()
    assert rodar(projeto, "--calcular-elusao", str(planilha)) == 0
    resumo = ultima_linha_json(capsys)
    assert resumo["taxa_elusao"] == 0.0 and resumo["config_habilitada"]["atualizada"] is False
    assert "congelada" in resumo["config_habilitada"]["motivo"] and cfg.read_bytes() == antes
    # configuração fora do projeto (fixture): nunca é editada
    outra = projeto / "02-triagem/validacao"
    for arq in outra.glob("elusao_*"):
        arq.unlink()
    assert rodar(projeto, "--config", str(DIR / "filtros_etiquetar.json"), "--amostra-elusao", "2", "--semente", "1") == 0
    codificar(planilha, [("excluir", None, None), ("incluir", None, None)])
    fixture = (DIR / "filtros_etiquetar.json").read_bytes()
    capsys.readouterr()
    assert rodar(projeto, "--calcular-elusao", str(planilha)) == 0
    assert "fora do projeto" in ultima_linha_json(capsys)["config_habilitada"]["motivo"]
    assert (DIR / "filtros_etiquetar.json").read_bytes() == fixture


def test_calcular_elusao_erros(projeto, capsys):
    pytest.importorskip("openpyxl")
    cfg = preparar_elusao(projeto)
    assert rodar(projeto, "--config", str(cfg), "--amostra-elusao", "2", "--semente", "3") == 0
    planilha = projeto / "02-triagem/validacao/elusao_filtros_v1_cega.xlsx"
    assert rodar(projeto, "--calcular-elusao", str(planilha)) == 1                   # nada codificado
    codificar(planilha, [("talvez sim", None, None), ("excluir", None, None)])
    assert rodar(projeto, "--calcular-elusao", str(planilha)) == 1                   # código ilegível
    wb, ws, _ = ler_planilha(planilha)
    ws.cell(row=2, column=2, value="RS9999")
    ws.cell(row=2, column=10, value="incluir")
    wb.save(planilha)
    assert rodar(projeto, "--calcular-elusao", str(planilha)) == 1                   # id fora da amostra
    assert rodar(projeto, "--calcular-elusao", str(planilha), "--config", str(cfg)) == 1
    assert rodar(projeto, "--calcular-elusao", "nao_existe.xlsx") == 1
    assert rodar(projeto) == 1                                                        # sem --config


# ---------------------------------------------------------------------------
# v1.3: âncoras e planilha de elusão pelo leitor único; indexada_em por nome da base; permissões
# ---------------------------------------------------------------------------
def test_ancoras_do_excel_cp1252_com_nome_da_base_e_id_de_busca_substituida(projeto_vazio, capsys):
    from rslib import estado
    raiz = projeto_vazio
    projeto_recall(raiz)
    est = estado.carregar_estado(raiz)
    # B03 foi substituída por B01 (mesma base, string nova): o id antigo no arquivo congelado ainda conta
    next(b for b in est["buscas"] if b["id"] == "B03")["substituida_por"] = "B01"
    estado.salvar_estado(raiz, est)
    ancoras = raiz / "00-protocolo" / "ancoras_validacao.csv"
    ancoras.write_bytes((
        "id;doi;titulo;ano;indexada_em;observação\r\n"
        "A1;10.5555/a1;;;scopus|wos;conferida por revisão\r\n"
        "A2;10.5555/a2;;;B03;id antigo, antes da substituição\r\n"
        "A4;;Missing anchor;2015;scopus;não achada\r\n").encode("cp1252"))
    cfg = escrever_config(raiz, [{"nome": "ano", "tipo": "ano", "min": 1900}])
    assert rodar(raiz, "--config", cfg, "--ancoras", str(ancoras)) == 0
    rec = json.loads((raiz / "01-busca" / "recall_ancoras.json").read_text(encoding="utf-8"))
    assert [a["ancora"] for a in rec["ancoras"]] == ["A1", "A2", "A4"]
    assert rec["por_base"]["B03"]["buscas"] == ["B01"] and rec["por_base"]["B03"]["recall"] == 1.0
    assert rec["por_base"]["scopus"]["n_ancoras"] == 2 and rec["por_base"]["scopus"]["perdidas"] == ["A4"]
    assert any("'B03': busca substituída por B01" in a and "nome da base" in a for a in rec["avisos"])
    assert not any("'B03': inativa" in a for a in rec["avisos"])
    assert any("cp1252" in a for a in rec["avisos"])
    assert (rec["combinado"]["n_ancoras"], rec["combinado"]["n_encontradas"]) == (3, 2)


def test_modelo_de_ancoras_recomenda_nome_da_base():
    modelo = (SKILL / "assets" / "templates" / "ancoras.csv").read_text(encoding="utf-8")
    assert "openalex|scopus|wos" in modelo and "B01" not in modelo


def test_calcular_elusao_de_csv_do_excel(projeto, capsys):
    pytest.importorskip("openpyxl")
    cfg = preparar_elusao(projeto)
    assert rodar(projeto, "--config", str(cfg), "--amostra-elusao", "3", "--semente", "7") == 0
    desenho_rel = "02-triagem/validacao/elusao_filtros_v1_desenho.json"
    desenho = json.loads((projeto / desenho_rel).read_text(encoding="utf-8"))
    planilha = projeto / "02-triagem/validacao/elusao_codificada.csv"
    linhas = ["ordem;ID_RS;título;decisao_h1;decisao_h2;decisao_consenso"]
    for k, (i, d) in enumerate(zip(desenho["ids"], ["Incluir", "excluir", "excluir"]), start=1):
        linhas.append(f"{k};{i.lower()};Título com acentuação;{d};;")
    planilha.write_bytes(("\r\n".join(linhas) + "\r\n").encode("cp1252"))
    capsys.readouterr()
    assert rodar(projeto, "--calcular-elusao", str(planilha), "--desenho", desenho_rel) == 0
    resumo = ultima_linha_json(capsys)
    assert resumo["n_revisados"] == 3 and resumo["taxa_elusao"] == round(1 / 3, 4)


@pytest.mark.skipif(__import__("os").name == "nt", reason="permissões POSIX")
def test_regressao_saidas_do_filtrar_com_permissao_de_arquivo_comum(projeto):
    import stat
    from rslib import esquema, estado
    assert rodar(projeto, "--config", str(DIR / "filtros_etiquetar.json")) == 0
    for rel in (esquema.ARQ_FILTRO_FORMAL, "02-triagem/filtro_formal_contagens.json"):
        assert stat.S_IMODE((projeto / rel).stat().st_mode) == estado.modo_arquivo_padrao(), rel
